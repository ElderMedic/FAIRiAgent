#!/usr/bin/env python3
"""Merge gate for FAIRiAgent auto retrieval/repair eval results.

The gate compares a fresh `auto` eval run against the current Shadow and Tuned
baselines. It intentionally fails when the auto result is missing: that is the
correct state before a full six-document eval has been run.
"""

from __future__ import annotations

import argparse
import json
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from fairifier.validation.metadata_json_format import (
    validate_field_datatypes,
    validate_json_structure,
    validate_source_grounding,
    validate_value_formats,
)


DEFAULT_ROOT = Path("evaluation/prototypes/auto_repair_classifier")
DEFAULT_SHADOW = Path("evaluation/runs/shadow_pro_tuned/results/evaluation_results.json")
DEFAULT_TUNED = Path("evaluation/runs/phase4_pro_tuned/results/evaluation_results.json")
DEFAULT_AUTO = Path("evaluation/runs/auto_pro_tuned/results/evaluation_results.json")


KEY_METRICS = [
    "schema_compliance",
    "row_alignment_f1",
    "sheet_placement_accuracy",
    "precision_excl_discoveries",
    "completeness",
    "value_match_rate",
    "llm_judge_score",
    "aggregate_score",
]

SINGLE_ROW_SHEETS = {"investigation", "study"}
FAIRDS_EXCEL_FILENAME = "metadata_fairds.xlsx"


def normalize_name(value: Any) -> str:
    return " ".join(str(value or "").strip().lower().replace("_", " ").split())


@dataclass(frozen=True)
class GateThresholds:
    min_common_documents: int = 6
    schema_tolerance: float = 0.01
    row_alignment_tolerance: float = 0.02
    sheet_placement_tolerance: float = 0.01
    precision_tolerance: float = 0.02
    min_precision_excl_discoveries: float = 0.97
    completeness_tolerance: float = 0.01
    value_match_tolerance_vs_tuned: float = 0.03


def resolve_run_dir(results_path: Path) -> Path:
    if results_path.name == "evaluation_results.json" and results_path.parent.name == "results":
        return results_path.parent.parent
    return results_path.parent


def load_json(path: Path) -> Dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        data = json.load(handle)
    if not isinstance(data, dict):
        raise ValueError(f"{path} did not contain a JSON object")
    return data


def first_model_name(results: Dict[str, Any]) -> str:
    per_model = results.get("per_model_results") or {}
    if not per_model:
        raise ValueError("evaluation_results.json has no per_model_results")
    return next(iter(per_model.keys()))


def first_model_result(results: Dict[str, Any]) -> Dict[str, Any]:
    return (results.get("per_model_results") or {})[first_model_name(results)]


def discover_model_run_root(run_dir: Path) -> Optional[Path]:
    if not run_dir.exists():
        return None
    model_name_candidates: List[Path] = []
    for child in run_dir.iterdir():
        if not child.is_dir() or child.name in {"results", "outputs", "artifacts"}:
            continue
        if any(child.glob("*/run_*/metadata.json")) or any(
            child.glob("*/run_*/metadata_json.json")
        ):
            model_name_candidates.append(child)
    if len(model_name_candidates) == 1:
        return model_name_candidates[0]
    if len(model_name_candidates) > 1:
        return sorted(model_name_candidates)[0]
    return None


def first_model_metrics(results: Dict[str, Any]) -> Dict[str, float]:
    metrics = (results.get("model_comparison") or {}).get("metrics") or {}
    if metrics:
        raw = next(iter(metrics.values()))
        return {str(k): as_float(v) for k, v in raw.items()}
    model = first_model_result(results)
    out: Dict[str, float] = {}
    for key in KEY_METRICS:
        value = model.get(key)
        if isinstance(value, (int, float)):
            out[key] = float(value)
    return out


def as_float(value: Any, default: float = 0.0) -> float:
    try:
        if value in (None, ""):
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


def nested_float(data: Dict[str, Any], path: Iterable[str], default: float = 0.0) -> float:
    current: Any = data
    for key in path:
        if not isinstance(current, dict):
            return default
        current = current.get(key)
    return as_float(current, default)


def doc_ids(model: Dict[str, Any]) -> List[str]:
    completeness = (model.get("completeness") or {}).get("per_document") or {}
    return sorted(str(doc_id) for doc_id in completeness.keys())


def doc_metrics(model: Dict[str, Any], doc_id: str) -> Dict[str, float]:
    completeness = (
        ((model.get("completeness") or {}).get("per_document") or {})
        .get(doc_id, {})
        .get("overall_metrics", {})
    )
    structural = (
        ((model.get("structural") or {}).get("per_document") or {})
        .get(doc_id, {})
        .get("summary_metrics", {})
    )
    schema = ((model.get("schema_validation") or {}).get("per_document") or {}).get(
        doc_id,
        {},
    )
    value = ((model.get("value_accuracy") or {}).get("per_document") or {}).get(
        doc_id,
        {},
    )
    return {
        "completeness": nested_float(completeness, ["overall_completeness"]),
        "required_completeness": nested_float(completeness, ["required_completeness"]),
        "extra_fields": nested_float(completeness, ["extra_fields"]),
        "schema_compliance": nested_float(schema, ["schema_compliance_rate"]),
        "sheet_placement_accuracy": nested_float(
            structural,
            ["sheet_placement_accuracy"],
        ),
        "row_alignment_f1": nested_float(structural, ["row_alignment_f1"]),
        "value_match_rate": nested_float(value, ["summary", "value_match_rate"]),
    }


def find_doc_run_dir(model_root: Path, doc_id: str) -> Optional[Path]:
    direct = model_root / doc_id / "run_1"
    if direct.exists():
        return direct
    matches = sorted(model_root.glob(f"**/{doc_id}/run_*"))
    return matches[0] if matches else None


def read_json_file(path: Path) -> Tuple[bool, str, Dict[str, Any]]:
    try:
        data = load_json(path)
    except FileNotFoundError:
        return False, "missing", {}
    except (OSError, json.JSONDecodeError, ValueError) as exc:
        return False, f"parse_error: {exc}", {}
    return True, "ok", data


def validate_metadata_file_format(metadata: Dict[str, Any]) -> Dict[str, Any]:
    errors: List[str] = []
    warnings: List[str] = []
    validate_json_structure(metadata, errors, warnings)
    validate_field_datatypes(metadata, errors, warnings)
    validate_value_formats(metadata, errors, warnings)
    source_grounding = validate_source_grounding(metadata, errors, warnings)
    return {
        "passed": not errors,
        "errors": errors,
        "warnings": warnings[:20],
        "source_grounding": source_grounding,
    }


def validate_runtime_config(runtime_config: Dict[str, Any]) -> List[str]:
    errors: List[str] = []
    cfg = runtime_config.get("config")
    if not isinstance(cfg, dict):
        return ["runtime_config_missing_config_section"]

    effective_mode = str(cfg.get("effective_retrieval_mode") or "").strip().lower()
    if effective_mode != "auto":
        errors.append(f"runtime_config_effective_mode_not_auto:{effective_mode or 'missing'}")

    retrieval_mode = str(cfg.get("retrieval_mode") or "").strip().lower()
    if retrieval_mode and retrieval_mode not in {"auto", "shadow", "tuned"}:
        errors.append(f"runtime_config_invalid_retrieval_mode:{retrieval_mode}")

    if cfg.get("auto_repair_enabled") is not True:
        errors.append("runtime_config_auto_repair_not_enabled")
    if cfg.get("auto_repair_apply_patches") is not True:
        errors.append("runtime_config_auto_repair_not_applying_patches")
    if not isinstance(cfg.get("auto_repair_classifier_shadow_enabled"), bool):
        errors.append("runtime_config_classifier_shadow_setting_missing")

    return errors


def validate_auto_repair_trace(trace: Dict[str, Any]) -> List[str]:
    errors: List[str] = []
    mode = str(trace.get("mode") or "").strip()
    if mode != "deterministic_exact_patch":
        errors.append(f"auto_repair_trace_unexpected_mode:{mode or 'missing'}")

    summary = trace.get("summary")
    if not isinstance(summary, dict):
        return errors + ["auto_repair_trace_missing_summary"]
    if summary.get("apply_patches") is not True:
        errors.append("auto_repair_trace_not_applying_patches")
    for key in ("candidate_count", "accepted_patch_count", "rejected_patch_count"):
        value = summary.get(key)
        if value is not None and not isinstance(value, int):
            errors.append(f"auto_repair_trace_non_integer_{key}")
    return errors


def accepted_field_records(
    *,
    metadata: Dict[str, Any],
    trace: Dict[str, Any],
) -> List[Dict[str, Any]]:
    summary = metadata.get("auto_repair_summary")
    if isinstance(summary, dict) and isinstance(summary.get("accepted_fields"), list):
        return [item for item in summary["accepted_fields"] if isinstance(item, dict)]
    accepted_patches = trace.get("accepted_patches")
    if isinstance(accepted_patches, list):
        return [item for item in accepted_patches if isinstance(item, dict)]
    return []


def sheet_payload(container: Dict[str, Any], sheet: str) -> Any:
    for key, value in container.items():
        if normalize_name(key).replace(" ", "") == normalize_name(sheet).replace(" ", ""):
            return value
    return None


def field_payload_value(field: Dict[str, Any]) -> Any:
    return field.get("value") if "value" in field else field.get("extracted_value")


def metadata_field_present(metadata: Dict[str, Any], sheet: str, field_name: str) -> bool:
    isa_structure = metadata.get("isa_structure")
    if not isinstance(isa_structure, dict):
        return False
    payload = sheet_payload(isa_structure, sheet)
    if not isinstance(payload, dict):
        return False
    fields = payload.get("fields")
    if not isinstance(fields, list):
        return False
    target = normalize_name(field_name)
    for item in fields:
        if not isinstance(item, dict):
            continue
        name = item.get("field_name") or item.get("name") or item.get("field")
        if normalize_name(name) == target and field_payload_value(item) not in (None, ""):
            return True
    return False


def single_row_value_present(container: Dict[str, Any], sheet: str, field_name: str) -> bool:
    payload = sheet_payload(container, sheet)
    if not isinstance(payload, dict):
        return False
    target = normalize_name(field_name)
    columns = payload.get("columns")
    if isinstance(columns, list) and target not in {normalize_name(item) for item in columns}:
        return False
    rows = payload.get("rows")
    if not isinstance(rows, list) or len(rows) != 1 or not isinstance(rows[0], dict):
        return False
    for key, value in rows[0].items():
        if normalize_name(key) == target and value not in (None, ""):
            return True
    return False


def excel_field_value_present(excel_path: Path, sheet: str, field_name: str) -> bool:
    from openpyxl import load_workbook

    workbook = load_workbook(excel_path, data_only=True, read_only=True)
    try:
        sheet_key = normalize_name(sheet).replace(" ", "")
        target_field = normalize_name(field_name)
        worksheet = None
        for candidate in workbook.worksheets:
            if normalize_name(candidate.title).replace(" ", "") == sheet_key:
                worksheet = candidate
                break
        if worksheet is None:
            return False

        header_to_col: Dict[str, int] = {}
        for col in range(1, worksheet.max_column + 1):
            header = worksheet.cell(1, col).value
            if header not in (None, ""):
                header_to_col[normalize_name(header)] = col
        col_idx = header_to_col.get(target_field)
        if col_idx is None:
            return False

        for row in range(2, worksheet.max_row + 1):
            value = worksheet.cell(row, col_idx).value
            if value not in (None, ""):
                return True
        return False
    finally:
        workbook.close()


def validate_fairds_excel_artifact(
    *,
    run_dir: Path,
    metadata: Dict[str, Any],
    trace: Dict[str, Any],
) -> Tuple[str, List[str]]:
    excel_path = run_dir / FAIRDS_EXCEL_FILENAME
    if not excel_path.exists():
        return "missing", ["metadata_fairds_xlsx_missing"]

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(excel_path, data_only=True, read_only=True)
        workbook.close()
    except Exception as exc:
        return f"parse_error: {exc}", [f"metadata_fairds_xlsx_parse_error:{exc}"]

    errors: List[str] = []
    summary = trace.get("summary") if isinstance(trace.get("summary"), dict) else {}
    accepted = int(summary.get("accepted_patch_count") or 0)
    if accepted > 0:
        records = accepted_field_records(metadata=metadata, trace=trace)
        for item in records:
            field_name = str(item.get("field") or item.get("field_name") or "").strip()
            sheet = str(item.get("isa_sheet") or "study").strip() or "study"
            if not field_name:
                continue
            try:
                present = excel_field_value_present(excel_path, sheet, field_name)
            except Exception as exc:
                return (
                    f"parse_error: {exc}",
                    [f"metadata_fairds_xlsx_parse_error:{exc}"],
                )
            if not present:
                errors.append(f"accepted_patch_missing_excel_value:{sheet}:{field_name}")
    return "ok", errors


def validate_accepted_patch_materialization(
    *,
    metadata: Dict[str, Any],
    isa_values: Dict[str, Any],
    trace: Dict[str, Any],
) -> List[str]:
    errors: List[str] = []
    summary = trace.get("summary") if isinstance(trace.get("summary"), dict) else {}
    accepted = int(summary.get("accepted_patch_count") or 0)
    if accepted <= 0:
        return errors

    records = accepted_field_records(metadata=metadata, trace=trace)
    if not records:
        return ["metadata_json_missing_auto_repair_accepted_fields"]

    for item in records:
        field_name = str(item.get("field") or item.get("field_name") or "").strip()
        sheet = str(item.get("isa_sheet") or "study").strip() or "study"
        if not field_name:
            errors.append("accepted_patch_missing_field_name")
            continue
        sheet_key = normalize_name(sheet).replace(" ", "")
        if not metadata_field_present(metadata, sheet, field_name):
            errors.append(f"accepted_patch_missing_metadata_field:{sheet}:{field_name}")
        if sheet_key in SINGLE_ROW_SHEETS:
            isa_structure = metadata.get("isa_structure")
            if not isinstance(isa_structure, dict) or not single_row_value_present(
                isa_structure,
                sheet,
                field_name,
            ):
                errors.append(f"accepted_patch_missing_metadata_row_value:{sheet}:{field_name}")
            if not single_row_value_present(isa_values, sheet, field_name):
                errors.append(f"accepted_patch_missing_isa_values_value:{sheet}:{field_name}")
    return errors


def validate_auto_artifacts(
    *,
    auto_run_dir: Optional[Path],
    doc_ids_to_check: List[str],
) -> Dict[str, Any]:
    if auto_run_dir is None:
        return {
            "checked": False,
            "passed": False,
            "reason": "auto_run_dir_not_provided",
            "documents": [],
            "summary": {},
        }
    model_root = discover_model_run_root(auto_run_dir)
    if model_root is None:
        return {
            "checked": False,
            "passed": False,
            "reason": f"model_run_root_not_found_under:{auto_run_dir}",
            "documents": [],
            "summary": {},
        }

    documents: List[Dict[str, Any]] = []
    accepted_patch_count = 0
    mutated_count = 0
    trace_only_count = 0
    non_production_trace_count = 0
    trace_mode_counts: Dict[str, int] = {}
    for doc_id in doc_ids_to_check:
        run_dir = find_doc_run_dir(model_root, doc_id)
        record: Dict[str, Any] = {
            "doc_id": doc_id,
            "run_dir": str(run_dir) if run_dir else None,
            "passed": False,
            "errors": [],
        }
        if run_dir is None:
            record["errors"].append("doc_run_dir_missing")
            documents.append(record)
            continue

        metadata_path = run_dir / "metadata.json"
        if not metadata_path.exists():
            metadata_path = run_dir / "metadata_json.json"
        metadata_ok, metadata_status, metadata = read_json_file(metadata_path)
        record["metadata_json"] = metadata_status
        if not metadata_ok:
            record["errors"].append(f"metadata_json_{metadata_status}")
        else:
            isa_structure = metadata.get("isa_structure")
            if not isinstance(isa_structure, dict) or not isa_structure:
                record["errors"].append("metadata_json_missing_isa_structure")
            format_report = validate_metadata_file_format(metadata)
            record["metadata_json_format"] = (
                "ok" if format_report["passed"] else "failed"
            )
            record["metadata_source_grounding"] = format_report.get("source_grounding", {})
            if not format_report["passed"]:
                record["errors"].extend(
                    f"metadata_json_format_error:{error}"
                    for error in format_report["errors"]
                )
            if format_report.get("warnings"):
                record["metadata_json_warnings"] = format_report["warnings"]

        workflow_ok, workflow_status, _workflow = read_json_file(
            run_dir / "workflow_report.json"
        )
        record["workflow_report"] = workflow_status
        if not workflow_ok:
            record["errors"].append(f"workflow_report_{workflow_status}")

        runtime_ok, runtime_status, runtime_config = read_json_file(
            run_dir / "runtime_config.json"
        )
        record["runtime_config"] = runtime_status
        if not runtime_ok:
            record["errors"].append(f"runtime_config_{runtime_status}")
        else:
            runtime_errors = validate_runtime_config(runtime_config)
            if runtime_errors:
                record["errors"].extend(runtime_errors)
            else:
                cfg = runtime_config.get("config", {})
                record["effective_retrieval_mode"] = cfg.get("effective_retrieval_mode")
                record["auto_repair_apply_patches_config"] = cfg.get(
                    "auto_repair_apply_patches"
                )
                record["auto_repair_classifier_shadow_enabled"] = cfg.get(
                    "auto_repair_classifier_shadow_enabled"
                )

        trace_ok, trace_status, trace = read_json_file(run_dir / "auto_repair_trace.json")
        record["auto_repair_trace"] = trace_status
        if not trace_ok:
            record["errors"].append(f"auto_repair_trace_{trace_status}")
        else:
            trace_errors = validate_auto_repair_trace(trace)
            trace_mode = str(trace.get("mode") or "missing")
            record["auto_repair_trace_mode"] = trace_mode
            trace_mode_counts[trace_mode] = trace_mode_counts.get(trace_mode, 0) + 1
            if trace_errors:
                record["errors"].extend(trace_errors)
                non_production_trace_count += 1
            summary = trace.get("summary") if isinstance(trace.get("summary"), dict) else {}
            accepted = int(summary.get("accepted_patch_count") or 0)
            accepted_patch_count += accepted
            if bool(summary.get("metadata_mutated")):
                mutated_count += 1
            if summary.get("apply_patches") is False:
                trace_only_count += 1
            record["accepted_patch_count"] = accepted
            record["metadata_mutated"] = bool(summary.get("metadata_mutated"))
            record["apply_patches"] = summary.get("apply_patches")
            if accepted > 0 and metadata_ok:
                repair_summary = metadata.get("auto_repair_summary")
                if not isinstance(repair_summary, dict):
                    record["errors"].append("metadata_json_missing_auto_repair_summary")
                elif int(repair_summary.get("accepted_patch_count") or 0) != accepted:
                    record["errors"].append(
                        "metadata_json_auto_repair_summary_count_mismatch"
                    )

        isa_values_path = run_dir / "isa_values_json.json"
        isa_values_ok, isa_values_status, isa_values = read_json_file(isa_values_path)
        record["isa_values_json"] = isa_values_status
        if not isa_values_ok:
            record["errors"].append(f"isa_values_json_{isa_values_status}")
        elif not isinstance(isa_values, dict) or not isa_values:
            record["errors"].append("isa_values_json_empty")
        if trace_ok and metadata_ok and isa_values_ok:
            record["errors"].extend(
                validate_accepted_patch_materialization(
                    metadata=metadata,
                    isa_values=isa_values,
                    trace=trace,
                )
            )
        excel_status = "not_checked"
        if metadata_ok and trace_ok:
            excel_status, excel_errors = validate_fairds_excel_artifact(
                run_dir=run_dir,
                metadata=metadata,
                trace=trace,
            )
            record["errors"].extend(excel_errors)
        record["metadata_fairds_xlsx"] = excel_status

        record["passed"] = not record["errors"]
        documents.append(record)

    failed_documents = [item for item in documents if not item.get("passed")]
    return {
        "checked": True,
        "passed": not failed_documents,
        "model_root": str(model_root),
        "documents": documents,
        "summary": {
            "document_count": len(documents),
            "failed_document_count": len(failed_documents),
            "accepted_patch_count": accepted_patch_count,
            "metadata_mutated_document_count": mutated_count,
            "trace_only_document_count": trace_only_count,
            "non_production_trace_document_count": non_production_trace_count,
            "trace_mode_counts": trace_mode_counts,
        },
    }


def metric_row(
    name: str,
    auto: Dict[str, float],
    shadow: Dict[str, float],
    tuned: Dict[str, float],
) -> Dict[str, float]:
    auto_value = as_float(auto.get(name))
    shadow_value = as_float(shadow.get(name))
    tuned_value = as_float(tuned.get(name))
    return {
        "auto": auto_value,
        "shadow": shadow_value,
        "tuned": tuned_value,
        "delta_vs_shadow": auto_value - shadow_value,
        "delta_vs_tuned": auto_value - tuned_value,
    }


def add_check(
    checks: List[Dict[str, Any]],
    name: str,
    passed: bool,
    detail: str,
    severity: str = "fail",
) -> None:
    checks.append(
        {
            "name": name,
            "passed": bool(passed),
            "severity": severity,
            "detail": detail,
        }
    )


def build_gate_report(
    *,
    auto_results: Optional[Dict[str, Any]],
    shadow_results: Dict[str, Any],
    tuned_results: Dict[str, Any],
    auto_run_dir: Optional[Path] = None,
    thresholds: GateThresholds = GateThresholds(),
) -> Dict[str, Any]:
    shadow_model = first_model_result(shadow_results)
    tuned_model = first_model_result(tuned_results)
    shadow_metrics = first_model_metrics(shadow_results)
    tuned_metrics = first_model_metrics(tuned_results)

    if auto_results is None:
        return {
            "status": "missing_auto_results",
            "passed": False,
            "checks": [
                {
                    "name": "auto_results_present",
                    "passed": False,
                    "severity": "fail",
                    "detail": "Fresh auto evaluation results are required before merge.",
                }
            ],
            "baselines": {
                "shadow_model": first_model_name(shadow_results),
                "tuned_model": first_model_name(tuned_results),
                "shadow_metrics": {
                    key: shadow_metrics.get(key) for key in KEY_METRICS
                },
                "tuned_metrics": {key: tuned_metrics.get(key) for key in KEY_METRICS},
            },
        }

    auto_model = first_model_result(auto_results)
    auto_metrics = first_model_metrics(auto_results)
    common_docs = sorted(
        set(doc_ids(auto_model)) & set(doc_ids(shadow_model)) & set(doc_ids(tuned_model))
    )

    checks: List[Dict[str, Any]] = []
    add_check(
        checks,
        "common_document_count",
        len(common_docs) >= thresholds.min_common_documents,
        f"{len(common_docs)} common documents; required {thresholds.min_common_documents}.",
    )

    rows = {
        key: metric_row(key, auto_metrics, shadow_metrics, tuned_metrics)
        for key in KEY_METRICS
    }

    add_check(
        checks,
        "schema_not_worse_than_shadow",
        rows["schema_compliance"]["delta_vs_shadow"] >= -thresholds.schema_tolerance,
        (
            f"auto={rows['schema_compliance']['auto']:.4f}, "
            f"shadow={rows['schema_compliance']['shadow']:.4f}, "
            f"tolerance={thresholds.schema_tolerance:.4f}"
        ),
    )
    add_check(
        checks,
        "row_alignment_not_worse_than_shadow",
        rows["row_alignment_f1"]["delta_vs_shadow"]
        >= -thresholds.row_alignment_tolerance,
        (
            f"auto={rows['row_alignment_f1']['auto']:.4f}, "
            f"shadow={rows['row_alignment_f1']['shadow']:.4f}, "
            f"tolerance={thresholds.row_alignment_tolerance:.4f}"
        ),
    )
    add_check(
        checks,
        "sheet_placement_not_worse_than_shadow",
        rows["sheet_placement_accuracy"]["delta_vs_shadow"]
        >= -thresholds.sheet_placement_tolerance,
        (
            f"auto={rows['sheet_placement_accuracy']['auto']:.4f}, "
            f"shadow={rows['sheet_placement_accuracy']['shadow']:.4f}, "
            f"tolerance={thresholds.sheet_placement_tolerance:.4f}"
        ),
    )
    precision = rows["precision_excl_discoveries"]["auto"]
    add_check(
        checks,
        "precision_remains_safe",
        (
            precision >= thresholds.min_precision_excl_discoveries
            and rows["precision_excl_discoveries"]["delta_vs_shadow"]
            >= -thresholds.precision_tolerance
        ),
        (
            f"auto={precision:.4f}, shadow={rows['precision_excl_discoveries']['shadow']:.4f}, "
            f"minimum={thresholds.min_precision_excl_discoveries:.4f}, "
            f"delta_tolerance={thresholds.precision_tolerance:.4f}"
        ),
    )
    add_check(
        checks,
        "completeness_not_worse_than_shadow",
        rows["completeness"]["delta_vs_shadow"] >= -thresholds.completeness_tolerance,
        (
            f"auto={rows['completeness']['auto']:.4f}, "
            f"shadow={rows['completeness']['shadow']:.4f}, "
            f"tolerance={thresholds.completeness_tolerance:.4f}"
        ),
    )
    add_check(
        checks,
        "value_match_near_tuned",
        rows["value_match_rate"]["delta_vs_tuned"]
        >= -thresholds.value_match_tolerance_vs_tuned,
        (
            f"auto={rows['value_match_rate']['auto']:.4f}, "
            f"tuned={rows['value_match_rate']['tuned']:.4f}, "
            f"tolerance={thresholds.value_match_tolerance_vs_tuned:.4f}"
        ),
        severity="warn",
    )

    doc_regressions: List[Dict[str, Any]] = []
    for doc_id in common_docs:
        auto_doc = doc_metrics(auto_model, doc_id)
        shadow_doc = doc_metrics(shadow_model, doc_id)
        failures: List[str] = []
        if (
            auto_doc["schema_compliance"] - shadow_doc["schema_compliance"]
            < -thresholds.schema_tolerance
        ):
            failures.append("schema")
        if (
            auto_doc["row_alignment_f1"] - shadow_doc["row_alignment_f1"]
            < -thresholds.row_alignment_tolerance
        ):
            failures.append("row_alignment")
        if (
            auto_doc["completeness"] - shadow_doc["completeness"]
            < -thresholds.completeness_tolerance
        ):
            failures.append("completeness")
        if failures:
            doc_regressions.append(
                {
                    "doc_id": doc_id,
                    "failures": failures,
                    "auto": auto_doc,
                    "shadow": shadow_doc,
                }
            )
    add_check(
        checks,
        "no_document_level_regressions",
        not doc_regressions,
        f"{len(doc_regressions)} documents regressed beyond tolerances.",
    )

    artifact_report = validate_auto_artifacts(
        auto_run_dir=auto_run_dir,
        doc_ids_to_check=common_docs,
    )
    add_check(
        checks,
        "auto_artifacts_complete",
        bool(artifact_report.get("passed")),
        artifact_report.get("reason")
        or (
            f"{artifact_report.get('summary', {}).get('failed_document_count', 0)} "
            "documents have missing or invalid artifacts."
        ),
    )

    failed = [item for item in checks if item["severity"] == "fail" and not item["passed"]]
    warned = [item for item in checks if item["severity"] == "warn" and not item["passed"]]
    return {
        "status": "pass" if not failed else "fail",
        "passed": not failed,
        "checks": checks,
        "warnings": warned,
        "models": {
            "auto": first_model_name(auto_results),
            "shadow": first_model_name(shadow_results),
            "tuned": first_model_name(tuned_results),
        },
        "common_documents": common_docs,
        "metrics": rows,
        "document_regressions": doc_regressions,
        "artifact_validation": artifact_report,
    }


def render_markdown(report: Dict[str, Any]) -> str:
    lines = [
        "# Auto Merge Gate Report",
        "",
        f"Status: **{report.get('status')}**",
        "",
    ]
    if report.get("models"):
        lines.extend(
            [
                "## Models",
                "",
                *[
                    f"- {name}: `{model}`"
                    for name, model in sorted(report["models"].items())
                ],
                "",
            ]
        )
    lines.extend(["## Checks", ""])
    for check in report.get("checks", []):
        icon = "PASS" if check.get("passed") else "FAIL"
        if check.get("severity") == "warn" and not check.get("passed"):
            icon = "WARN"
        lines.append(f"- {icon} `{check.get('name')}`: {check.get('detail')}")
    if report.get("metrics"):
        lines.extend(["", "## Metrics", ""])
        lines.append("| Metric | Auto | Shadow | Tuned | Delta vs Shadow | Delta vs Tuned |")
        lines.append("|---|---:|---:|---:|---:|---:|")
        for key in KEY_METRICS:
            row = report["metrics"].get(key, {})
            lines.append(
                f"| {key} | {row.get('auto', 0):.4f} | "
                f"{row.get('shadow', 0):.4f} | {row.get('tuned', 0):.4f} | "
                f"{row.get('delta_vs_shadow', 0):+.4f} | "
                f"{row.get('delta_vs_tuned', 0):+.4f} |"
            )
    regressions = report.get("document_regressions") or []
    if regressions:
        lines.extend(["", "## Document Regressions", ""])
        for item in regressions:
            failures = ", ".join(item.get("failures") or [])
            lines.append(f"- `{item.get('doc_id')}`: {failures}")
    artifact = report.get("artifact_validation") or {}
    if artifact:
        lines.extend(["", "## Artifact Validation", ""])
        summary = artifact.get("summary") or {}
        if artifact.get("checked"):
            lines.append(f"- model root: `{artifact.get('model_root')}`")
            lines.append(f"- documents checked: {summary.get('document_count', 0)}")
            lines.append(f"- failed documents: {summary.get('failed_document_count', 0)}")
            lines.append(f"- accepted patches: {summary.get('accepted_patch_count', 0)}")
            lines.append(
                "- metadata mutated documents: "
                f"{summary.get('metadata_mutated_document_count', 0)}"
            )
            lines.append(
                f"- trace-only documents: {summary.get('trace_only_document_count', 0)}"
            )
            lines.append(
                "- non-production trace documents: "
                f"{summary.get('non_production_trace_document_count', 0)}"
            )
            trace_modes = summary.get("trace_mode_counts") or {}
            if trace_modes:
                rendered_modes = ", ".join(
                    f"{mode}={count}" for mode, count in sorted(trace_modes.items())
                )
                lines.append(f"- trace modes: {rendered_modes}")
            failed_docs = [
                item for item in artifact.get("documents", [])
                if not item.get("passed")
            ]
            if failed_docs:
                lines.extend(["", "### Failed Artifact Documents", ""])
                for item in failed_docs[:20]:
                    errors = item.get("errors") or []
                    rendered_errors = ", ".join(str(error) for error in errors[:8])
                    if len(errors) > 8:
                        rendered_errors += f", ... ({len(errors)} total)"
                    lines.append(f"- `{item.get('doc_id')}`: {rendered_errors}")
        else:
            lines.append(f"- not checked: {artifact.get('reason')}")
    return "\n".join(lines) + "\n"


def parse_args(argv: Optional[List[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--auto-results", type=Path, default=DEFAULT_AUTO)
    parser.add_argument(
        "--auto-run-dir",
        type=Path,
        help="Run directory containing the auto model/doc/run_*/ artifacts.",
    )
    parser.add_argument("--shadow-results", type=Path, default=DEFAULT_SHADOW)
    parser.add_argument("--tuned-results", type=Path, default=DEFAULT_TUNED)
    parser.add_argument(
        "--output-json",
        type=Path,
        default=DEFAULT_ROOT / "artifacts" / "auto_merge_gate_report.json",
    )
    parser.add_argument(
        "--output-md",
        type=Path,
        default=DEFAULT_ROOT / "artifacts" / "auto_merge_gate_report.md",
    )
    parser.add_argument("--min-common-documents", type=int, default=6)
    parser.add_argument("--schema-tolerance", type=float, default=0.01)
    parser.add_argument("--row-alignment-tolerance", type=float, default=0.02)
    parser.add_argument("--sheet-placement-tolerance", type=float, default=0.01)
    parser.add_argument("--precision-tolerance", type=float, default=0.02)
    parser.add_argument("--min-precision-excl-discoveries", type=float, default=0.97)
    parser.add_argument("--completeness-tolerance", type=float, default=0.01)
    parser.add_argument("--value-match-tolerance-vs-tuned", type=float, default=0.03)
    parser.add_argument(
        "--no-fail",
        action="store_true",
        help="Write reports but exit 0 even when the gate fails.",
    )
    return parser.parse_args(argv)


def main(argv: Optional[List[str]] = None) -> int:
    args = parse_args(argv)
    thresholds = GateThresholds(
        min_common_documents=args.min_common_documents,
        schema_tolerance=args.schema_tolerance,
        row_alignment_tolerance=args.row_alignment_tolerance,
        sheet_placement_tolerance=args.sheet_placement_tolerance,
        precision_tolerance=args.precision_tolerance,
        min_precision_excl_discoveries=args.min_precision_excl_discoveries,
        completeness_tolerance=args.completeness_tolerance,
        value_match_tolerance_vs_tuned=args.value_match_tolerance_vs_tuned,
    )

    shadow_results = load_json(args.shadow_results)
    tuned_results = load_json(args.tuned_results)
    auto_results = load_json(args.auto_results) if args.auto_results.exists() else None
    auto_run_dir = args.auto_run_dir
    if auto_run_dir is None and auto_results is not None:
        auto_run_dir = resolve_run_dir(args.auto_results)
    report = build_gate_report(
        auto_results=auto_results,
        shadow_results=shadow_results,
        tuned_results=tuned_results,
        auto_run_dir=auto_run_dir,
        thresholds=thresholds,
    )

    args.output_json.parent.mkdir(parents=True, exist_ok=True)
    args.output_json.write_text(json.dumps(report, indent=2), encoding="utf-8")
    args.output_md.write_text(render_markdown(report), encoding="utf-8")
    print(f"wrote {args.output_json}")
    print(f"wrote {args.output_md}")
    print(f"status={report['status']}")

    if args.no_fail or report.get("passed"):
        return 0
    return 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
