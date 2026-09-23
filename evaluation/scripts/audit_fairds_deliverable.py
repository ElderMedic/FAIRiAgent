#!/usr/bin/env python3
"""Audit one FAIRiAgent run at the generated workbook boundary.

The audit deliberately checks both layers of the deliverable:

* metadata/schema: packages, FAIR-DS column labels, mandatory-field exposure;
* values/matrix: entity cardinality, parent links, meaningful row payloads, and
  exact JSON-to-XLSX materialization.

It is source-agnostic. No document name, organism, assay, package, or identifier
pattern is encoded here; package applicability comes from the run's persisted
source-to-schema contract audit.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

from openpyxl import load_workbook


ISA_LEVELS: Tuple[str, ...] = (
    "investigation",
    "study",
    "observationunit",
    "sample",
    "assay",
)
PARENT_LINKS = {
    "study": ("investigation identifier", "investigation"),
    "observationunit": ("study identifier", "study"),
    "sample": ("observation unit identifier", "observationunit"),
    "assay": ("sample identifier", "sample"),
}
IDENTIFIER_FIELDS = {
    "investigation": "investigation identifier",
    "study": "study identifier",
    "observationunit": "observation unit identifier",
    "sample": "sample identifier",
    "assay": "assay identifier",
}
NAME_FIELDS = {
    "investigation": "investigation title",
    "study": "study title",
    "observationunit": "observation unit name",
    "sample": "sample name",
    "assay": "assay name",
}
REQUIREMENT_SUFFIX_RE = re.compile(r"\s*\((?:M|O|R)\)\s*$", re.IGNORECASE)
LIST_DELIMITER_RE = re.compile(r"[,;|\n]")


def _load_json(path: Path) -> Dict[str, Any]:
    parsed = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(parsed, dict):
        raise ValueError(f"Expected a JSON object in {path}")
    return parsed


def _sheet_level(title: str) -> Optional[str]:
    base = title.split(" - ", 1)[0].strip().lower().replace(" ", "")
    if base == "help":
        return "help"
    return base if base in ISA_LEVELS else None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _trim_trailing_empty(values: Sequence[Any]) -> List[str]:
    normalized = [_cell_text(value).strip() for value in values]
    while normalized and not normalized[-1]:
        normalized.pop()
    return normalized


def _issue(
    issues: List[Dict[str, Any]],
    severity: str,
    code: str,
    message: str,
    **details: Any,
) -> None:
    record: Dict[str, Any] = {
        "severity": severity,
        "code": code,
        "message": message,
    }
    if details:
        record["details"] = details
    issues.append(record)


def _matrix_sheet(matrix: Mapping[str, Any], level: str) -> Dict[str, Any]:
    raw = matrix.get(level) or {}
    return raw if isinstance(raw, dict) else {}


def _row_signature(row: Mapping[str, Any], keys: Iterable[str]) -> Tuple[str, ...]:
    return tuple(_cell_text(row.get(key)).strip() for key in keys)


def audit_run(run_dir: Path) -> Dict[str, Any]:
    run_dir = run_dir.resolve()
    deliverables = run_dir / "deliverables"
    if not deliverables.is_dir():
        deliverables = run_dir
    reports = run_dir / "reports"
    metadata_path = deliverables / "metadata.json"
    matrix_path = deliverables / "isa_values.json"
    workbook_path = deliverables / "metadata_fairds.xlsx"
    entity_plan_path = reports / "entity_plan.json"
    workflow_report_path = run_dir / "workflow_report.json"

    issues: List[Dict[str, Any]] = []
    required_artifacts = (metadata_path, matrix_path, workbook_path)
    for path in required_artifacts:
        if not path.is_file():
            _issue(
                issues,
                "error",
                "missing_artifact",
                f"Required deliverable is missing: {path.name}",
                path=str(path),
            )
    if (run_dir / ".running").exists():
        _issue(
            issues,
            "error",
            "run_not_terminal",
            "Run still has a .running marker.",
        )
    if any(issue["code"] == "missing_artifact" for issue in issues):
        return _finalize_report(run_dir, issues, {}, {})

    try:
        metadata = _load_json(metadata_path)
        matrix = _load_json(matrix_path)
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        _issue(issues, "error", "invalid_json_artifact", str(exc))
        return _finalize_report(run_dir, issues, {}, {})

    packages_used = [
        str(name).strip()
        for name in metadata.get("packages_used", []) or []
        if str(name).strip()
    ]
    if not packages_used:
        _issue(
            issues,
            "error",
            "packages_missing",
            "metadata.json does not declare any FAIR-DS package.",
        )

    _audit_persisted_field_contracts(issues, metadata, matrix, packages_used)

    package_trace = metadata.get("package_selection_trace")
    if isinstance(package_trace, dict) and package_trace:
        if not package_trace.get("stable", False):
            _issue(
                issues,
                "error",
                "package_selection_unstable",
                "Package selection did not pass the bounded contract audit.",
                uncovered_levels=package_trace.get("uncovered_schema_levels", []),
            )
        if package_trace.get("coverage_resolved") is False:
            _issue(
                issues,
                "error",
                "package_schema_coverage_unresolved",
                "Strong source-to-schema matches remain uncovered.",
                uncovered_levels=package_trace.get("uncovered_schema_levels", []),
            )
    else:
        _issue(
            issues,
            "error",
            "package_selection_trace_missing",
            "Package choice cannot be independently audited from metadata.json.",
        )

    metadata_matrix = metadata.get("isa_values") or {}
    if metadata_matrix != matrix:
        _issue(
            issues,
            "error",
            "matrix_artifacts_diverge",
            "metadata.json isa_values and isa_values.json are not identical.",
        )

    isa_structure = metadata.get("isa_structure") or {}
    for level in ISA_LEVELS:
        columns = list(_matrix_sheet(matrix, level).get("columns") or [])
        structure_columns = list(
            (isa_structure.get(level) or {}).get("columns") or []
        )
        if structure_columns != columns:
            _issue(
                issues,
                "error",
                "structure_columns_diverge",
                f"isa_structure and isa_values columns differ for {level}.",
                level=level,
            )

    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    except Exception as exc:  # openpyxl raises several format-specific classes
        _issue(issues, "error", "workbook_unreadable", str(exc))
        return _finalize_report(run_dir, issues, metadata, matrix)

    formula_cells: List[str] = []
    error_cells: List[str] = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formula_cells.append(f"{worksheet.title}!{cell.coordinate}")
                elif cell.data_type == "e":
                    error_cells.append(f"{worksheet.title}!{cell.coordinate}")
    if formula_cells:
        _issue(
            issues,
            "error",
            "formula_cells_present",
            "The metadata deliverable must materialize values, not formulas.",
            cells=formula_cells[:20],
        )
    if error_cells:
        _issue(
            issues,
            "error",
            "excel_error_cells_present",
            "The workbook contains Excel error-valued cells.",
            cells=error_cells[:20],
        )

    sheet_map: Dict[str, List[Any]] = defaultdict(list)
    unknown_sheets: List[str] = []
    for worksheet in workbook.worksheets:
        level = _sheet_level(worksheet.title)
        if level is None:
            unknown_sheets.append(worksheet.title)
        else:
            sheet_map[level].append(worksheet)
    if "person" in {title.strip().lower() for title in workbook.sheetnames}:
        _issue(
            issues,
            "error",
            "person_sheet_forbidden",
            "Investigators must be rows in Investigation, not a Person sheet.",
        )
    for title in unknown_sheets:
        _issue(
            issues,
            "error",
            "unexpected_sheet",
            f"Unexpected workbook sheet: {title}",
        )
    for level in ISA_LEVELS:
        worksheets = sheet_map.get(level, [])
        if len(worksheets) != 1:
            _issue(
                issues,
                "error",
                "sheet_cardinality",
                f"Expected exactly one {level} sheet, found {len(worksheets)}.",
                level=level,
            )
            continue
        _audit_worksheet(
            issues,
            level,
            worksheets[0],
            _matrix_sheet(matrix, level),
            isa_structure.get(level) or {},
            bool(metadata.get("needs_review", False)),
        )
    if not sheet_map.get("help"):
        _issue(
            issues,
            "warning",
            "help_sheet_missing",
            "Workbook has no Help sheet describing FAIR-DS terms.",
        )

    _audit_entity_links(issues, matrix)
    _audit_controlled_values(issues, metadata, matrix)
    _audit_entity_plan(issues, matrix, entity_plan_path)
    _audit_workflow_status(issues, workflow_report_path)

    metrics = _build_metrics(metadata, matrix, workbook.sheetnames)
    return _finalize_report(run_dir, issues, metrics, {})


def _audit_persisted_field_contracts(
    issues: List[Dict[str, Any]],
    metadata: Mapping[str, Any],
    matrix: Mapping[str, Any],
    packages_used: Sequence[str],
) -> None:
    """Verify workbook columns against retrieved FAIR-DS definitions.

    Comparing the matrix with ``isa_structure`` alone is circular because both
    artifacts are produced by FAIRiAgent.  The persisted ``_field_definitions``
    records are the run's independently retrieved FAIR-DS contract boundary.
    Every emitted column must match one of those exact level/term pairs and the
    defining package must be among ``packages_used``.
    """
    definitions = [
        item
        for item in metadata.get("_field_definitions") or []
        if isinstance(item, Mapping)
    ]
    if not definitions:
        _issue(
            issues,
            "error",
            "field_contract_definitions_missing",
            "No persisted FAIR-DS field definitions are available to validate columns.",
        )
        return

    selected = {str(package).strip().casefold() for package in packages_used}
    by_key: Dict[Tuple[str, str], List[Mapping[str, Any]]] = defaultdict(list)
    for definition in definitions:
        level = str(definition.get("isa_sheet") or "").strip().lower()
        term = str(
            definition.get("term")
            or definition.get("field_name")
            or definition.get("name")
            or ""
        ).strip()
        if level in ISA_LEVELS and term:
            by_key[(level, term)].append(definition)

    missing: List[Dict[str, str]] = []
    package_mismatches: List[Dict[str, str]] = []
    for level in ISA_LEVELS:
        for raw_column in _matrix_sheet(matrix, level).get("columns") or []:
            column = str(raw_column).strip()
            matches = by_key.get((level, column), [])
            if not matches:
                missing.append({"level": level, "column": column})
                continue
            if any(bool(item.get("source_extension")) for item in matches):
                continue
            packages = {
                str(item.get("package") or item.get("package_name") or "").strip()
                for item in matches
                if str(item.get("package") or item.get("package_name") or "").strip()
            }
            if packages and not any(package.casefold() in selected for package in packages):
                package_mismatches.append(
                    {
                        "level": level,
                        "column": column,
                        "defining_packages": " | ".join(sorted(packages)),
                    }
                )
    if missing:
        _issue(
            issues,
            "error",
            "columns_missing_fairds_contract",
            "Workbook columns are absent from persisted FAIR-DS field definitions.",
            count=len(missing),
            examples=missing[:20],
        )
    if package_mismatches:
        _issue(
            issues,
            "error",
            "column_package_not_selected",
            "Workbook columns come from FAIR-DS packages not declared by the run.",
            count=len(package_mismatches),
            examples=package_mismatches[:20],
        )


def _audit_worksheet(
    issues: List[Dict[str, Any]],
    level: str,
    worksheet: Any,
    matrix_sheet: Mapping[str, Any],
    structure_sheet: Mapping[str, Any],
    needs_review: bool,
) -> None:
    matrix_columns = [str(value) for value in matrix_sheet.get("columns") or []]
    matrix_rows = [
        row for row in matrix_sheet.get("rows") or [] if isinstance(row, dict)
    ]
    header_values = next(
        worksheet.iter_rows(min_row=1, max_row=1, values_only=True),
        (),
    )
    headers = _trim_trailing_empty(header_values)
    if headers != matrix_columns:
        _issue(
            issues,
            "error",
            "excel_headers_diverge",
            f"Excel headers do not match isa_values columns for {level}.",
            level=level,
            excel=headers,
            isa_values=matrix_columns,
        )
    duplicate_headers = [
        name for name, count in Counter(headers).items() if name and count > 1
    ]
    if duplicate_headers:
        _issue(
            issues,
            "error",
            "duplicate_headers",
            f"Duplicate column names in {level}.",
            columns=duplicate_headers,
        )
    marked_headers = [name for name in headers if REQUIREMENT_SUFFIX_RE.search(name)]
    if marked_headers:
        _issue(
            issues,
            "error",
            "nonstandard_requirement_suffix",
            f"Non-standard (M)/(O)/(R) suffixes found in {level} headers.",
            columns=marked_headers,
        )

    declared_fields = {
        str(field.get("field_name") or "").strip()
        for field in structure_sheet.get("fields") or []
        if isinstance(field, dict) and str(field.get("field_name") or "").strip()
    }
    allowed_linkage = {PARENT_LINKS[level][0]} if level in PARENT_LINKS else set()
    undeclared = [
        column
        for column in matrix_columns
        if column not in declared_fields and column not in allowed_linkage
    ]
    if undeclared:
        _issue(
            issues,
            "error",
            "undeclared_metadata_columns",
            f"Columns in {level} are not declared FAIR-DS fields or ISA linkage.",
            columns=undeclared,
        )

    excel_rows: List[Dict[str, str]] = []
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        row_values = list(values[: len(headers)])
        if not any(_cell_text(value).strip() for value in row_values):
            continue
        excel_rows.append(
            {
                header: _cell_text(value).strip()
                for header, value in zip(headers, row_values)
            }
        )
    normalized_matrix_rows = [
        {column: _cell_text(row.get(column)).strip() for column in matrix_columns}
        for row in matrix_rows
    ]
    if excel_rows != normalized_matrix_rows:
        _issue(
            issues,
            "error",
            "excel_values_diverge",
            f"Excel values do not exactly materialize isa_values rows for {level}.",
            level=level,
            excel_row_count=len(excel_rows),
            isa_row_count=len(normalized_matrix_rows),
        )

    required_fields = {
        str(field.get("field_name") or "").strip()
        for field in structure_sheet.get("fields") or []
        if isinstance(field, dict)
        and field.get("required") is True
        and str(field.get("field_name") or "").strip()
    }
    missing_required_columns = sorted(required_fields - set(matrix_columns))
    if missing_required_columns:
        _issue(
            issues,
            "error",
            "mandatory_columns_missing",
            f"Mandatory FAIR-DS fields are absent from {level} columns.",
            columns=missing_required_columns,
        )
    for column in sorted(required_fields & set(matrix_columns)):
        blank_count = sum(not _cell_text(row.get(column)).strip() for row in matrix_rows)
        if blank_count:
            severity = "warning" if needs_review else "error"
            _issue(
                issues,
                severity,
                "mandatory_values_blank",
                f"Mandatory field {level}.{column} is blank in {blank_count} row(s).",
                level=level,
                column=column,
                blank_rows=blank_count,
            )


def _audit_entity_links(
    issues: List[Dict[str, Any]], matrix: Mapping[str, Any]
) -> None:
    identifiers: Dict[str, set[str]] = {}
    for level in ISA_LEVELS:
        rows = _matrix_sheet(matrix, level).get("rows") or []
        identifier_field = IDENTIFIER_FIELDS[level]
        values = [
            _cell_text(row.get(identifier_field)).strip()
            for row in rows
            if isinstance(row, dict)
        ]
        identifiers[level] = {value for value in values if value}
        blanks = sum(not value for value in values)
        if blanks:
            _issue(
                issues,
                "error",
                "blank_entity_identifier",
                f"{level} contains {blanks} blank identifier(s).",
            )
        list_like = [value for value in values if LIST_DELIMITER_RE.search(value)]
        if list_like:
            _issue(
                issues,
                "error",
                "list_like_identifier",
                f"{level} identifier cells contain list delimiters.",
                examples=list_like[:5],
            )
        if level != "investigation" and len(values) != len(set(values)):
            _issue(
                issues,
                "error",
                "duplicate_entity_identifier",
                f"{level} contains duplicate entity identifiers.",
            )

        if level not in {"investigation", "study"} and len(rows) > 1:
            parent_field = PARENT_LINKS[level][0]
            payload_columns = [
                column
                for column in _matrix_sheet(matrix, level).get("columns") or []
                if column not in {identifier_field, parent_field}
            ]
            if not payload_columns:
                _issue(
                    issues,
                    "error",
                    "entity_rows_without_payload_columns",
                    f"{level} has multiple rows but no descriptive/measurement columns.",
                )
            else:
                signatures = [
                    _row_signature(row, payload_columns)
                    for row in rows
                    if isinstance(row, dict)
                ]
                if signatures and len(set(signatures)) == 1:
                    _issue(
                        issues,
                        "error",
                        "meaningless_repeated_entity_rows",
                        f"All {level} rows are identical after identifiers and links are removed.",
                        row_count=len(signatures),
                    )
                empty_payload_rows = sum(not any(signature) for signature in signatures)
                if empty_payload_rows:
                    _issue(
                        issues,
                        "error",
                        "empty_entity_payload",
                        f"{level} has {empty_payload_rows} row(s) with no meaningful payload.",
                    )

    for child_level, (link_field, parent_level) in PARENT_LINKS.items():
        rows = _matrix_sheet(matrix, child_level).get("rows") or []
        missing_links: List[int] = []
        orphan_links: List[str] = []
        for index, row in enumerate(rows, start=1):
            if not isinstance(row, dict):
                continue
            value = _cell_text(row.get(link_field)).strip()
            if not value:
                missing_links.append(index)
            elif value not in identifiers[parent_level]:
                orphan_links.append(value)
        if missing_links:
            _issue(
                issues,
                "error",
                "missing_parent_link",
                f"{child_level} rows are missing {link_field}.",
                rows=missing_links[:20],
            )
        if orphan_links:
            _issue(
                issues,
                "error",
                "orphan_parent_link",
                f"{child_level} rows reference unknown {parent_level} identifiers.",
                values=list(dict.fromkeys(orphan_links))[:20],
            )


def _audit_controlled_values(
    issues: List[Dict[str, Any]],
    metadata: Mapping[str, Any],
    matrix: Mapping[str, Any],
) -> None:
    """Check all matrix cells against persisted FAIR-DS regex contracts."""
    contracts: Dict[Tuple[str, str], str] = {}
    for field in metadata.get("_field_definitions") or []:
        if not isinstance(field, dict):
            continue
        level = str(field.get("isa_sheet") or "").strip().lower()
        name = str(
            field.get("term")
            or field.get("field_name")
            or field.get("name")
            or ""
        ).strip().lower()
        pattern = str(field.get("regex") or "").strip()
        if level in ISA_LEVELS and name and pattern:
            contracts[(level, name)] = pattern

    invalid: List[Dict[str, Any]] = []
    invalid_patterns: List[Dict[str, str]] = []
    for (level, name), pattern in contracts.items():
        try:
            compiled = re.compile(pattern)
        except re.error as exc:
            invalid_patterns.append(
                {"level": level, "field": name, "regex": pattern, "error": str(exc)}
            )
            continue
        for row_index, row in enumerate(
            _matrix_sheet(matrix, level).get("rows") or [], start=1
        ):
            if not isinstance(row, dict):
                continue
            value = _cell_text(row.get(name)).strip()
            if value and compiled.fullmatch(value) is None:
                invalid.append(
                    {
                        "level": level,
                        "field": name,
                        "row": row_index,
                        "value": value,
                    }
                )
    if invalid_patterns:
        _issue(
            issues,
            "error",
            "invalid_fairds_contract_regex",
            "Persisted FAIR-DS field contracts contain invalid regular expressions.",
            examples=invalid_patterns[:10],
        )
    if invalid:
        _issue(
            issues,
            "error",
            "controlled_value_contract_violation",
            "Workbook cells violate persisted FAIR-DS controlled-value contracts.",
            count=len(invalid),
            examples=invalid[:20],
        )


def _audit_entity_plan(
    issues: List[Dict[str, Any]],
    matrix: Mapping[str, Any],
    entity_plan_path: Path,
) -> None:
    if not entity_plan_path.is_file():
        _issue(
            issues,
            "warning",
            "entity_plan_missing",
            "Entity plan is unavailable for cardinality/contact reconciliation.",
        )
        return
    try:
        entity_plan = _load_json(entity_plan_path)
    except Exception as exc:
        _issue(issues, "error", "entity_plan_invalid", str(exc))
        return
    validation = entity_plan.get("validation") or {}
    plan = entity_plan.get("plan") or {}
    if validation.get("passed") is False:
        _issue(
            issues,
            "error",
            "entity_plan_validation_failed",
            "Entity-plan validation failed.",
            errors=validation.get("errors", []),
        )
    expected_counts = validation.get("row_counts") or {}
    for level in ISA_LEVELS:
        actual = len(_matrix_sheet(matrix, level).get("rows") or [])
        expected = expected_counts.get(level)
        if level == "investigation":
            expected = validation.get("investigation_contact_count", expected)
        if expected is not None and int(expected) != actual:
            _issue(
                issues,
                "error",
                "entity_plan_row_count_mismatch",
                f"{level} row count does not match the validated entity plan.",
                expected=int(expected),
                actual=actual,
            )
    expected_contacts = validation.get("investigation_contact_count")
    if expected_contacts is not None:
        investigation_rows = _matrix_sheet(matrix, "investigation").get("rows") or []
        people = {
            (
                _cell_text(row.get("firstname")).strip(),
                _cell_text(row.get("lastname")).strip(),
            )
            for row in investigation_rows
            if isinstance(row, dict)
        }
        people.discard(("", ""))
        if len(people) != int(expected_contacts):
            _issue(
                issues,
                "error",
                "investigator_count_mismatch",
                "Investigation rows do not preserve every planned investigator.",
                expected=int(expected_contacts),
                actual=len(people),
            )

    # When a selected package exposes a level-specific name column, the name
    # must identify the corresponding planned entity.  This catches a common
    # projection failure where one LLM example is broadcast to every row even
    # though identifiers and descriptions differ.
    level_plans = {
        str(item.get("level") or "").strip().lower(): item
        for item in plan.get("levels") or []
        if isinstance(item, dict)
    }
    for level, name_field in NAME_FIELDS.items():
        block = _matrix_sheet(matrix, level)
        columns = {
            str(column).strip().lower() for column in block.get("columns") or []
        }
        if name_field not in columns:
            continue
        rows = [row for row in block.get("rows") or [] if isinstance(row, dict)]
        rows_by_id: Dict[str, List[Mapping[str, Any]]] = defaultdict(list)
        id_field = IDENTIFIER_FIELDS[level]
        for row in rows:
            rows_by_id[_cell_text(row.get(id_field)).strip()].append(row)
        mismatches: List[Dict[str, str]] = []
        for entity in (level_plans.get(level) or {}).get("entities") or []:
            if not isinstance(entity, dict):
                continue
            row_id = _cell_text(entity.get("row_id")).strip()
            expected_id = (
                _cell_text(entity.get("external_identifier")).strip() or row_id
            )
            expected_name = _cell_text(entity.get("label")).strip() or row_id
            candidates = rows_by_id.get(expected_id) or []
            if not candidates:
                continue
            actual_names = {
                _cell_text(row.get(name_field)).strip() for row in candidates
            }
            mapped_names = {
                _cell_text(attribute.get("value")).strip()
                for attribute in entity.get("attributes") or []
                if isinstance(attribute, dict)
                and _cell_text(attribute.get("field_name")).strip().lower()
                == name_field
                and _cell_text(attribute.get("value")).strip()
            }
            allowed_names = {expected_name, *mapped_names}
            if not actual_names or not actual_names.issubset(allowed_names):
                mismatches.append(
                    {
                        "identifier": expected_id,
                        "expected": expected_name,
                        "actual": " | ".join(sorted(actual_names)),
                    }
                )
        if mismatches:
            _issue(
                issues,
                "error",
                "entity_name_plan_mismatch",
                f"{level} names do not identify their planned entities.",
                level=level,
                field=name_field,
                examples=mismatches[:10],
            )


def _audit_workflow_status(
    issues: List[Dict[str, Any]], workflow_report_path: Path
) -> None:
    if not workflow_report_path.is_file():
        _issue(
            issues,
            "warning",
            "workflow_report_missing",
            "Workflow report is unavailable.",
        )
        return
    try:
        workflow = _load_json(workflow_report_path)
    except Exception as exc:
        _issue(issues, "error", "workflow_report_invalid", str(exc))
        return
    if str(workflow.get("workflow_status") or "").lower() != "completed":
        _issue(
            issues,
            "error",
            "workflow_not_completed",
            "Workflow report does not declare completed status.",
            status=workflow.get("workflow_status"),
        )


def _build_metrics(
    metadata: Mapping[str, Any],
    matrix: Mapping[str, Any],
    sheet_names: Sequence[str],
) -> Dict[str, Any]:
    row_counts = {
        level: len(_matrix_sheet(matrix, level).get("rows") or [])
        for level in ISA_LEVELS
    }
    column_counts = {
        level: len(_matrix_sheet(matrix, level).get("columns") or [])
        for level in ISA_LEVELS
    }
    return {
        "packages_used": metadata.get("packages_used", []),
        "needs_review": bool(metadata.get("needs_review", False)),
        "overall_confidence": metadata.get("overall_confidence"),
        "row_counts": row_counts,
        "column_counts": column_counts,
        "sheet_names": list(sheet_names),
    }


def _finalize_report(
    run_dir: Path,
    issues: List[Dict[str, Any]],
    metrics: Mapping[str, Any],
    _unused: Mapping[str, Any],
) -> Dict[str, Any]:
    error_count = sum(issue["severity"] == "error" for issue in issues)
    warning_count = sum(issue["severity"] == "warning" for issue in issues)
    return {
        "schema_version": "fairds_deliverable_audit.v1",
        "run_dir": str(run_dir),
        # ``passed`` means ready without unresolved review findings.  Keep the
        # structural result separate so callers cannot mistake a warning-heavy
        # workbook for an accepted deliverable.
        "passed": error_count == 0 and warning_count == 0,
        "structural_passed": error_count == 0,
        "review_required": warning_count > 0,
        "error_count": error_count,
        "warning_count": warning_count,
        "metrics": dict(metrics),
        "issues": issues,
    }


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("run_dir", type=Path)
    parser.add_argument(
        "--output",
        type=Path,
        help="Optional JSON report path. The audit always prints JSON to stdout.",
    )
    args = parser.parse_args(argv)
    report = audit_run(args.run_dir)
    rendered = json.dumps(report, indent=2, ensure_ascii=False)
    print(rendered)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(rendered + "\n", encoding="utf-8")
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    sys.exit(main())
