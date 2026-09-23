#!/usr/bin/env python3
"""Compare a FAIR-DS deliverable with an external expected-outcome contract.

The evaluator is intentionally generic.  Input-specific expectations live in
versioned JSON fixtures, never in FAIRiAgent's production extraction code.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import sys
from collections import Counter
from pathlib import Path
from typing import Any, Dict, List, Mapping, Optional, Sequence

from jsonschema import Draft202012Validator
from openpyxl import load_workbook

from evaluation.scripts.audit_fairds_deliverable import audit_run


SCHEMA_PATH = Path(__file__).parents[1] / "schemas" / "fairds_expected_outcome.schema.json"
ID_FIELDS = {
    "investigation": "investigation identifier",
    "study": "study identifier",
    "observationunit": "observation unit identifier",
    "sample": "sample identifier",
    "assay": "assay identifier",
}


def _load_object(path: Path) -> Dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"Expected JSON object: {path}")
    return value


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _error(errors: List[Dict[str, Any]], code: str, message: str, **details: Any) -> None:
    item: Dict[str, Any] = {"code": code, "message": message}
    if details:
        item["details"] = details
    errors.append(item)


def _rows(matrix: Mapping[str, Any], level: str) -> List[Mapping[str, Any]]:
    return [
        row
        for row in ((matrix.get(level) or {}).get("rows") or [])
        if isinstance(row, Mapping)
    ]


def _columns(matrix: Mapping[str, Any], level: str) -> List[str]:
    return [str(item).strip() for item in ((matrix.get(level) or {}).get("columns") or [])]


def _validate_spec(spec: Mapping[str, Any]) -> None:
    schema = _load_object(SCHEMA_PATH)
    errors = sorted(Draft202012Validator(schema).iter_errors(spec), key=lambda item: list(item.path))
    if errors:
        rendered = "; ".join(error.message for error in errors[:10])
        raise ValueError(f"Invalid expected-outcome contract: {rendered}")


def evaluate_expected_outcome(
    run_dir: Path,
    expectation_path: Path,
    *,
    project_root: Path,
) -> Dict[str, Any]:
    run_dir = run_dir.resolve()
    expectation_path = expectation_path.resolve()
    project_root = project_root.resolve()
    spec = _load_object(expectation_path)
    _validate_spec(spec)
    errors: List[Dict[str, Any]] = []

    source_assets = [spec["source"], *(spec.get("source_assets") or [])]
    checked_assets: set[str] = set()
    for asset in source_assets:
        relative_path = str(asset["path"])
        if relative_path in checked_assets:
            continue
        checked_assets.add(relative_path)
        source_path = (project_root / relative_path).resolve()
        if not source_path.is_file():
            _error(
                errors,
                "source_missing",
                "Expected source asset is missing.",
                path=str(source_path),
            )
            continue
        digest = hashlib.sha256(source_path.read_bytes()).hexdigest()
        if digest != asset["sha256"]:
            _error(
                errors,
                "source_checksum_mismatch",
                "Expected-outcome fixture does not match the current source asset.",
                path=relative_path,
                expected=asset["sha256"],
                actual=digest,
            )

    deliverables = run_dir / "deliverables"
    if not deliverables.is_dir():
        deliverables = run_dir
    metadata_path = deliverables / "metadata.json"
    matrix_path = deliverables / "isa_values.json"
    if not metadata_path.is_file() or not matrix_path.is_file():
        _error(errors, "deliverables_missing", "metadata.json or isa_values.json is missing.")
        return _finish(spec, run_dir, expectation_path, audit_run(run_dir), errors)
    metadata = _load_object(metadata_path)
    matrix = _load_object(matrix_path)
    audit = audit_run(run_dir)

    packages = {_text(item) for item in metadata.get("packages_used") or [] if _text(item)}
    required_packages = set(spec.get("packages", {}).get("required") or [])
    forbidden_packages = set(spec.get("packages", {}).get("forbidden") or [])
    if required_packages - packages:
        _error(
            errors,
            "required_packages_missing",
            "Required FAIR-DS packages were not selected.",
            missing=sorted(required_packages - packages),
        )
    for alternatives in spec.get("packages", {}).get("required_any_of") or []:
        if not packages.intersection(alternatives):
            _error(
                errors,
                "required_package_alternative_missing",
                "None of the acceptable FAIR-DS package alternatives was selected.",
                alternatives=alternatives,
            )
    if forbidden_packages & packages:
        _error(
            errors,
            "forbidden_packages_selected",
            "Forbidden FAIR-DS packages were selected.",
            packages=sorted(forbidden_packages & packages),
        )

    sheet_names = {
        (
            str(name).split(" - ", 1)[0].strip().lower()
            if str(name).strip().lower() != "help"
            else "Help"
        )
        for name in audit.get("metrics", {}).get("sheet_names") or []
    }
    def normalized_sheet_name(value: Any) -> str:
        text = str(value).split(" - ", 1)[0].strip()
        return "Help" if text.lower() == "help" else text.lower()

    required_sheet_names = {
        normalized_sheet_name(name) for name in spec.get("required_sheets") or []
    }
    forbidden_sheet_names = {
        normalized_sheet_name(name) for name in spec.get("forbidden_sheets") or []
    }
    missing_sheets = required_sheet_names - sheet_names
    forbidden_sheets = forbidden_sheet_names & sheet_names
    if missing_sheets:
        _error(errors, "required_sheets_missing", "Required sheets are missing.", sheets=sorted(missing_sheets))
    if forbidden_sheets:
        _error(errors, "forbidden_sheets_present", "Forbidden sheets are present.", sheets=sorted(forbidden_sheets))

    for level, level_spec in (spec.get("levels") or {}).items():
        rows = _rows(matrix, level)
        columns = _columns(matrix, level)
        if "row_count" in level_spec and len(rows) != int(level_spec["row_count"]):
            _error(
                errors,
                "row_count_mismatch",
                f"{level} row count differs from expectation.",
                level=level,
                expected=int(level_spec["row_count"]),
                actual=len(rows),
            )
        missing_columns = set(level_spec.get("required_columns") or []) - set(columns)
        extra_forbidden = set(level_spec.get("forbidden_columns") or []) & set(columns)
        if missing_columns:
            _error(errors, "required_columns_missing", f"{level} is missing expected columns.", level=level, columns=sorted(missing_columns))
        if extra_forbidden:
            _error(errors, "forbidden_columns_present", f"{level} contains forbidden columns.", level=level, columns=sorted(extra_forbidden))
        for field in level_spec.get("unique_fields") or []:
            values = [_text(row.get(field)) for row in rows]
            if any(not value for value in values) or len(values) != len(set(values)):
                _error(errors, "field_not_unique", f"{level}.{field} is blank or non-unique.", level=level, field=field)
        for expected_row in level_spec.get("expected_rows") or []:
            if not any(all(_text(row.get(field)) == _text(value) for field, value in expected_row.items()) for row in rows):
                _error(errors, "expected_row_missing", f"{level} lacks an expected row.", level=level, expected_row=expected_row)
        for field, field_spec in (level_spec.get("fields") or {}).items():
            values = [_text(row.get(field)) for row in rows]
            nonempty = [value for value in values if value]
            if "nonempty_count" in field_spec and len(nonempty) != int(field_spec["nonempty_count"]):
                _error(errors, "nonempty_count_mismatch", f"{level}.{field} non-empty count differs.", level=level, field=field, expected=int(field_spec["nonempty_count"]), actual=len(nonempty))
            if "min_nonempty_count" in field_spec and len(nonempty) < int(field_spec["min_nonempty_count"]):
                _error(errors, "minimum_nonempty_count_not_met", f"{level}.{field} has insufficient source-supported coverage.", level=level, field=field, expected_minimum=int(field_spec["min_nonempty_count"]), actual=len(nonempty))
            if "max_nonempty_count" in field_spec and len(nonempty) > int(field_spec["max_nonempty_count"]):
                _error(errors, "maximum_nonempty_count_exceeded", f"{level}.{field} exceeds the permitted coverage boundary.", level=level, field=field, expected_maximum=int(field_spec["max_nonempty_count"]), actual=len(nonempty))
            blank_count = len(values) - len(nonempty)
            if "blank_count" in field_spec and blank_count != int(field_spec["blank_count"]):
                _error(errors, "blank_count_mismatch", f"{level}.{field} blank count differs.", level=level, field=field, expected=int(field_spec["blank_count"]), actual=blank_count)
            if "distinct_nonempty_count" in field_spec and len(set(nonempty)) != int(field_spec["distinct_nonempty_count"]):
                _error(errors, "distinct_value_count_mismatch", f"{level}.{field} distinct-value count differs.", level=level, field=field, expected=int(field_spec["distinct_nonempty_count"]), actual=len(set(nonempty)))
            if "value_counts" in field_spec:
                actual_counts = Counter(values)
                expected_counts = Counter({str(key): int(value) for key, value in field_spec["value_counts"].items()})
                if actual_counts != expected_counts:
                    _error(errors, "value_distribution_mismatch", f"{level}.{field} value distribution differs.", level=level, field=field, expected=dict(expected_counts), actual=dict(actual_counts))
            if "allowed_values" in field_spec:
                disallowed = sorted(set(nonempty) - set(field_spec["allowed_values"]))
                if disallowed:
                    _error(errors, "disallowed_values", f"{level}.{field} contains unexpected values.", level=level, field=field, values=disallowed[:20])
            for substring in field_spec.get("required_substrings") or []:
                if not any(str(substring).casefold() in value.casefold() for value in nonempty):
                    _error(errors, "required_substring_missing", f"{level}.{field} lacks required source-supported content.", level=level, field=field, substring=substring)
            if field_spec.get("value_regex"):
                pattern = re.compile(str(field_spec["value_regex"]))
                invalid = [value for value in nonempty if pattern.fullmatch(value) is None]
                if invalid:
                    _error(errors, "field_value_pattern_mismatch", f"{level}.{field} contains values outside the expected pattern.", level=level, field=field, values=invalid[:20])

    for relationship in spec.get("relationships") or []:
        child_level = relationship["child_level"]
        parent_level = relationship["parent_level"]
        link_field = relationship["link_field"]
        parent_id_field = ID_FIELDS[parent_level]
        parents = {_text(row.get(parent_id_field)): row for row in _rows(matrix, parent_level)}
        children = _rows(matrix, child_level)
        if relationship.get("one_child_per_parent"):
            link_counts = Counter(_text(row.get(link_field)) for row in children)
            failures = sorted(parent_id for parent_id in parents if link_counts[parent_id] != 1)
            if failures or len(children) != len(parents):
                _error(errors, "relationship_not_one_to_one", f"{child_level}->{parent_level} is not one-to-one.", child_level=child_level, parent_level=parent_level, failures=failures[:20])
        for equality in relationship.get("linked_field_equalities") or []:
            compared = 0
            mismatches: List[Dict[str, str]] = []
            for child in children:
                parent_id = _text(child.get(link_field))
                parent = parents.get(parent_id)
                if parent is None:
                    continue
                parent_value = _text(parent.get(equality["parent_field"]))
                if equality.get("compare_when_parent_nonempty") and not parent_value:
                    continue
                child_value = _text(child.get(equality["child_field"]))
                compared += 1
                if child_value != parent_value:
                    mismatches.append({"parent_id": parent_id, "parent": parent_value, "child": child_value})
            if "expected_comparison_count" in equality and compared != int(equality["expected_comparison_count"]):
                _error(errors, "linked_comparison_count_mismatch", "Linked-field comparison count differs.", expected=int(equality["expected_comparison_count"]), actual=compared, child_field=equality["child_field"], parent_field=equality["parent_field"])
            if mismatches:
                _error(errors, "linked_field_mismatch", "Linked parent/child field values differ.", child_field=equality["child_field"], parent_field=equality["parent_field"], count=len(mismatches), examples=mismatches[:20])

    for check in spec.get("source_table_checks") or []:
        table_path = (project_root / check["path"]).resolve()
        comparison = check.get("comparison", "exact_multiset")
        source_values: List[str] = []
        source_by_key: Dict[str, str] = {}
        try:
            workbook = load_workbook(table_path, read_only=True, data_only=True)
            worksheet = workbook[check["sheet"]]
            header_values = next(
                worksheet.iter_rows(
                    min_row=int(check["header_row"]),
                    max_row=int(check["header_row"]),
                    values_only=True,
                )
            )
            headers = [_text(value) for value in header_values]
            required_source_fields = {check["source_field"]}
            if check.get("filter"):
                required_source_fields.add(check["filter"]["field"])
            if check.get("source_key_field"):
                required_source_fields.add(check["source_key_field"])
            missing_source_fields = required_source_fields - set(headers)
            if missing_source_fields:
                _error(
                    errors,
                    "source_table_fields_missing",
                    "A source-table expectation references missing columns.",
                    path=str(table_path),
                    columns=sorted(missing_source_fields),
                )
                continue
            indexes = {name: headers.index(name) for name in required_source_fields}
            for values in worksheet.iter_rows(
                min_row=int(check["header_row"]) + 1, values_only=True
            ):
                table_filter = check.get("filter")
                if table_filter and _text(
                    values[indexes[table_filter["field"]]]
                ) != _text(table_filter["value"]):
                    continue
                value = _text(values[indexes[check["source_field"]]])
                value = (check.get("source_value_map") or {}).get(value, value)
                if comparison in {"exact_keyed", "numeric_keyed"}:
                    key = _text(values[indexes[check["source_key_field"]]])
                    if not key:
                        _error(
                            errors,
                            "source_table_key_blank",
                            "A keyed source-table check has a blank source key.",
                            path=str(table_path),
                            source_key_field=check["source_key_field"],
                        )
                    elif key in source_by_key:
                        _error(
                            errors,
                            "source_table_key_duplicate",
                            "A keyed source-table check has duplicate source keys.",
                            path=str(table_path),
                            source_key_field=check["source_key_field"],
                            key=key,
                        )
                    else:
                        source_by_key[key] = value
                elif value:
                    source_values.append(value)
            workbook.close()
        except (OSError, KeyError, StopIteration, ValueError) as exc:
            _error(
                errors,
                "source_table_unreadable",
                "A source-table expectation could not be evaluated.",
                path=str(table_path),
                error=str(exc),
            )
            continue

        target_rows = _rows(matrix, check["target_level"])
        if comparison in {"exact_keyed", "numeric_keyed"}:
            lookup = check.get("target_key_lookup")
            lookup_values: Dict[str, str] = {}
            if lookup:
                for row in _rows(matrix, lookup["level"]):
                    lookup_id = _text(row.get(lookup["lookup_id_field"]))
                    lookup_value = _text(row.get(lookup["lookup_value_field"]))
                    if lookup_id:
                        lookup_values[lookup_id] = lookup_value
            target_by_key: Dict[str, str] = {}
            duplicate_target_keys: List[str] = []
            for row in target_rows:
                if lookup:
                    key = lookup_values.get(
                        _text(row.get(lookup["target_link_field"])), ""
                    )
                else:
                    key = _text(row.get(check["target_key_field"]))
                if not key:
                    continue
                if key in target_by_key:
                    duplicate_target_keys.append(key)
                target_by_key[key] = _text(row.get(check["target_field"]))
            if duplicate_target_keys:
                _error(
                    errors,
                    "target_key_duplicate",
                    "A keyed source-table check has duplicate target keys.",
                    target_level=check["target_level"],
                    keys=sorted(set(duplicate_target_keys))[:20],
                )
            missing_keys = sorted(set(source_by_key) - set(target_by_key))
            unexpected_keys = sorted(set(target_by_key) - set(source_by_key))
            mismatches: List[Dict[str, str]] = []
            for key in sorted(set(source_by_key) & set(target_by_key)):
                source_value = source_by_key[key]
                target_value = target_by_key[key]
                if comparison == "numeric_keyed" and source_value and target_value:
                    try:
                        matches = math.isclose(
                            float(source_value),
                            float(target_value),
                            rel_tol=float(check.get("relative_tolerance", 0.0)),
                            abs_tol=float(check.get("absolute_tolerance", 0.0)),
                        )
                    except ValueError:
                        matches = False
                else:
                    matches = source_value == target_value
                if not matches:
                    mismatches.append(
                        {
                            "key": key,
                            "source": source_value,
                            "target": target_value,
                        }
                    )
            if missing_keys or unexpected_keys or mismatches:
                _error(
                    errors,
                    "source_table_keyed_projection_mismatch",
                    "Deliverable values do not match their keyed source-table records.",
                    source_field=check["source_field"],
                    target_level=check["target_level"],
                    target_field=check["target_field"],
                    source_count=len(source_by_key),
                    target_count=len(target_by_key),
                    missing_keys=missing_keys[:20],
                    unexpected_keys=unexpected_keys[:20],
                    value_mismatches=mismatches[:20],
                )
            continue

        target_values = [
            _text(row.get(check["target_field"]))
            for row in target_rows
            if _text(row.get(check["target_field"]))
        ]
        matches = set(source_values) == set(target_values) if comparison == "exact_set" else Counter(source_values) == Counter(target_values)
        if not matches:
            _error(
                errors,
                "source_table_projection_mismatch",
                "Deliverable values do not exactly match the selected source-table records.",
                source_field=check["source_field"],
                target_level=check["target_level"],
                target_field=check["target_field"],
                source_count=len(source_values),
                target_count=len(target_values),
                missing=sorted(set(source_values) - set(target_values))[:20],
                unexpected=sorted(set(target_values) - set(source_values))[:20],
            )

    allowed_warnings = set(spec.get("allowed_audit_warning_codes") or [])
    unexpected_warnings = [issue for issue in audit.get("issues") or [] if issue.get("severity") == "warning" and issue.get("code") not in allowed_warnings]
    if unexpected_warnings:
        _error(errors, "unexpected_audit_warnings", "Deliverable has warnings outside the expected review boundary.", codes=sorted({item.get("code") for item in unexpected_warnings}))
    return _finish(spec, run_dir, expectation_path, audit, errors)


def _finish(
    spec: Mapping[str, Any],
    run_dir: Path,
    expectation_path: Path,
    audit: Mapping[str, Any],
    errors: List[Dict[str, Any]],
) -> Dict[str, Any]:
    structural_passed = bool(audit.get("structural_passed", False))
    asserted_fields = sum(
        len((level_spec or {}).get("fields") or {})
        for level_spec in (spec.get("levels") or {}).values()
    )
    return {
        "schema_version": "fairiagent.fairds_expected_outcome_result.v1",
        "case_id": spec.get("case_id"),
        "run_dir": str(run_dir),
        "expectation_path": str(expectation_path),
        "passed": structural_passed and not errors,
        "review_required": bool(audit.get("review_required", False)),
        "generic_audit": {
            "structural_passed": structural_passed,
            "error_count": audit.get("error_count"),
            "warning_count": audit.get("warning_count"),
        },
        "contract_coverage": {
            "pinned_source_asset_count": len(
                {
                    str(item.get("path"))
                    for item in [spec.get("source") or {}, *(spec.get("source_assets") or [])]
                    if item.get("path")
                }
            ),
            "asserted_level_count": len(spec.get("levels") or {}),
            "asserted_field_count": asserted_fields,
            "source_table_check_count": len(spec.get("source_table_checks") or []),
            "coverage_ledger_entry_count": len(spec.get("coverage_ledger") or []),
        },
        "expectation_error_count": len(errors),
        "expectation_errors": errors,
    }


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("run_dir", type=Path)
    parser.add_argument("expectation", type=Path)
    parser.add_argument("--project-root", type=Path, default=Path.cwd())
    parser.add_argument("--output", type=Path)
    args = parser.parse_args(argv)
    report = evaluate_expected_outcome(
        args.run_dir,
        args.expectation,
        project_root=args.project_root,
    )
    rendered = json.dumps(report, indent=2, ensure_ascii=False)
    print(rendered)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(rendered + "\n", encoding="utf-8")
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    sys.exit(main())
