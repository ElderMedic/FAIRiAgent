from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

import fairifier.agents.entity_structure_planner as entity_planner_module
from fairifier.agents.entity_structure_planner import EntityStructurePlannerAgent
from fairifier.agents.isa_value_mapper import ISAValueMapperAgent
from fairifier.agents.response_models import EntityPlanScopeAuditResponse
from fairifier.graph.excel import _generate_xlsx_local
from fairifier.graph.nodes import OrchestrateNode
from fairifier.utils.entity_plan import (
    contacts_from_source_text,
    materialize_entity_claims,
    normalize_claim_measurement_units,
    normalize_source_measurement_value,
    normalize_entity_plan,
    project_matrix_onto_entity_plan,
    validate_entity_matrix_against_plan,
)
from fairifier.utils.isa_matrix_projection import apply_matrix_to_metadata


MICHAEL_SOURCE = """Contributor(s)\tNodine M, Hofmann F, Schon M
Citation(s)\tHofmann F, Schon MA, Nodine MD (2019) RNA-seq analysis.
Overall design\tEight stages of embryo development with three biological
replicates. Bent cotyledon RNA was tested at four input amounts with three
library preparation methods.
"""


def test_entity_planner_reuses_provider_configured_model_without_local_binding(
    monkeypatch,
):
    class ProviderConfiguredModel:
        def bind(self, **kwargs):
            raise AssertionError(f"agent-local sampler binding is forbidden: {kwargs}")

    class SharedHelper:
        def __init__(self):
            self.model = ProviderConfiguredModel()
            self.llm = self.model
            self.llm_responses = []

        def get_llm(self):
            return self.model

    shared_helper = SharedHelper()
    monkeypatch.setattr(
        entity_planner_module, "get_llm_helper", lambda: shared_helper
    )

    planner = EntityStructurePlannerAgent()

    assert planner.llm_helper.llm is shared_helper.model
    assert planner.llm_helper.llm_responses is shared_helper.llm_responses


def _embryo_study_claims():
    stages = [
        "pre-globular",
        "globular",
        "heart",
        "torpedo",
        "bent cotyledon",
        "mature green",
        "post-mature green",
        "dry seed",
    ]
    return {
        "investigations": [
            {
                "row_id": "investigation_001",
                "label": "GSE126024",
                "parent_row_id": None,
                "external_identifier": "GSE126024",
                "evidence": "GEO series GSE126024",
            }
        ],
        "studies": [
            {
                "row_id": "study_001",
                "label": "Arabidopsis embryo RNA-seq",
                "parent_row_id": "investigation_001",
                "external_identifier": None,
                "evidence": "Overall design",
            }
        ],
        "design_groups": [
            {
                "group_id": "development",
                "study_row_id": "study_001",
                "evidence": "Eight stages with three biological replicates",
                "dimensions": [
                    {
                        "name": "developmental_stage",
                        "level_count": 8,
                        "explicit_values": stages,
                        "origin": "explicit",
                        "field_name": "developmental stage",
                        "applies_to": ["observationunit", "sample", "assay"],
                    },
                    {
                        "name": "biological_replicate",
                        "level_count": 3,
                        "explicit_values": [],
                        "origin": "derived",
                        "field_name": "biological replicate",
                        "applies_to": ["sample", "assay"],
                    },
                ],
                "unresolved_ambiguities": [],
            },
            {
                "group_id": "low_input",
                "study_row_id": "study_001",
                "evidence": "Four input amounts crossed with three methods",
                "dimensions": [
                    {
                        "name": "source_stage",
                        "level_count": 1,
                        "explicit_values": ["bent cotyledon"],
                        "origin": "explicit",
                        "field_name": "developmental stage",
                        "applies_to": ["observationunit", "sample", "assay"],
                    },
                    {
                        "name": "input_amount",
                        "level_count": 4,
                        "explicit_values": ["5 ng", "1 ng", "0.5 ng", "0.1 ng"],
                        "origin": "explicit",
                        "field_name": "input amount",
                        "applies_to": ["sample", "assay"],
                    },
                    {
                        "name": "library_method",
                        "level_count": 3,
                        "explicit_values": [],
                        "origin": "derived",
                        "field_name": "library preparation method",
                        "applies_to": ["sample", "assay"],
                    },
                ],
                "unresolved_ambiguities": ["Method names are not in the abstract."],
            },
        ],
        "design_summary": [],
        "unresolved_ambiguities": [],
        "confidence": 0.9,
    }


def _embryo_study_plan():
    contacts = contacts_from_source_text(MICHAEL_SOURCE)
    return normalize_entity_plan(
        materialize_entity_claims(_embryo_study_claims()),
        investigation_contacts=contacts,
    )


def test_contacts_use_parser_identities_and_enrich_only_explicit_details():
    source = """Authors: Alice M. Example $^{1}$ and Bob Tester $^{1}$

Affiliations: $^{1}$ Example Research Institute, Example City

Corresponding author: Alice M. Example (email: alice.example@institute.test)
"""

    contacts = contacts_from_source_text(
        source,
        fallback_names=["Alice M. Example", "Bob Tester"],
    )

    assert [(row["firstname"], row["lastname"]) for row in contacts] == [
        ("Alice M.", "Example"),
        ("Bob", "Tester"),
    ]
    assert contacts[0]["email address"] == "alice.example@institute.test"
    assert contacts[1]["email address"] == ""
    assert {row["organization"] for row in contacts} == {
        "Example Research Institute, Example City"
    }


def test_contacts_do_not_guess_organization_from_competing_affiliations():
    source = """Authors: Alice Example and Bob Tester
Affiliations: First University; Second Research Institute
"""

    contacts = contacts_from_source_text(
        source,
        fallback_names=["Alice Example", "Bob Tester"],
    )

    assert all(row["organization"] == "" for row in contacts)


def _seed_matrix():
    return {
        "investigation": {
            "columns": [
                "investigation identifier",
                "investigation title",
                "investigation description",
                "firstname",
                "lastname",
                "email address",
                "orcid",
                "organization",
                "department",
            ],
            "rows": [{"investigation identifier": "GSE126024"}],
            "fields": [
                {
                    "field_name": "investigation identifier",
                    "package_source": "default",
                    "requirement": "MANDATORY",
                }
            ],
        },
        "study": {
            "columns": ["study identifier", "study title", "investigation identifier"],
            "rows": [{"study identifier": "GSE126024"}],
            "fields": [{"field_name": "study identifier", "package_source": "default"}],
        },
        "observationunit": {
            "columns": [
                "observation unit identifier",
                "observation unit name",
                "study identifier",
                "developmental stage",
            ],
            "rows": [{"observation unit identifier": "collapsed stages"}],
            "fields": [
                {"field_name": "observation unit identifier", "package_source": "default"}
            ],
        },
        "sample": {
            "columns": [
                "sample identifier",
                "sample name",
                "observation unit identifier",
                "developmental stage",
                "biological replicate",
                "input amount",
            ],
            "rows": [{"sample identifier": "sample_a, sample_b"}],
            "fields": [
                {
                    "field_name": "sample identifier",
                    "package_source": "Crop Plant sample enhanced annotation checklist",
                }
            ],
        },
        "assay": {
            "columns": ["assay identifier", "assay description", "sample identifier"],
            "rows": [{"assay identifier": "GSE126024"}],
            "fields": [{"field_name": "assay identifier", "package_source": "default"}],
        },
    }


def test_embryo_study_design_materializes_fairds_cardinalities_and_links():
    plan = _embryo_study_plan()
    counts = {level["level"]: level["cardinality"] for level in plan["levels"]}

    assert counts == {
        "investigation": 1,
        "study": 1,
        "observationunit": 9,
        "sample": 36,
        "assay": 36,
    }
    assert plan["materialization_errors"] == []
    samples = next(
        level["entities"] for level in plan["levels"] if level["level"] == "sample"
    )
    assays = next(
        level["entities"] for level in plan["levels"] if level["level"] == "assay"
    )
    sample_ids = {entity["row_id"] for entity in samples}
    assert all(entity["parent_row_id"] in sample_ids for entity in assays)


def test_source_declared_unit_symbol_normalizes_standalone_quantities():
    source = "Inputs ranged from 0.1 to 5 nanograms (ng) of total RNA."

    assert normalize_source_measurement_value("5 nanograms", source) == "5 ng"
    assert normalize_source_measurement_value("0.1 nanograms", source) == "0.1 ng"
    assert normalize_source_measurement_value("1 nanogram", source) == "1 ng"
    assert normalize_source_measurement_value("five nanograms", source) == (
        "five nanograms"
    )
    assert normalize_source_measurement_value("5 nanograms total RNA", source) == (
        "5 nanograms total RNA"
    )


def test_claim_quantities_use_source_declared_unit_symbol():
    claims = {
        "design_groups": [
            {
                "dimensions": [
                    {
                        "explicit_values": ["5 nanograms", "1 nanograms"],
                        "known_values": [
                            {"position": 1, "value": "0.5 nanograms"}
                        ],
                    }
                ]
            }
        ]
    }

    normalized = normalize_claim_measurement_units(
        claims, "nanograms (ng)"
    )

    dimension = normalized["design_groups"][0]["dimensions"][0]
    assert dimension["explicit_values"] == ["5 ng", "1 ng"]
    assert dimension["known_values"][0]["value"] == "0.5 ng"


def test_unnamed_derived_singleton_does_not_pollute_entity_labels():
    claims = _embryo_study_claims()
    claims["design_groups"][0]["dimensions"].append(
        {
            "name": "unnamed_method",
            "level_count": 1,
            "explicit_values": [],
            "known_values": [],
            "origin": "derived",
            "field_mappings": [],
            "applies_to": ["sample", "assay"],
        }
    )

    plan = materialize_entity_claims(claims)
    samples = next(
        item["entities"]
        for item in plan["levels"]
        if item["level"] == "sample"
    )

    assert len(samples) == 36
    assert all("unnamed method" not in entity["label"] for entity in samples)
    assert all(
        all(
            attribute["dimension_name"] != "unnamed_method"
            for attribute in entity["attributes"]
        )
        for entity in samples
    )
    sample_level = next(
        item for item in plan["levels"] if item["level"] == "sample"
    )
    assert any(
        "one unnamed derived level" in ambiguity
        for ambiguity in sample_level["unresolved_ambiguities"]
    )


def test_child_links_on_shared_dimensions_when_parent_has_level_only_constant():
    claims = _embryo_study_claims()
    for group in claims["design_groups"]:
        group["dimensions"].append(
            {
                "name": "sample_only_property",
                "level_count": 1,
                "explicit_values": ["source-literal"],
                "origin": "explicit",
                "field_name": "sampling strategy",
                "applies_to": ["sample"],
            }
        )

    materialized = materialize_entity_claims(claims)

    assert materialized["materialization_errors"] == []
    assays = next(
        level["entities"]
        for level in materialized["levels"]
        if level["level"] == "assay"
    )
    assert all(entity["parent_row_id"] for entity in assays)


def test_shared_dimension_parent_join_fails_when_parent_only_factor_is_varying():
    claims = _embryo_study_claims()
    claims["design_groups"][0]["dimensions"].append(
        {
            "name": "unpropagated_parent_factor",
            "level_count": 2,
            "explicit_values": ["left", "right"],
            "origin": "explicit",
            "field_name": "sampling strategy",
            "applies_to": ["sample"],
        }
    )

    materialized = materialize_entity_claims(claims)

    assert any(
        "matches 2 sample parents" in error
        for error in materialized["materialization_errors"]
    )


def test_contacts_repeat_in_investigation_without_person_level():
    plan = _embryo_study_plan()
    matrix = project_matrix_onto_entity_plan(_seed_matrix(), plan)
    validation = validate_entity_matrix_against_plan(matrix, plan)

    assert validation["passed"] is True
    assert validation["investigation_contact_count"] == 3
    assert "person" not in matrix
    assert [
        (row["firstname"], row["lastname"])
        for row in matrix["investigation"]["rows"]
    ] == [("MD", "Nodine"), ("F", "Hofmann"), ("MA", "Schon")]
    assert validation["row_counts"] == {
        "investigation": 3,
        "study": 1,
        "observationunit": 9,
        "sample": 36,
        "assay": 36,
    }
    assert all(
        "," not in row["sample identifier"] for row in matrix["sample"]["rows"]
    )
    assert "assay name" not in matrix["assay"]["columns"]


def test_selected_assay_name_is_compiled_from_each_planned_entity():
    seed = _seed_matrix()
    seed["assay"]["columns"].append("assay name")
    seed["assay"]["rows"][0]["assay name"] = "one broadcast example"

    plan = _embryo_study_plan()
    matrix = project_matrix_onto_entity_plan(seed, plan)
    validation = validate_entity_matrix_against_plan(matrix, plan)
    assay_entities = next(
        level["entities"] for level in plan["levels"] if level["level"] == "assay"
    )

    assert validation["passed"] is True
    assert [row["assay name"] for row in matrix["assay"]["rows"]] == [
        entity["label"] for entity in assay_entities
    ]


def test_empty_mapper_level_retains_confirmed_shared_fields_after_plan_projection():
    from fairifier.agents.isa_value_mapper import ISAValueMapperAgent

    seed = _seed_matrix()
    seed["assay"]["columns"].append("library strategy")
    seed["assay"]["rows"] = []
    plan = _embryo_study_plan()
    fields = {
        "assay": [
            {
                "field_name": "library strategy",
                "value": "RNA-Seq",
                "confidence": 0.95,
                "status": "confirmed",
                "evidence": "source_001 explicitly states mRNA sequencing",
                "value_scope": "level",
            }
        ]
    }

    seed = ISAValueMapperAgent._seed_missing_planned_levels(seed, plan)
    seed = ISAValueMapperAgent._restore_high_confidence_shared_values(
        seed, fields, entity_plan=plan
    )
    matrix = project_matrix_onto_entity_plan(seed, plan)

    assert len(matrix["assay"]["rows"]) == 36
    assert {
        row["library strategy"] for row in matrix["assay"]["rows"]
    } == {"RNA-Seq"}


def test_materialized_plan_retains_unmapped_design_dimensions_in_every_entity():
    claims = _embryo_study_claims()
    claims["design_groups"][0]["dimensions"][1]["field_name"] = None

    plan = normalize_entity_plan(materialize_entity_claims(claims))
    samples = next(
        level["entities"] for level in plan["levels"] if level["level"] == "sample"
    )
    development = [row for row in samples if row["source_group"] == "development"]

    assert len(development) == 24
    assert all(
        any(
            attribute["dimension_name"] == "biological_replicate"
            and attribute["field_name"] is None
            for attribute in row["attributes"]
        )
        for row in development
    )


def test_source_native_name_need_not_repeat_dimension_present_in_description():
    plan = _embryo_study_plan()
    samples = next(
        level["entities"] for level in plan["levels"] if level["level"] == "sample"
    )
    samples[0]["label"] = "source-native-sample-id"
    samples[0]["external_identifier"] = "source-native-sample-id"
    samples[0]["attributes"].append(
        {
            "dimension_name": "source series",
            "field_name": None,
            "value": "developmental time course",
            "origin": "explicit",
        }
    )

    matrix = project_matrix_onto_entity_plan(_seed_matrix(), plan)
    validation = validate_entity_matrix_against_plan(matrix, plan)

    first = matrix["sample"]["rows"][0]
    assert first["sample name"] == "source-native-sample-id"
    assert "developmental time course" in first["sample description"]
    assert validation["passed"] is True


def test_ordinal_placeholders_do_not_populate_semantic_fairds_terms():
    claims = _embryo_study_claims()
    claims["design_groups"][1]["dimensions"][2]["field_mappings"] = [
        {"level": "assay", "field_name": "protocol"}
    ]

    plan = materialize_entity_claims(
        claims,
        field_catalog={"assay": ["protocol"]},
    )
    assays = next(
        level["entities"] for level in plan["levels"] if level["level"] == "assay"
    )
    low_input = [row for row in assays if row["source_group"] == "low_input"]

    assert low_input
    assert all(
        any(
            attribute["dimension_name"] == "library_method"
            and attribute["origin"] == "derived"
            and attribute["field_name"] is None
            for attribute in row["attributes"]
        )
        for row in low_input
    )


def test_partial_factor_labels_preserve_declared_cardinality():
    claims = _embryo_study_claims()
    stage = claims["design_groups"][0]["dimensions"][0]
    stage["explicit_values"] = []
    stage["known_values"] = [
        {"position": 1, "value": "8-cell/16-cell"},
        {"position": 8, "value": "mature green"},
    ]
    stage["origin"] = "derived"

    plan = materialize_entity_claims(claims)
    counts = {item["level"]: item["cardinality"] for item in plan["levels"]}
    stage_entities = [
        row
        for item in plan["levels"]
        if item["level"] == "observationunit"
        for row in item["entities"]
        if row["source_group"] == "development"
    ]

    assert plan["materialization_errors"] == []
    assert counts == {
        "investigation": 1,
        "study": 1,
        "observationunit": 9,
        "sample": 36,
        "assay": 36,
    }
    assert len(stage_entities) == 8
    stage_attributes = [
        next(
            attribute
            for attribute in row["attributes"]
            if attribute["dimension_name"] == "developmental_stage"
        )
        for row in stage_entities
    ]
    assert stage_attributes[0] == {
        "dimension_name": "developmental_stage",
        "field_name": "developmental stage",
        "value": "8-cell/16-cell",
        "origin": "explicit",
    }
    assert stage_attributes[-1]["value"] == "mature green"
    assert stage_attributes[-1]["origin"] == "explicit"
    assert all(item["origin"] == "derived" for item in stage_attributes[1:-1])
    assert all(item["field_name"] is None for item in stage_attributes[1:-1])


def test_dimension_terms_are_validated_separately_for_each_isa_level():
    claims = _embryo_study_claims()
    claims["design_groups"][0]["dimensions"][0]["field_mappings"] = [
        {"level": "observationunit", "field_name": "study treatment"},
        {"level": "sample", "field_name": "study treatment"},
        {"level": "assay", "field_name": "study treatment"},
    ]
    claims["design_groups"][0]["dimensions"][0]["field_name"] = None
    catalog = {
        "observationunit": ["study treatment"],
        "sample": ["sample treatment"],
        "assay": ["protocol"],
    }

    plan = normalize_entity_plan(
        materialize_entity_claims(claims, field_catalog=catalog)
    )
    by_level = {
        level["level"]: level["entities"] for level in plan["levels"]
    }

    assert any(
        attribute["field_name"] == "study treatment"
        for attribute in by_level["observationunit"][0]["attributes"]
    )
    assert all(
        attribute["field_name"] is None
        for entity in by_level["sample"]
        if entity["source_group"] == "development"
        for attribute in entity["attributes"]
        if attribute["dimension_name"] == "developmental_stage"
    )
    assert all(
        attribute["field_name"] is None
        for entity in by_level["assay"]
        if entity["source_group"] == "development"
        for attribute in entity["attributes"]
        if attribute["dimension_name"] == "developmental_stage"
    )


def test_design_dimensions_cannot_overwrite_reserved_identity_fields():
    claims = _embryo_study_claims()
    dimension = claims["design_groups"][0]["dimensions"][1]
    dimension["field_name"] = None
    dimension["field_mappings"] = [
        {"level": "sample", "field_name": "sample name"}
    ]

    plan = normalize_entity_plan(
        materialize_entity_claims(
            claims,
            field_catalog={
                "observationunit": ["observation unit name"],
                "sample": ["sample name"],
                "assay": ["assay description"],
            },
        )
    )
    samples = next(
        level["entities"]
        for level in plan["levels"]
        if level["level"] == "sample"
    )

    assert all(
        attribute["field_name"] is None
        for entity in samples
        if entity["source_group"] == "development"
        for attribute in entity["attributes"]
        if attribute["dimension_name"] == "biological_replicate"
    )


def test_design_group_label_replaces_internal_group_id_in_entity_names():
    claims = _embryo_study_claims()
    claims["design_groups"][0]["label"] = "developmental-stage series"

    plan = normalize_entity_plan(materialize_entity_claims(claims))
    observations = next(
        level["entities"]
        for level in plan["levels"]
        if level["level"] == "observationunit"
    )

    assert observations[0]["label"].startswith("developmental-stage series")
    assert not observations[0]["label"].startswith("development observationunit")


def test_entity_planner_prompt_renders_level_specific_field_mapping_example():
    planner = object.__new__(EntityStructurePlannerAgent)

    prompt = planner._prompt(
        {
            "document_content": "A two-condition experiment.",
            "document_info": {"title": "Test"},
            "retrieved_knowledge": [],
        }
    )

    assert '"field_mappings"' in prompt
    assert '"level": "observationunit"' in prompt


def test_scope_audit_restores_exact_constant_condition_without_changing_cardinality():
    claims = {
        "design_groups": [
            {
                "group_id": "group_002",
                "evidence": "RNA from bent-cotyledon embryos was diluted to 1 ng.",
                "dimensions": [
                    {
                        "name": "dose",
                        "level_count": 1,
                        "explicit_values": ["1 ng"],
                        "known_values": [],
                    }
                ],
            }
        ]
    }
    audit = {
        "missing_conditions": [
            {
                "group_id": "group_002",
                "dimension_name": "developmental_stage",
                "source_value": "bent-cotyledon",
                "level": "sample",
                "field_name": "dev_stage",
            }
        ]
    }

    added = EntityStructurePlannerAgent._apply_scope_audit_findings(
        claims, audit, {"sample": ["dev_stage"]}
    )

    assert added == ["group_002:sample.dev_stage=bent-cotyledon"]
    restored = claims["design_groups"][0]["dimensions"][-1]
    assert restored["level_count"] == 1
    assert restored["explicit_values"] == ["bent-cotyledon"]
    assert restored["applies_to"] == ["sample", "assay"]


def test_scope_audit_rejects_values_not_verbatim_in_group_evidence():
    claims = {
        "design_groups": [
            {"group_id": "group_001", "evidence": "control samples", "dimensions": []}
        ]
    }
    audit = {
        "missing_conditions": [
            {
                "group_id": "group_001",
                "dimension_name": "genotype",
                "source_value": "invented mutant",
                "level": "sample",
                "field_name": "genotype",
            }
        ]
    }

    added = EntityStructurePlannerAgent._apply_scope_audit_findings(
        claims, audit, {"sample": ["genotype"]}
    )

    assert added == []
    assert claims["design_groups"][0]["dimensions"] == []


def test_plan_validation_rejects_many_assays_per_sample_without_reuse_evidence():
    plan = _embryo_study_plan()
    plan["design_claims"] = {
        "design_groups": [
            {
                "group_id": "low_input",
                "evidence": "Four inputs were prepared with three methods.",
                "reuse_same_sample_for_multiple_assays": False,
                "sample_reuse_evidence": None,
            }
        ]
    }
    sample_level = next(
        item for item in plan["levels"] if item["level"] == "sample"
    )
    assay_level = next(
        item for item in plan["levels"] if item["level"] == "assay"
    )
    sample_level["entities"] = [
        entity
        for entity in sample_level["entities"]
        if entity.get("source_group") != "low_input" or entity["row_id"].endswith("_01")
    ]
    sample_level["cardinality"] = len(sample_level["entities"])

    errors = EntityStructurePlannerAgent._plan_errors(plan)

    assert any("without an explicit source quote" in error for error in errors)


def test_plan_validation_allows_many_assays_when_reuse_quote_is_verbatim():
    plan = _embryo_study_plan()
    plan["design_claims"] = {
        "design_groups": [
            {
                "group_id": "low_input",
                "evidence": "The same identified RNA sample was measured three times.",
                "reuse_same_sample_for_multiple_assays": True,
                "sample_reuse_evidence": "same identified RNA sample",
            }
        ]
    }
    sample_level = next(
        item for item in plan["levels"] if item["level"] == "sample"
    )
    sample_level["entities"] = [
        entity
        for entity in sample_level["entities"]
        if entity.get("source_group") != "low_input" or entity["row_id"].endswith("_01")
    ]
    sample_level["cardinality"] = len(sample_level["entities"])

    errors = EntityStructurePlannerAgent._plan_errors(plan)

    assert not any("without an explicit source quote" in error for error in errors)


def test_plan_validation_rejects_incomplete_authoritative_table_coverage():
    plan = _embryo_study_plan()
    plan["source_coverage"] = {
        "complete": False,
        "tables": [
            {
                "source_path": "metadata.xlsx",
                "table_name": "records",
                "unmapped_columns": ["read count"],
            }
        ],
    }

    errors = EntityStructurePlannerAgent._plan_errors(plan)

    assert any("coverage is incomplete" in error for error in errors)


def test_projection_does_not_broadcast_one_entity_narrative_to_siblings():
    seed = _seed_matrix()
    seed["sample"]["columns"].extend(["sample description", "sampling timepoint"])
    seed["sample"]["rows"] = [
        {
            "sample identifier": "example",
            "sample description": "pre-globular biological replicate 1",
            "sampling timepoint": "pre-globular",
        }
    ]

    matrix = project_matrix_onto_entity_plan(seed, _embryo_study_plan())
    sample_rows = matrix["sample"]["rows"]

    assert len({row["sample description"] for row in sample_rows}) == 36
    assert all(
        "pre-globular" not in row["sampling timepoint"]
        for row in sample_rows
        if "pre-globular" not in row["sample name"]
    )
    assert {
        row["developmental stage"]
        for row in sample_rows
        if row["sample identifier"].startswith("development_")
    } == {
        "pre-globular",
        "globular",
        "heart",
        "torpedo",
        "bent cotyledon",
        "mature green",
        "post-mature green",
        "dry seed",
    }
    assert all(
        row["assay description"] == f"Assay for sample {row['sample identifier']}"
        for row in matrix["assay"]["rows"]
    )


def test_validation_rejects_collapsed_mapped_dimension_values():
    plan = _embryo_study_plan()
    matrix = project_matrix_onto_entity_plan(_seed_matrix(), plan)
    for row in matrix["sample"]["rows"]:
        if row["sample identifier"].startswith("development_"):
            row["developmental stage"] = "pre-globular"

    validation = validate_entity_matrix_against_plan(matrix, plan)

    assert validation["passed"] is False
    assert validation["mapped_attribute_errors"]
    assert any(
        "does not preserve mapped dimension" in error
        for error in validation["errors"]
    )


def test_matrix_projection_removes_legacy_person_level_from_metadata():
    payload = {
        "isa_structure": {
            "investigation": {"fields": []},
            "person": {"fields": [{"field_name": "person name"}]},
        }
    }
    matrix = project_matrix_onto_entity_plan(_seed_matrix(), _embryo_study_plan())

    updated = apply_matrix_to_metadata(payload, matrix)

    assert "person" not in updated["isa_structure"]
    assert "person" not in updated["isa_values"]


def test_contact_details_stay_with_the_matching_investigation_row():
    seed = _seed_matrix()
    seed["investigation"]["rows"] = [
        {
            "investigation identifier": "GSE126024",
            "firstname": "F",
            "lastname": "Hofmann",
            "email address": "hofmann@example.org",
        },
        {
            "investigation identifier": "GSE126024",
            "firstname": "MA",
            "lastname": "Schon",
            "email address": "schon@example.org",
        },
    ]

    rows = project_matrix_onto_entity_plan(seed, _embryo_study_plan())["investigation"]["rows"]

    assert rows[0]["email address"] == ""
    assert rows[1]["email address"] == "hofmann@example.org"
    assert rows[2]["email address"] == "schon@example.org"


def test_document_info_authors_are_contact_fallback_when_no_author_line_exists():
    contacts = EntityStructurePlannerAgent._source_contacts(
        {
            "document_content": "An abstract without an author header.",
            "document_info": {"authors": ["Hofmann F", "Schon MA", "Nodine MD"]},
        }
    )

    assert [(item["firstname"], item["lastname"]) for item in contacts] == [
        ("F", "Hofmann"),
        ("MA", "Schon"),
        ("MD", "Nodine"),
    ]


def test_local_excel_uses_five_package_qualified_sheets_and_raw_headers():
    seed = _seed_matrix()
    matrix = project_matrix_onto_entity_plan(seed, _embryo_study_plan())
    for level_name, block in matrix.items():
        block["fields"] = seed[level_name].get("fields", [])
    workbook = load_workbook(BytesIO(_generate_xlsx_local(matrix)), data_only=True)

    assert all("person" not in name.lower() for name in workbook.sheetnames)
    assert workbook.sheetnames[:5] == [
        "investigation - default",
        "study - default",
        "observationunit - default",
        "sample - Crop Plant sample enha",
        "assay - default",
    ]
    for worksheet in workbook.worksheets[:5]:
        headers = [
            worksheet.cell(1, column).value
            for column in range(1, worksheet.max_column + 1)
        ]
        assert all(not str(header).endswith(("(M)", "(R)", "(O)")) for header in headers)
    investigation = workbook["investigation - default"]
    assert [investigation.cell(row, 4).value for row in range(2, 5)] == ["MD", "F", "MA"]


def test_human_gold_workbooks_define_the_same_five_level_contract():
    root = Path(__file__).resolve().parents[1]
    gold_paths = [
        root / "evaluation/datasets/raw/biorem/BIOREM_Metadata.xlsx",
        root / "evaluation/datasets/raw/biosensor/Whole-cell_biosensor_metadata.xlsx",
        root / "evaluation/datasets/raw/earthworm/Diagonal_RNAseq_Earthworms.xlsx",
    ]
    investigation_rows = []
    for path in gold_paths:
        workbook = load_workbook(path, read_only=True, data_only=True)
        assert all("person" not in name.lower() for name in workbook.sheetnames)
        data_sheets = [name for name in workbook.sheetnames if name.lower() != "help"]
        assert len(data_sheets) >= 5
        for sheet_name in data_sheets:
            worksheet = workbook[sheet_name]
            headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
            assert all(not str(header).endswith(("(M)", "(R)", "(O)")) for header in headers)
        investigation = next(
            workbook[name]
            for name in workbook.sheetnames
            if name.lower().startswith("investigation - ")
        )
        investigation_rows.append(investigation.max_row - 1)

    assert investigation_rows == [7, 1, 16]


def test_structure_planner_catalog_excludes_field_selectors_but_keeps_value_fields():
    knowledge = [
        {
            "term": "design variable selector",
            "metadata": {
                "isa_sheet": "sample",
                "regex": "(dose|replicate|dev_stage)",
            },
        },
        {"term": "dose", "metadata": {"isa_sheet": "sample", "regex": ".*"}},
        {
            "term": "replicate",
            "metadata": {"isa_sheet": "sample", "regex": ".*"},
        },
        {
            "term": "dev_stage",
            "metadata": {"isa_sheet": "sample", "regex": ".*"},
        },
    ]

    catalog = EntityStructurePlannerAgent._field_catalog(knowledge)

    assert "design variable selector" not in catalog["sample"]
    assert {"dose", "replicate", "dev_stage"}.issubset(catalog["sample"])


def test_structure_planner_drops_dimension_mapping_that_violates_field_contract():
    plan = {
        "levels": [
            {
                "level": "assay",
                "entities": [
                    {
                        "source_group": "developmental_profile",
                        "attributes": [
                            {
                                "dimension_name": "developmental stage",
                                "field_name": "library source",
                                "value": "mature green",
                                "origin": "explicit",
                            }
                        ],
                    }
                ],
            }
        ],
        "unresolved_ambiguities": [],
    }
    knowledge = [
        {
            "term": "library source",
            "metadata": {
                "isa_sheet": "assay",
                "regex": "(GENOMIC|TRANSCRIPTOMIC|METAGENOMIC|SYNTHETIC)",
            },
        }
    ]

    result = EntityStructurePlannerAgent._enforce_dimension_field_contracts(
        plan, knowledge
    )

    attribute = result["levels"][0]["entities"][0]["attributes"][0]
    assert attribute["field_name"] is None
    assert attribute["value"] == "mature green"
    assert "do not satisfy" in result["unresolved_ambiguities"][0]


def test_contract_rejection_preserves_source_coverage_when_column_has_other_role():
    plan = {
        "levels": [
            {
                "level": "assay",
                "entities": [
                    {
                        "source_group": "profile",
                        "attributes": [
                            {
                                "dimension_name": "method",
                                "field_name": "library source",
                                "value": "unsupported method",
                                "origin": "explicit",
                            }
                        ],
                    }
                ],
            }
        ],
        "source_coverage": {
            "complete": True,
            "tables": [
                {
                    "coverage_complete": True,
                    "unmapped_columns": [],
                    "fairds_mappings": [
                        {
                            "level": "assay",
                            "column": "method",
                            "field_name": "library source",
                        }
                    ],
                    "column_roles": [
                        {
                            "column": "method",
                            "roles": ["assay_context", "fairds_mapping:assay"],
                        }
                    ],
                }
            ],
        },
        "unresolved_ambiguities": [],
    }
    knowledge = [
        {
            "term": "library source",
            "metadata": {
                "isa_sheet": "assay",
                "regex": "(GENOMIC|TRANSCRIPTOMIC)",
            },
        }
    ]

    result = EntityStructurePlannerAgent._enforce_dimension_field_contracts(
        plan, knowledge
    )

    coverage = result["source_coverage"]
    assert coverage["complete"] is True
    assert coverage["tables"][0]["unmapped_columns"] == []
    assert coverage["tables"][0]["column_roles"] == [
        {"column": "method", "roles": ["assay_context"]}
    ]


def test_contract_rejection_demotes_uncovered_source_column_to_extension():
    plan = {
        "levels": [
            {
                "level": "assay",
                "entities": [
                    {
                        "source_group": "profile",
                        "attributes": [
                            {
                                "dimension_name": "estimated.size",
                                "field_name": "insert size",
                                "value": "84.7",
                                "origin": "explicit",
                            }
                        ],
                    }
                ],
            }
        ],
        "source_coverage": {
            "complete": True,
            "tables": [
                {
                    "source_id": "source_001",
                    "source_path": "metadata.xlsx",
                    "table_name": "records",
                    "columns": ["estimated.size"],
                    "coverage_complete": True,
                    "unmapped_columns": [],
                    "fairds_mappings": [
                        {
                            "level": "assay",
                            "column": "estimated.size",
                            "field_name": "insert size",
                        }
                    ],
                    "column_roles": [
                        {
                            "column": "estimated.size",
                            "roles": ["fairds_mapping:assay"],
                        }
                    ],
                }
            ],
        },
        "source_extension_fields": [],
        "unresolved_ambiguities": [],
    }
    knowledge = [
        {
            "term": "insert size",
            "metadata": {"isa_sheet": "assay", "regex": r"[0-9]+ bp"},
        }
    ]

    result = EntityStructurePlannerAgent._enforce_dimension_field_contracts(
        plan, knowledge
    )

    attribute = result["levels"][0]["entities"][0]["attributes"][0]
    assert attribute["field_name"] == "estimated size"
    assert result["source_coverage"]["complete"] is True
    assert result["source_coverage"]["tables"][0]["unmapped_columns"] == []
    assert result["source_extension_fields"][0]["data_type"] == "number"
    assert result["source_extension_fields"][0]["demoted_from_fairds_mapping"] == (
        "insert size"
    )


def test_candidate_rank_prefers_authoritative_materialization_over_error_count():
    fallback = {"levels": [{"level": "sample", "entities": [{}]}]}
    authoritative = {
        "levels": [{"level": "sample", "entities": [{}]}],
        "source_coverage": {"tables": [{}], "complete": True},
    }

    fallback_rank = EntityStructurePlannerAgent._candidate_rank(
        fallback,
        ["one error"],
        used_authoritative_records=False,
    )
    authoritative_rank = EntityStructurePlannerAgent._candidate_rank(
        authoritative,
        ["one error", "second error"],
        used_authoritative_records=True,
    )

    assert authoritative_rank > fallback_rank


def test_record_column_audit_replaces_loose_mapping_with_typed_extension():
    claims = {
        "record_table_plans": [
            {
                "source_id": "source_001",
                "table_name": "records",
                "covers_complete_focal_study": True,
                "filters": [{"column": "project", "value": "study-1"}],
                "sample_identifier_column": "sample",
                "assay_identifier_column": "run",
                "observation_unit_columns": ["condition"],
                "column_mappings": [
                    {
                        "column": "adapter class",
                        "level": "assay",
                        "field_name": "sample preparation",
                    }
                ],
                "source_extension_mappings": [],
                "excluded_columns": [],
            }
        ]
    }
    audit = {
        "record_column_decisions": [
            {
                "source_id": "source_001",
                "table_name": "records",
                "column": column,
                "disposition": "structural_only",
            }
            for column in ("sample", "project", "run")
        ]
        + [
            {
                "source_id": "source_001",
                "table_name": "records",
                "column": "condition",
                "disposition": "source_extension",
                "level": "sample",
                "rationale": "Source-defined condition label.",
            },
            {
                "source_id": "source_001",
                "table_name": "records",
                "column": "adapter class",
                "disposition": "source_extension",
                "level": "assay",
                "data_type": "string",
                "rationale": "Adapter annotation, not a preparation protocol.",
            },
        ]
    }
    profiles = [
        {
            "source_id": "source_001",
            "table_name": "records",
            "columns": ["sample", "project", "run", "condition", "adapter class"],
        }
    ]

    applied, errors = EntityStructurePlannerAgent._apply_record_column_audit(
        claims,
        audit,
        {"assay": ["sample preparation"]},
        profiles,
    )

    assert errors == []
    record_plan = claims["record_table_plans"][0]
    assert record_plan["column_mappings"] == []
    assert record_plan["source_extension_mappings"] == [
        {
            "column": "condition",
            "level": "sample",
            "data_type": "string",
            "definition": "Source-defined condition label.",
        },
        {
            "column": "adapter class",
            "level": "assay",
            "data_type": "string",
            "definition": "Adapter annotation, not a preparation protocol.",
        },
    ]
    assert "source_001:adapter class->assay.source_extension" in applied


def test_scope_audit_accepts_semantic_decision_alias_and_advisory_extras():
    parsed = EntityPlanScopeAuditResponse.model_validate(
        {
            "missing_conditions": [
                {
                    "group_id": "group_001",
                    "dimension_name": "condition",
                    "source_value": "treated",
                    "level": "sample",
                    "field_name": "treatment",
                    "rationale": "Provider explanation",
                    "suggested_level": "sample",
                }
            ],
            "record_column_decisions": [
                {
                    "column": "adapter class",
                    "decision": "source_extension",
                    "level": "assay",
                }
            ],
        }
    ).model_dump()

    assert parsed["record_column_decisions"][0]["disposition"] == "source_extension"
    assert parsed["record_column_decisions"][0]["source_id"] == ""


def test_record_column_audit_infers_omitted_table_context_only_when_unique():
    claims = {
        "record_table_plans": [
            {
                "source_id": "source_001",
                "table_name": "records",
                "covers_complete_focal_study": True,
                "filters": [],
                "sample_identifier_column": "sample",
                "assay_identifier_column": "run",
                "group_column": "branch",
                "column_mappings": [],
                "source_extension_mappings": [],
                "excluded_columns": [],
            }
        ]
    }
    audit = {
        "record_column_decisions": [
            {"column": "sample", "disposition": "structural_only"},
            {"column": "run", "disposition": "structural_only"},
            {"column": "branch", "disposition": "structural_only"},
            {
                "column": "adapter class",
                "disposition": "source_extension",
                "level": "assay",
            },
        ]
    }
    profiles = [
        {
            "source_id": "source_001",
            "table_name": "records",
            "columns": ["sample", "run", "branch", "adapter class"],
        }
    ]

    _, errors = EntityStructurePlannerAgent._apply_record_column_audit(
        claims, audit, {}, profiles
    )

    assert errors == []
    assert claims["record_table_plans"][0]["source_extension_mappings"][0][
        "column"
    ] == "adapter class"


def test_record_column_audit_retains_populated_focal_metadata_when_excluded():
    claims = {
        "record_table_plans": [
            {
                "source_id": "source_001",
                "table_name": "records",
                "covers_complete_focal_study": True,
                "filters": [],
                "sample_identifier_column": "sample",
                "assay_identifier_column": "run",
                "observation_unit_columns": [],
                "column_mappings": [],
                "source_extension_mappings": [],
                "excluded_columns": [],
            }
        ]
    }
    audit = {
        "record_column_decisions": [
            {
                "source_id": "source_001",
                "table_name": "records",
                "column": "sample",
                "disposition": "structural_only",
            },
            {
                "source_id": "source_001",
                "table_name": "records",
                "column": "run",
                "disposition": "structural_only",
            },
            {
                "source_id": "source_001",
                "table_name": "records",
                "column": "source category",
                "disposition": "excluded",
                "rationale": "Not used in the primary analysis.",
            },
        ]
    }
    profiles = [
        {
            "source_id": "source_001",
            "table_name": "records",
            "columns": ["sample", "run", "source category"],
            "column_distributions": [
                {"column": "source category", "nonblank": 12}
            ],
        }
    ]

    applied, errors = EntityStructurePlannerAgent._apply_record_column_audit(
        claims, audit, {}, profiles
    )

    assert errors == []
    assert claims["record_table_plans"][0]["excluded_columns"] == []
    assert claims["record_table_plans"][0]["source_extension_mappings"] == [
        {
            "column": "source category",
            "level": "sample",
            "data_type": "string",
            "definition": (
                "Populated source-record metadata retained instead of the proposed "
                "exclusion: Not used in the primary analysis."
            ),
        }
    ]
    assert applied[-1] == "source_001:source category->sample.source_extension"


def test_structure_planner_promotes_mapped_field_from_selected_package_contract():
    selected = {
        "term": "sample name",
        "metadata": {"isa_sheet": "sample", "package": "default"},
    }
    dose = {
        "term": "dose",
        "metadata": {
            "isa_sheet": "sample",
            "package": "Lean Plant Contract",
            "regex": ".*",
        },
    }
    state = {
        "retrieved_knowledge": [selected],
        "api_capabilities": {"selected_package_field_contracts": [dose]},
    }
    plan = {
        "levels": [
            {
                "level": "sample",
                "entities": [
                    {
                        "attributes": [
                            {
                                "dimension_name": "input amount",
                                "field_name": "dose",
                                "value": "0.1 ng",
                                "origin": "explicit",
                            }
                        ]
                    }
                ],
            }
        ]
    }

    planning = EntityStructurePlannerAgent._planning_knowledge(state)
    promoted = EntityStructurePlannerAgent._promote_planned_fields(
        state, plan, planning
    )

    assert "dose" in EntityStructurePlannerAgent._field_catalog(planning)["sample"]
    assert promoted == ["sample.dose"]
    assert [item["term"] for item in state["retrieved_knowledge"]] == [
        "sample name",
        "dose",
    ]


def test_mapping_audit_replaces_generic_description_with_specific_contract_field():
    claims = {
        "design_groups": [
            {
                "group_id": "group_001",
                "dimensions": [
                    {
                        "name": "input_amount",
                        "explicit_values": ["5 ng", "1 ng"],
                        "known_values": [],
                        "applies_to": ["sample", "assay"],
                        "field_mappings": [
                            {"level": "sample", "field_name": "sample description"}
                        ],
                    }
                ],
            }
        ]
    }
    audit = {
        "mapping_corrections": [
            {
                "group_id": "group_001",
                "dimension_name": "input_amount",
                "level": "sample",
                "field_name": "dose",
            }
        ]
    }
    knowledge = [
        {"term": "dose", "metadata": {"isa_sheet": "sample", "regex": ".*"}}
    ]

    applied = EntityStructurePlannerAgent._apply_mapping_audit_findings(
        claims,
        audit,
        {"sample": ["sample description", "dose"]},
        knowledge,
    )

    assert applied == ["group_001:input_amount->sample.dose"]
    assert claims["design_groups"][0]["dimensions"][0]["field_mappings"] == [
        {"level": "sample", "field_name": "dose"}
    ]


def test_complete_mapping_audit_requires_one_decision_per_dimension():
    claims = {
        "design_groups": [
            {
                "group_id": "group_001",
                "dimensions": [
                    {
                        "name": "input_amount",
                        "explicit_values": ["5 ng", "1 ng"],
                        "known_values": [],
                        "applies_to": ["sample", "assay"],
                        "field_mappings": [],
                    }
                ],
            }
        ]
    }

    applied, errors = EntityStructurePlannerAgent._apply_complete_mapping_audit(
        claims, {"mapping_decisions": []}, {"assay": ["sample preparation"]}, []
    )

    assert applied == []
    assert errors == [
        "Independent mapping audit omitted or duplicated group_001:input_amount."
    ]
    assert claims["design_groups"][0]["dimensions"][0]["field_mappings"] == []


def test_complete_mapping_audit_replaces_draft_mapping_with_validated_decision():
    claims = {
        "design_groups": [
            {
                "group_id": "group_001",
                "dimensions": [
                    {
                        "name": "input_amount",
                        "explicit_values": ["5 ng", "1 ng"],
                        "known_values": [],
                        "applies_to": ["sample", "assay"],
                        "field_mappings": [
                            {"level": "sample", "field_name": "dose"}
                        ],
                    }
                ],
            }
        ]
    }
    audit = {
        "mapping_decisions": [
            {
                "group_id": "group_001",
                "dimension_name": "input_amount",
                "mappings": [
                    {
                        "level": "assay",
                        "field_name": "sample preparation",
                    }
                ],
            }
        ]
    }
    knowledge = [
        {
            "term": "sample preparation",
            "metadata": {"isa_sheet": "assay", "regex": ".*"},
        }
    ]

    applied, errors = EntityStructurePlannerAgent._apply_complete_mapping_audit(
        claims,
        audit,
        {"sample": ["dose"], "assay": ["sample preparation"]},
        knowledge,
    )

    assert errors == []
    assert applied == [
        "group_001:input_amount->assay.sample preparation"
    ]
    assert claims["design_groups"][0]["dimensions"][0]["field_mappings"] == [
        {"level": "assay", "field_name": "sample preparation"}
    ]


def test_complete_mapping_audit_rejects_bad_target_but_keeps_cardinality_claim():
    claims = {
        "design_groups": [
            {
                "group_id": "group_001",
                "dimensions": [
                    {
                        "name": "input_amount",
                        "level_count": 2,
                        "explicit_values": ["5 ng", "1 ng"],
                        "known_values": [],
                        "applies_to": ["sample", "assay"],
                        "field_mappings": [
                            {"level": "sample", "field_name": "dose"}
                        ],
                    }
                ],
            }
        ]
    }
    audit = {
        "mapping_decisions": [
            {
                "group_id": "group_001",
                "dimension_name": "input_amount",
                "mappings": [
                    {"level": "assay", "field_name": "not in catalog"}
                ],
            }
        ]
    }

    applied, warnings = EntityStructurePlannerAgent._apply_complete_mapping_audit(
        claims, audit, {"sample": ["dose"], "assay": ["sample preparation"]}, []
    )

    assert applied == ["group_001:input_amount->unmapped"]
    assert warnings and "rejected mapping target" in warnings[0]
    dimension = claims["design_groups"][0]["dimensions"][0]
    assert dimension["level_count"] == 2
    assert dimension["field_mappings"] == []


@pytest.mark.anyio
async def test_orchestrator_stops_before_json_when_entity_plan_is_invalid(monkeypatch):
    calls = []

    class Planner:
        async def execute(self, state):
            state["entity_plan_validation"] = {
                "passed": False,
                "errors": ["sample parent is unresolved"],
            }
            return state

    class JSONGenerator:
        async def execute(self, state):
            calls.append("json")
            return state

    node = OrchestrateNode(
        document_parser=object(),
        knowledge_retriever=object(),
        entity_structure_planner=Planner(),
        json_generator=JSONGenerator(),
        max_step_retries=0,
        max_global_retries=0,
    )

    async def pass_through(state, agent, agent_name, check_output_fn):
        calls.append(agent_name)
        if agent_name == "DocumentParser":
            state["document_info"] = {"title": "source"}
        elif agent_name == "KnowledgeRetriever":
            state["retrieved_knowledge"] = [{"term": "study title"}]
        elif agent_name == "JSONGenerator":
            return await agent.execute(state)
        return state

    async def plan_workflow(state):
        return state

    monkeypatch.setattr(node, "_execute_agent_with_retry", pass_through)
    monkeypatch.setattr(node, "_plan_workflow_internal", plan_workflow)

    result = await node({"context": {}, "execution_history": [], "errors": []})

    assert calls == ["DocumentParser", "KnowledgeRetriever"]
    assert result["entity_plan_hard_gate_failed"] is True
    assert result["needs_human_review"] is True


def test_mapper_clears_broadcast_values_outside_plan_mapped_entities():
    matrix = {
        "sample": {
            "columns": ["sample identifier", "dose"],
            "rows": [
                {"sample identifier": "s1", "dose": "5 ng"},
                {"sample identifier": "s2", "dose": "0.1-5 ng"},
            ],
        }
    }
    plan = {
        "levels": [
            {
                "level": "sample",
                "entities": [
                    {
                        "row_id": "s1",
                        "attributes": [
                            {"field_name": "dose", "value": "5 ng"}
                        ],
                    },
                    {"row_id": "s2", "attributes": []},
                ],
            }
        ]
    }

    cleaned, counts = ISAValueMapperAgent._clear_values_outside_plan_field_scope(
        matrix, plan
    )

    assert cleaned["sample"]["rows"][0]["dose"] == "5 ng"
    assert cleaned["sample"]["rows"][1]["dose"] == ""
    assert counts == {"sample.dose": 1}


def test_mapper_prunes_sheet_wide_repeated_optional_narrative():
    matrix = {
        "assay": {
            "columns": ["assay identifier", "notes"],
            "rows": [
                {"assay identifier": "a1", "notes": "whole-study summary"},
                {"assay identifier": "a2", "notes": "whole-study summary"},
            ],
        }
    }
    fields = {
        "assay": [
            {
                "field_name": "notes",
                "requirement": "OPTIONAL",
                "status": "confirmed",
                "evidence": "abstract",
            }
        ]
    }

    cleaned, pruned = ISAValueMapperAgent._prune_uninformative_optional_columns(
        matrix, fields, {"levels": []}
    )

    assert cleaned["assay"]["columns"] == ["assay identifier"]
    assert pruned == {"assay": ["notes"]}


def test_mapper_keeps_distinct_field_when_only_value_substring_overlaps_plan():
    matrix = {
        "sample": {
            "columns": ["sample identifier", "specific field", "broad field"],
            "rows": [
                {
                    "sample identifier": "s1",
                    "specific field": "Line-A (reference)",
                    "broad field": "Line-A",
                },
                {
                    "sample identifier": "s2",
                    "specific field": "",
                    "broad field": "",
                },
            ],
        }
    }
    fields = {
        "sample": [
            {
                "field_name": "specific field",
                "requirement": "OPTIONAL",
                "status": "confirmed",
                "evidence": "source",
            },
            {
                "field_name": "broad field",
                "requirement": "OPTIONAL",
                "status": "confirmed",
                "evidence": "source",
            },
        ]
    }
    plan = {
        "design_claims": {
            "design_groups": [
                {
                    "group_id": "g1",
                    "dimensions": [
                        {"name": "line", "applies_to": ["sample", "assay"]}
                    ],
                }
            ]
        },
        "levels": [
            {
                "level": "sample",
                "entities": [
                    {
                        "source_group": "g1",
                        "attributes": [
                            {
                                "dimension_name": "line",
                                "field_name": "specific field",
                                "value": "Line-A (reference)",
                                "origin": "explicit",
                            }
                        ],
                    },
                    {"source_group": "g2", "attributes": []},
                ],
            }
        ],
    }

    cleaned, pruned = ISAValueMapperAgent._prune_uninformative_optional_columns(
        matrix, fields, plan, source_text="Line-A (reference)"
    )

    assert "specific field" in cleaned["sample"]["columns"]
    assert "broad field" in cleaned["sample"]["columns"]
    assert pruned == {}


def test_mapper_prunes_optional_cells_not_supported_by_confirmed_field_value():
    matrix = {
        "assay": {
            "columns": ["assay identifier", "optional label"],
            "rows": [
                {"assay identifier": "a1", "optional label": "derived label 01"},
                {"assay identifier": "a2", "optional label": "derived label 02"},
            ],
        }
    }
    fields = {
        "assay": [
            {
                "field_name": "optional label",
                "value": "source assay type",
                "requirement": "OPTIONAL",
                "status": "confirmed",
                "evidence": "source assay type",
            }
        ]
    }

    cleaned, pruned = ISAValueMapperAgent._prune_uninformative_optional_columns(
        matrix, fields, {"levels": []}, source_text="source assay type"
    )

    assert cleaned["assay"]["columns"] == ["assay identifier"]
    assert pruned == {"assay": ["optional label"]}


def test_mapper_keeps_optional_values_verified_against_source_text():
    matrix = {
        "sample": {
            "columns": ["sample identifier", "specific factor"],
            "rows": [
                {"sample identifier": "s1", "specific factor": "condition alpha"},
                {"sample identifier": "s2", "specific factor": "condition alpha"},
            ],
        }
    }
    fields = {
        "sample": [
            {
                "field_name": "specific factor",
                "requirement": "OPTIONAL",
                "status": "provisional",
                "evidence": "",
            }
        ]
    }

    cleaned, pruned = ISAValueMapperAgent._prune_uninformative_optional_columns(
        matrix,
        fields,
        {"levels": []},
        source_text="Samples were collected under condition alpha.",
    )

    assert cleaned["sample"]["columns"] == [
        "sample identifier",
        "specific factor",
    ]
    assert pruned == {}


def test_mapper_keeps_optional_values_verified_against_explicit_plan_values():
    matrix = {
        "assay": {
            "columns": ["assay identifier", "preparation amount"],
            "rows": [
                {"assay identifier": "a1", "preparation amount": "5 ng"},
                {"assay identifier": "a2", "preparation amount": "1 ng"},
            ],
        }
    }
    fields = {
        "assay": [
            {
                "field_name": "preparation amount",
                "requirement": "OPTIONAL",
                "status": "provisional",
                "evidence": "",
            }
        ]
    }
    plan = {
        "design_claims": {
            "design_groups": [
                {
                    "group_id": "g1",
                    "dimensions": [
                        {
                            "name": "amount",
                            "applies_to": ["sample", "assay"],
                        }
                    ],
                }
            ]
        },
        "levels": [
            {
                "level": "assay",
                "entities": [
                    {
                        "source_group": "g1",
                        "attributes": [
                            {
                                "dimension_name": "amount",
                                "field_name": None,
                                "value": "5 ng",
                                "origin": "explicit",
                            }
                        ],
                    },
                    {
                        "source_group": "g1",
                        "attributes": [
                            {
                                "dimension_name": "amount",
                                "field_name": None,
                                "value": "1 ng",
                                "origin": "explicit",
                            }
                        ],
                    },
                ],
            }
        ],
    }

    cleaned, pruned = ISAValueMapperAgent._prune_uninformative_optional_columns(
        matrix, fields, plan
    )

    assert cleaned["assay"]["columns"] == [
        "assay identifier",
        "preparation amount",
    ]
    assert pruned == {}


def test_mapper_preserves_placeholder_shaped_value_owned_by_entity_plan():
    matrix = {
        "assay": {
            "columns": ["assay identifier", "adapter trim"],
            "rows": [
                {"assay identifier": "a1", "adapter trim": "none"},
                {"assay identifier": "a2", "adapter trim": "nextera"},
            ],
        }
    }
    fields = {
        "assay": [
            {
                "field_name": "adapter trim",
                "requirement": "OPTIONAL",
                "status": "confirmed",
                "evidence": "authoritative source table",
            }
        ]
    }
    plan = {
        "levels": [
            {
                "level": "assay",
                "entities": [
                    {
                        "attributes": [
                            {
                                "field_name": "adapter trim",
                                "value": "none",
                                "origin": "explicit",
                            }
                        ]
                    },
                    {
                        "attributes": [
                            {
                                "field_name": "adapter trim",
                                "value": "nextera",
                                "origin": "explicit",
                            }
                        ]
                    },
                ],
            }
        ]
    }

    cleaned, pruned = ISAValueMapperAgent._prune_uninformative_optional_columns(
        matrix, fields, plan
    )

    assert cleaned["assay"]["columns"] == ["assay identifier", "adapter trim"]
    assert [row["adapter trim"] for row in cleaned["assay"]["rows"]] == [
        "none",
        "nextera",
    ]
    assert pruned == {}


def test_mapper_restores_confirmed_level_value_despite_synthetic_entity_id():
    matrix = {
        "assay": {
            "columns": ["assay identifier", "library strategy"],
            "rows": [
                {"assay identifier": "a1", "library strategy": ""},
                {"assay identifier": "a2", "library strategy": ""},
            ],
        }
    }
    fields = {
        "assay": [
            {
                "field_name": "library strategy",
                "value": "RNA-Seq",
                "status": "confirmed",
                "confidence": 0.95,
                "evidence": "source_001: mRNA-seq",
                "value_scope": "level",
                "entity_id": "assay_001",
            }
        ]
    }

    restored = ISAValueMapperAgent._restore_high_confidence_shared_values(
        matrix, fields, entity_plan={"levels": []}
    )

    assert [row["library strategy"] for row in restored["assay"]["rows"]] == [
        "RNA-Seq",
        "RNA-Seq",
    ]


def test_reuse_audit_fails_closed_when_shared_pool_does_not_prove_reuse():
    claims = {
        "design_groups": [
            {
                "group_id": "group_001",
                "evidence": "RNA was diluted and prepared with three methods.",
                "reuse_same_sample_for_multiple_assays": True,
                "sample_reuse_evidence": "prepared with three methods",
            }
        ]
    }
    audit = {
        "reuse_validations": [
            {
                "group_id": "group_001",
                "approved": False,
                "evidence_quote": None,
                "rationale": "Distinct preparation paths are not explicit reuse.",
            }
        ]
    }

    rejected = EntityStructurePlannerAgent._enforce_independent_reuse_audit(
        claims, audit
    )

    assert rejected == ["group_001"]
    assert claims["design_groups"][0]["reuse_same_sample_for_multiple_assays"] is False
    assert claims["design_groups"][0]["sample_reuse_evidence"] is None


def test_reuse_audit_accepts_only_verbatim_independent_proof():
    claims = {
        "design_groups": [
            {
                "group_id": "group_001",
                "evidence": "The same identified sample was measured in two modalities.",
                "reuse_same_sample_for_multiple_assays": True,
                "sample_reuse_evidence": None,
            }
        ]
    }
    audit = {
        "reuse_validations": [
            {
                "group_id": "group_001",
                "approved": True,
                "evidence_quote": "same identified sample",
                "rationale": "Explicitly the same material.",
            }
        ]
    }

    rejected = EntityStructurePlannerAgent._enforce_independent_reuse_audit(
        claims, audit
    )

    assert rejected == []
    assert claims["design_groups"][0]["sample_reuse_evidence"] == "same identified sample"


def test_dimension_scope_audit_requires_consensus_before_changing_cardinality():
    claims = {
        "design_groups": [
            {
                "group_id": "group_001",
                "dimensions": [
                    {
                        "name": "input_amount",
                        "applies_to": ["observationunit", "sample", "assay"],
                    }
                ],
            }
        ]
    }
    audit = {
        "dimension_scopes": [
            {
                "group_id": "group_001",
                "dimension_name": "input_amount",
                "applies_to": ["sample", "assay"],
            }
        ]
    }

    applied, errors = EntityStructurePlannerAgent._apply_dimension_scope_audit(
        claims, audit
    )

    assert applied == []
    assert errors == [
        "Independent scope audit disagreed with the draft scope for "
        "group_001:input_amount: draft=['observationunit', 'sample', 'assay'], "
        "audit=['sample', 'assay']."
    ]
    assert claims["design_groups"][0]["dimensions"][0]["applies_to"] == [
        "observationunit",
        "sample",
        "assay",
    ]


def test_dimension_scope_audit_can_repair_disagreement_with_authoritative_records():
    claims = {
        "record_table_plans": [{"covers_complete_focal_study": True}],
        "design_groups": [
            {
                "group_id": "group_001",
                "reuse_same_sample_for_multiple_assays": False,
                "dimensions": [
                    {
                        "name": "input_amount",
                        "level_count": 5,
                        "semantic_role": "observation_condition",
                        "applies_to": ["observationunit", "sample", "assay"],
                    }
                ],
            }
        ],
    }
    audit = {
        "dimension_scopes": [
            {
                "group_id": "group_001",
                "dimension_name": "input_amount",
                "semantic_role": "material_path",
                "applies_to": ["sample", "assay"],
            }
        ]
    }

    applied, errors = EntityStructurePlannerAgent._apply_dimension_scope_audit(
        claims, audit
    )

    assert errors == []
    assert applied == [
        "group_001:input_amount.role=material_path",
        "group_001:input_amount=sample/assay",
    ]
    dimension = claims["design_groups"][0]["dimensions"][0]
    assert dimension["semantic_role"] == "material_path"
    assert dimension["applies_to"] == ["sample", "assay"]


def test_dimension_scope_audit_can_relocate_constant_without_changing_cardinality():
    claims = {
        "design_groups": [
            {
                "group_id": "group_001",
                "dimensions": [
                    {
                        "name": "material_state",
                        "level_count": 1,
                        "applies_to": ["sample", "assay"],
                    }
                ],
            }
        ]
    }
    audit = {
        "dimension_scopes": [
            {
                "group_id": "group_001",
                "dimension_name": "material_state",
                "applies_to": ["observationunit", "sample", "assay"],
            }
        ]
    }

    applied, errors = EntityStructurePlannerAgent._apply_dimension_scope_audit(
        claims, audit
    )

    assert errors == []
    assert applied == [
        "group_001:material_state=observationunit/sample/assay"
    ]
    assert claims["design_groups"][0]["dimensions"][0]["applies_to"] == [
        "observationunit",
        "sample",
        "assay",
    ]


def test_semantic_role_consensus_repairs_varying_condition_scope():
    claims = {
        "design_groups": [
            {
                "group_id": "group_001",
                "reuse_same_sample_for_multiple_assays": False,
                "dimensions": [
                    {
                        "name": "condition_axis",
                        "level_count": 8,
                        "semantic_role": "observation_condition",
                        "applies_to": ["sample", "assay"],
                    }
                ],
            }
        ]
    }
    audit = {
        "dimension_scopes": [
            {
                "group_id": "group_001",
                "dimension_name": "condition_axis",
                "semantic_role": "observation_condition",
                "applies_to": ["observationunit", "sample", "assay"],
            }
        ]
    }

    applied, errors = EntityStructurePlannerAgent._apply_dimension_scope_audit(
        claims, audit
    )

    assert errors == []
    assert applied == [
        "group_001:condition_axis=observationunit/sample/assay"
    ]
    assert claims["design_groups"][0]["dimensions"][0]["applies_to"] == [
        "observationunit",
        "sample",
        "assay",
    ]


def test_semantic_role_disagreement_blocks_scope_arbitration():
    claims = {
        "design_groups": [
            {
                "group_id": "group_001",
                "dimensions": [
                    {
                        "name": "factor_axis",
                        "level_count": 3,
                        "semantic_role": "observation_condition",
                        "applies_to": ["observationunit", "sample", "assay"],
                    }
                ],
            }
        ]
    }
    audit = {
        "dimension_scopes": [
            {
                "group_id": "group_001",
                "dimension_name": "factor_axis",
                "semantic_role": "material_path",
                "applies_to": ["sample", "assay"],
            }
        ]
    }

    applied, errors = EntityStructurePlannerAgent._apply_dimension_scope_audit(
        claims, audit
    )

    assert applied == []
    assert errors == [
        "Independent semantic-role audit disagreed with the draft for "
        "group_001:factor_axis: draft='observation_condition', "
        "audit='material_path'."
    ]


def test_scope_audit_schema_accepts_provider_dimension_aliases_and_redundant_context():
    from fairifier.agents.response_models import EntityPlanScopeAuditResponse

    parsed = EntityPlanScopeAuditResponse.model_validate(
        {
            "mapping_decisions": [
                {
                    "group_id": "group_001",
                    "dimension": "developmental_stage",
                    "mappings": [
                        {"level": "sample", "field_name": "dev_stage"}
                    ],
                    "applies_to": ["sample", "assay"],
                }
            ],
            "dimension_scopes": [
                {
                    "group_id": "group_001",
                    "dimension": "developmental_stage",
                    "applies_to": ["sample", "assay"],
                }
            ],
            "reuse_validations": [
                {
                    "group_id": "group_001",
                    "approved": False,
                    "evidence_quote": None,
                    "reasoning": "No explicit reuse evidence.",
                }
            ],
        }
    )

    payload = parsed.model_dump()
    assert payload["mapping_decisions"][0]["dimension_name"] == "developmental_stage"
    assert payload["dimension_scopes"][0]["dimension_name"] == "developmental_stage"
    assert payload["reuse_validations"][0]["rationale"] == "No explicit reuse evidence."


def test_scope_audit_keeps_core_decisions_when_optional_finding_is_incomplete():
    from fairifier.agents.response_models import EntityPlanScopeAuditResponse

    parsed = EntityPlanScopeAuditResponse.model_validate(
        {
            "missing_conditions": [
                {
                    "group_id": "group_001",
                    "source_value": "source-stated condition",
                    "level": "sample",
                    "field_name": "organism part",
                    "evidence": "Source sentence supporting the condition.",
                }
            ],
            "mapping_decisions": [
                {
                    "group_id": "group_001",
                    "dimension": "material_state",
                    "mappings": [],
                }
            ],
            "dimension_scopes": [
                {
                    "group_id": "group_001",
                    "dimension": "material_state",
                    "applies_to": ["sample", "assay"],
                }
            ],
        }
    ).model_dump()

    assert parsed["missing_conditions"][0]["dimension_name"] == ""
    assert parsed["missing_conditions"][0]["evidence"] == (
        "Source sentence supporting the condition."
    )
    assert parsed["mapping_decisions"][0]["dimension_name"] == "material_state"
    assert parsed["dimension_scopes"][0]["dimension_name"] == "material_state"


def test_design_claim_response_allows_unknown_contextual_dimension_count():
    from fairifier.agents.response_models import EntityDesignClaimsResponse

    parsed = EntityDesignClaimsResponse.model_validate(
        {
            "investigations": [],
            "studies": [],
            "design_groups": [
                {
                    "group_id": "context_group",
                    "study_row_id": "study_001",
                    "evidence": "Replicate count varies across reused datasets.",
                    "dimensions": [
                        {
                            "name": "replicate",
                            "level_count": None,
                            "origin": "derived",
                            "semantic_role": "sample_replicate",
                            "applies_to": ["sample"],
                        }
                    ],
                }
            ],
            "record_table_plans": [],
            "design_summary": [],
            "unresolved_ambiguities": [],
            "confidence": 0.5,
        }
    )

    assert parsed.design_groups[0].dimensions[0].level_count is None


def test_rejected_reuse_propagates_varying_assay_branch_through_sample():
    claims = {
        "design_groups": [
            {
                "group_id": "group_001",
                "evidence": "Three preparations were made from a shared pool.",
                "reuse_same_sample_for_multiple_assays": True,
                "sample_reuse_evidence": "shared pool",
                "dimensions": [
                    {
                        "name": "preparation_method",
                        "level_count": 3,
                        "applies_to": ["assay"],
                    }
                ],
            }
        ]
    }
    audit = {
        "reuse_validations": [
            {
                "group_id": "group_001",
                "approved": False,
                "evidence_quote": None,
            }
        ]
    }

    rejected = EntityStructurePlannerAgent._enforce_independent_reuse_audit(
        claims, audit
    )

    assert rejected == ["group_001"]
    group = claims["design_groups"][0]
    assert group["reuse_same_sample_for_multiple_assays"] is False
    assert group["dimensions"][0]["applies_to"] == ["sample", "assay"]


def test_scope_consensus_normalizes_varying_assay_only_factor_without_reuse():
    claims = {
        "design_groups": [
            {
                "group_id": "group_001",
                "reuse_same_sample_for_multiple_assays": False,
                "dimensions": [
                    {
                        "name": "preparation_method",
                        "level_count": 3,
                        "applies_to": ["sample", "assay"],
                    }
                ],
            }
        ]
    }
    audit = {
        "dimension_scopes": [
            {
                "group_id": "group_001",
                "dimension_name": "preparation_method",
                "applies_to": ["assay"],
            }
        ]
    }

    applied, errors = EntityStructurePlannerAgent._apply_dimension_scope_audit(
        claims, audit
    )

    assert applied == []
    assert errors == []


def test_dimension_scope_audit_fails_when_any_dimension_is_unreviewed():
    claims = {
        "design_groups": [
            {
                "group_id": "group_001",
                "dimensions": [
                    {"name": "condition", "applies_to": ["sample", "assay"]}
                ],
            }
        ]
    }

    applied, errors = EntityStructurePlannerAgent._apply_dimension_scope_audit(
        claims, {"dimension_scopes": []}
    )

    assert applied == []
    assert errors == [
        "Independent scope audit omitted or duplicated group_001:condition."
    ]
