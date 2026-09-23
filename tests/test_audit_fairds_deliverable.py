from __future__ import annotations

import json
import hashlib
from pathlib import Path

from openpyxl import Workbook

from evaluation.scripts.audit_fairds_deliverable import audit_run
from evaluation.scripts.evaluate_fairds_expected_outcome import (
    evaluate_expected_outcome,
)


LEVEL_ROWS = {
    "investigation": [
        {
            "investigation identifier": "inv-1",
            "investigation title": "Study",
            "firstname": "Ada",
            "lastname": "Lovelace",
        }
    ],
    "study": [
        {
            "study identifier": "study-1",
            "study title": "Study",
            "investigation identifier": "inv-1",
        }
    ],
    "observationunit": [
        {
            "observation unit identifier": "ou-1",
            "observation unit description": "experimental unit",
            "study identifier": "study-1",
        }
    ],
    "sample": [
        {
            "sample identifier": "sample-1",
            "sample description": "biological material",
            "observation unit identifier": "ou-1",
        }
    ],
    "assay": [
        {
            "assay identifier": "assay-1",
            "assay description": "measurement",
            "sample identifier": "sample-1",
        }
    ],
}


def _write_run(tmp_path: Path, *, add_person: bool = False) -> Path:
    run_dir = tmp_path / "run"
    deliverables = run_dir / "deliverables"
    reports = run_dir / "reports"
    deliverables.mkdir(parents=True)
    reports.mkdir()

    matrix = {}
    isa_structure = {}
    for level, rows in LEVEL_ROWS.items():
        columns = list(rows[0])
        matrix[level] = {"columns": columns, "rows": rows}
        isa_structure[level] = {
            "columns": columns,
            "rows": rows,
            "fields": [
                {
                    "field_name": column,
                    "required": True,
                }
                for column in columns
                if column
                not in {
                    "investigation identifier"
                    if level == "study"
                    else "",
                    "study identifier"
                    if level == "observationunit"
                    else "",
                    "observation unit identifier"
                    if level == "sample"
                    else "",
                    "sample identifier" if level == "assay" else "",
                }
            ],
        }

    metadata = {
        "packages_used": ["default", "focused-assay"],
        "package_selection_trace": {
            "stable": True,
            "coverage_resolved": True,
            "uncovered_schema_levels": [],
        },
        "needs_review": False,
        "overall_confidence": 0.9,
        "isa_structure": isa_structure,
        "isa_values": matrix,
        "_field_definitions": [
            {
                "term": column,
                "isa_sheet": level,
                "package": "default",
                "requirement": "MANDATORY",
                "required": True,
            }
            for level, sheet in matrix.items()
            for column in sheet["columns"]
        ],
    }
    (deliverables / "metadata.json").write_text(
        json.dumps(metadata), encoding="utf-8"
    )
    (deliverables / "isa_values.json").write_text(
        json.dumps(matrix), encoding="utf-8"
    )
    entity_plan = {
        "validation": {
            "passed": True,
            "errors": [],
            "row_counts": {level: len(rows) for level, rows in LEVEL_ROWS.items()},
            "investigation_contact_count": 1,
        }
    }
    (reports / "entity_plan.json").write_text(
        json.dumps(entity_plan), encoding="utf-8"
    )
    (run_dir / "workflow_report.json").write_text(
        json.dumps({"workflow_status": "completed"}), encoding="utf-8"
    )

    workbook = Workbook()
    workbook.remove(workbook.active)
    for level, sheet in matrix.items():
        worksheet = workbook.create_sheet(level)
        worksheet.append(sheet["columns"])
        for row in sheet["rows"]:
            worksheet.append([row.get(column, "") for column in sheet["columns"]])
    workbook.create_sheet("Help").append(["FAIR-DS terms"])
    if add_person:
        workbook.create_sheet("Person").append(["firstname", "lastname"])
    workbook.save(deliverables / "metadata_fairds.xlsx")
    return run_dir


def test_audit_accepts_consistent_fairds_workbook(tmp_path: Path):
    run_dir = _write_run(tmp_path)

    report = audit_run(run_dir)

    assert report["passed"] is True
    assert report["error_count"] == 0
    assert report["metrics"]["row_counts"]["sample"] == 1


def test_audit_rejects_formula_cells(tmp_path: Path):
    run_dir = _write_run(tmp_path)
    workbook_path = run_dir / "deliverables" / "metadata_fairds.xlsx"
    workbook = load_workbook_for_test(workbook_path)
    workbook["sample"].cell(row=2, column=2, value='="biological material"')
    workbook.save(workbook_path)

    report = audit_run(run_dir)

    assert "formula_cells_present" in {
        issue["code"] for issue in report["issues"]
    }


def test_audit_rejects_person_sheet_and_meaningless_repeated_entities(tmp_path: Path):
    run_dir = _write_run(tmp_path, add_person=True)
    deliverables = run_dir / "deliverables"
    matrix = json.loads((deliverables / "isa_values.json").read_text())
    matrix["assay"]["rows"].append(
        {
            "assay identifier": "assay-2",
            "assay description": "measurement",
            "sample identifier": "sample-1",
        }
    )
    (deliverables / "isa_values.json").write_text(json.dumps(matrix))
    metadata = json.loads((deliverables / "metadata.json").read_text())
    metadata["isa_values"] = matrix
    metadata["isa_structure"]["assay"]["rows"] = matrix["assay"]["rows"]
    (deliverables / "metadata.json").write_text(json.dumps(metadata))
    entity_plan = json.loads((run_dir / "reports" / "entity_plan.json").read_text())
    entity_plan["validation"]["row_counts"]["assay"] = 2
    (run_dir / "reports" / "entity_plan.json").write_text(json.dumps(entity_plan))
    workbook = load_workbook_for_test(deliverables / "metadata_fairds.xlsx")
    workbook["assay"].append(["assay-2", "measurement", "sample-1"])
    workbook.save(deliverables / "metadata_fairds.xlsx")

    report = audit_run(run_dir)
    codes = {issue["code"] for issue in report["issues"]}

    assert report["passed"] is False
    assert "person_sheet_forbidden" in codes
    assert "unexpected_sheet" in codes
    assert "meaningless_repeated_entity_rows" in codes


def test_audit_rejects_broadcast_entity_name_against_plan(tmp_path: Path):
    run_dir = _write_run(tmp_path)
    deliverables = run_dir / "deliverables"
    matrix = json.loads((deliverables / "isa_values.json").read_text())
    matrix["assay"]["columns"].append("assay name")
    matrix["assay"]["rows"][0]["assay name"] = "broadcast name"
    metadata = json.loads((deliverables / "metadata.json").read_text())
    metadata["isa_values"] = matrix
    metadata["isa_structure"]["assay"]["columns"] = matrix["assay"]["columns"]
    metadata["isa_structure"]["assay"]["fields"].append(
        {"field_name": "assay name", "required": False}
    )
    (deliverables / "metadata.json").write_text(json.dumps(metadata))
    (deliverables / "isa_values.json").write_text(json.dumps(matrix))
    entity_plan = json.loads((run_dir / "reports" / "entity_plan.json").read_text())
    entity_plan["plan"] = {
        "levels": [
            {
                "level": "assay",
                "entities": [
                    {
                        "row_id": "assay-1",
                        "external_identifier": None,
                        "label": "planned assay name",
                    }
                ],
            }
        ]
    }
    (run_dir / "reports" / "entity_plan.json").write_text(json.dumps(entity_plan))
    workbook = load_workbook_for_test(deliverables / "metadata_fairds.xlsx")
    sheet = workbook["assay"]
    sheet.cell(row=1, column=4, value="assay name")
    sheet.cell(row=2, column=4, value="broadcast name")
    workbook.save(deliverables / "metadata_fairds.xlsx")

    report = audit_run(run_dir)
    codes = {issue["code"] for issue in report["issues"]}

    assert report["passed"] is False
    assert "entity_name_plan_mismatch" in codes


def test_audit_warnings_require_review_instead_of_passing(tmp_path: Path):
    run_dir = _write_run(tmp_path)
    metadata_path = run_dir / "deliverables" / "metadata.json"
    metadata = json.loads(metadata_path.read_text())
    metadata["needs_review"] = True
    metadata["isa_values"]["sample"]["rows"][0]["sample description"] = ""
    metadata["isa_structure"]["sample"]["rows"][0]["sample description"] = ""
    metadata_path.write_text(json.dumps(metadata))
    matrix_path = run_dir / "deliverables" / "isa_values.json"
    matrix = json.loads(matrix_path.read_text())
    matrix["sample"]["rows"][0]["sample description"] = ""
    matrix_path.write_text(json.dumps(matrix))
    workbook = load_workbook_for_test(
        run_dir / "deliverables" / "metadata_fairds.xlsx"
    )
    workbook["sample"].cell(row=2, column=2, value="")
    workbook.save(run_dir / "deliverables" / "metadata_fairds.xlsx")

    report = audit_run(run_dir)

    assert report["structural_passed"] is True
    assert report["review_required"] is True
    assert report["passed"] is False


def test_audit_rejects_noncanonical_controlled_values(tmp_path: Path):
    run_dir = _write_run(tmp_path)
    deliverables = run_dir / "deliverables"
    matrix = json.loads((deliverables / "isa_values.json").read_text())
    matrix["assay"]["columns"].append("library strategy")
    matrix["assay"]["rows"][0]["library strategy"] = "mRNA-Seq"
    (deliverables / "isa_values.json").write_text(json.dumps(matrix))

    metadata = json.loads((deliverables / "metadata.json").read_text())
    metadata["isa_values"] = matrix
    metadata["isa_structure"]["assay"]["columns"] = matrix["assay"]["columns"]
    metadata["isa_structure"]["assay"]["rows"] = matrix["assay"]["rows"]
    metadata["isa_structure"]["assay"]["fields"].append(
        {"field_name": "library strategy", "required": False}
    )
    metadata["_field_definitions"] = [
        {
            "term": "library strategy",
            "isa_sheet": "assay",
            "regex": "(WGS|RNA-Seq|AMPLICON)",
        }
    ]
    (deliverables / "metadata.json").write_text(json.dumps(metadata))

    workbook = load_workbook_for_test(deliverables / "metadata_fairds.xlsx")
    sheet = workbook["assay"]
    column = sheet.max_column + 1
    sheet.cell(row=1, column=column, value="library strategy")
    sheet.cell(row=2, column=column, value="mRNA-Seq")
    workbook.save(deliverables / "metadata_fairds.xlsx")

    report = audit_run(run_dir)
    codes = {issue["code"] for issue in report["issues"]}

    assert report["structural_passed"] is False
    assert "controlled_value_contract_violation" in codes


def test_audit_rejects_column_without_persisted_fairds_contract(tmp_path: Path):
    run_dir = _write_run(tmp_path)
    metadata_path = run_dir / "deliverables" / "metadata.json"
    metadata = json.loads(metadata_path.read_text())
    metadata["_field_definitions"] = [
        item
        for item in metadata["_field_definitions"]
        if not (
            item["isa_sheet"] == "sample"
            and item["term"] == "sample description"
        )
    ]
    metadata_path.write_text(json.dumps(metadata))

    report = audit_run(run_dir)
    codes = {issue["code"] for issue in report["issues"]}

    assert report["structural_passed"] is False
    assert "columns_missing_fairds_contract" in codes


def test_audit_rejects_column_from_undeclared_package(tmp_path: Path):
    run_dir = _write_run(tmp_path)
    metadata_path = run_dir / "deliverables" / "metadata.json"
    metadata = json.loads(metadata_path.read_text())
    definition = next(
        item
        for item in metadata["_field_definitions"]
        if item["isa_sheet"] == "assay"
        and item["term"] == "assay description"
    )
    definition["package"] = "not-selected"
    metadata_path.write_text(json.dumps(metadata))

    report = audit_run(run_dir)
    codes = {issue["code"] for issue in report["issues"]}

    assert report["structural_passed"] is False
    assert "column_package_not_selected" in codes


def test_audit_accepts_explicit_source_table_extension(tmp_path: Path):
    run_dir = _write_run(tmp_path)
    deliverables = run_dir / "deliverables"
    metadata_path = deliverables / "metadata.json"
    metadata = json.loads(metadata_path.read_text())
    matrix = metadata["isa_values"]
    matrix["sample"]["columns"].append("source read count")
    matrix["sample"]["rows"][0]["source read count"] = "101"
    metadata["isa_structure"]["sample"]["columns"].append("source read count")
    metadata["isa_structure"]["sample"]["rows"][0]["source read count"] = "101"
    metadata["_field_definitions"].append(
        {
            "term": "source read count",
            "isa_sheet": "sample",
            "package": "Source metadata extension",
            "source_extension": True,
            "source_column": "input.reads",
            "requirement": "OPTIONAL",
        }
    )
    metadata_path.write_text(json.dumps(metadata))
    (deliverables / "isa_values.json").write_text(json.dumps(matrix))
    workbook = load_workbook_for_test(deliverables / "metadata_fairds.xlsx")
    worksheet = workbook["sample"]
    worksheet.cell(row=1, column=worksheet.max_column + 1, value="source read count")
    worksheet.cell(row=2, column=worksheet.max_column, value="101")
    workbook.save(deliverables / "metadata_fairds.xlsx")

    report = audit_run(run_dir)
    codes = {issue["code"] for issue in report["issues"]}

    assert "columns_missing_fairds_contract" not in codes
    assert "column_package_not_selected" not in codes


def load_workbook_for_test(path: Path):
    from openpyxl import load_workbook

    return load_workbook(path)


def _write_expectation(tmp_path: Path) -> Path:
    source = tmp_path / "source.txt"
    source.write_text("A minimal source", encoding="utf-8")
    spec = {
        "schema_version": "fairiagent.fairds_expected_outcome.v1",
        "case_id": "minimal",
        "source": {
            "path": "source.txt",
            "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        },
        "packages": {
            "required": ["default", "focused-assay"],
            "forbidden": ["unrelated"],
        },
        "required_sheets": [
            "investigation",
            "study",
            "observationunit",
            "sample",
            "assay",
            "Help",
        ],
        "forbidden_sheets": ["Person"],
        "allowed_audit_warning_codes": [],
        "levels": {
            "sample": {
                "row_count": 1,
                "required_columns": [
                    "sample identifier",
                    "sample description",
                    "observation unit identifier",
                ],
                "unique_fields": ["sample identifier"],
                "fields": {
                    "sample description": {
                        "value_counts": {"biological material": 1}
                    }
                },
            },
            "assay": {
                "row_count": 1,
                "unique_fields": ["assay identifier", "sample identifier"],
            },
        },
        "relationships": [
            {
                "child_level": "assay",
                "parent_level": "sample",
                "link_field": "sample identifier",
                "one_child_per_parent": True,
            }
        ],
    }
    path = tmp_path / "expected.json"
    path.write_text(json.dumps(spec), encoding="utf-8")
    return path


def test_expected_outcome_accepts_matching_deliverable(tmp_path: Path):
    run_dir = _write_run(tmp_path)
    expectation = _write_expectation(tmp_path)

    result = evaluate_expected_outcome(
        run_dir, expectation, project_root=tmp_path
    )

    assert result["passed"] is True
    assert result["expectation_error_count"] == 0


def test_expected_outcome_accepts_fairds_package_sheet_suffixes(tmp_path: Path):
    run_dir = _write_run(tmp_path)
    expectation = _write_expectation(tmp_path)
    workbook_path = run_dir / "deliverables" / "metadata_fairds.xlsx"
    workbook = load_workbook_for_test(workbook_path)
    for worksheet in workbook.worksheets:
        if worksheet.title != "Help":
            worksheet.title = f"{worksheet.title} - default"
    workbook.save(workbook_path)

    result = evaluate_expected_outcome(
        run_dir, expectation, project_root=tmp_path
    )

    assert result["passed"] is True
    assert result["expectation_error_count"] == 0


def test_expected_outcome_checks_exact_source_table_projection(tmp_path: Path):
    run_dir = _write_run(tmp_path)
    expectation = _write_expectation(tmp_path)
    source_table = tmp_path / "source_table.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "records"
    worksheet.append(["group", "source value"])
    worksheet.append(["focal", "biological material"])
    worksheet.append(["other", "ignored material"])
    workbook.save(source_table)
    spec = json.loads(expectation.read_text())
    spec["source_table_checks"] = [
        {
            "path": "source_table.xlsx",
            "sheet": "records",
            "header_row": 1,
            "filter": {"field": "group", "value": "focal"},
            "source_field": "source value",
            "target_level": "sample",
            "target_field": "sample description",
            "comparison": "exact_multiset",
        }
    ]
    expectation.write_text(json.dumps(spec), encoding="utf-8")

    result = evaluate_expected_outcome(
        run_dir, expectation, project_root=tmp_path
    )

    assert result["passed"] is True
    worksheet.cell(row=2, column=2, value="different material")
    workbook.save(source_table)
    result = evaluate_expected_outcome(
        run_dir, expectation, project_root=tmp_path
    )
    assert "source_table_projection_mismatch" in {
        error["code"] for error in result["expectation_errors"]
    }


def test_expected_outcome_checks_contract_breadth_and_keyed_projection(
    tmp_path: Path,
):
    run_dir = _write_run(tmp_path)
    expectation = _write_expectation(tmp_path)
    second_source = tmp_path / "supplement.txt"
    second_source.write_text("Pinned supplementary source", encoding="utf-8")
    source_table = tmp_path / "source_table.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "records"
    worksheet.append(["sample key", "assay annotation"])
    worksheet.append(["biological material", "measurement"])
    workbook.save(source_table)
    spec = json.loads(expectation.read_text())
    spec["schema_version"] = "fairiagent.fairds_expected_outcome.v2"
    spec["source_assets"] = [
        {
            "path": "supplement.txt",
            "sha256": hashlib.sha256(second_source.read_bytes()).hexdigest(),
            "role": "supplementary_methods",
        }
    ]
    spec["packages"]["required_any_of"] = [
        ["alternative-package", "focused-assay"]
    ]
    spec["levels"]["sample"]["fields"]["sample description"] = {
        "min_nonempty_count": 1,
        "max_nonempty_count": 1,
        "required_substrings": ["biological"],
        "value_regex": "biological material",
    }
    spec["source_table_checks"] = [
        {
            "path": "source_table.xlsx",
            "sheet": "records",
            "header_row": 1,
            "source_key_field": "sample key",
            "source_field": "assay annotation",
            "target_level": "assay",
            "target_field": "assay description",
            "target_key_lookup": {
                "level": "sample",
                "target_link_field": "sample identifier",
                "lookup_id_field": "sample identifier",
                "lookup_value_field": "sample description",
            },
            "comparison": "exact_keyed",
        }
    ]
    expectation.write_text(json.dumps(spec), encoding="utf-8")

    result = evaluate_expected_outcome(run_dir, expectation, project_root=tmp_path)

    assert result["passed"] is True
    assert result["contract_coverage"] == {
        "pinned_source_asset_count": 2,
        "asserted_level_count": 2,
        "asserted_field_count": 1,
        "source_table_check_count": 1,
        "coverage_ledger_entry_count": 0,
    }


def test_expected_outcome_checks_numeric_values_by_source_key(tmp_path: Path):
    run_dir = _write_run(tmp_path)
    expectation = _write_expectation(tmp_path)
    deliverables = run_dir / "deliverables"
    matrix = json.loads((deliverables / "isa_values.json").read_text())
    matrix["sample"]["rows"][0]["sample description"] = "1.0001"
    (deliverables / "isa_values.json").write_text(json.dumps(matrix))
    metadata = json.loads((deliverables / "metadata.json").read_text())
    metadata["isa_values"] = matrix
    metadata["isa_structure"]["sample"]["rows"] = matrix["sample"]["rows"]
    (deliverables / "metadata.json").write_text(json.dumps(metadata))
    workbook = load_workbook_for_test(deliverables / "metadata_fairds.xlsx")
    workbook["sample"].cell(row=2, column=2, value="1.0001")
    workbook.save(deliverables / "metadata_fairds.xlsx")
    source_table = tmp_path / "source_table.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "records"
    worksheet.append(["sample", "measurement"])
    worksheet.append(["sample-1", 1.0])
    workbook.save(source_table)
    spec = json.loads(expectation.read_text())
    spec["levels"]["sample"]["fields"]["sample description"] = {
        "value_counts": {"1.0001": 1}
    }
    spec["source_table_checks"] = [
        {
            "path": "source_table.xlsx",
            "sheet": "records",
            "header_row": 1,
            "source_key_field": "sample",
            "source_field": "measurement",
            "target_level": "sample",
            "target_key_field": "sample identifier",
            "target_field": "sample description",
            "comparison": "numeric_keyed",
            "absolute_tolerance": 0.001,
        }
    ]
    expectation.write_text(json.dumps(spec), encoding="utf-8")

    result = evaluate_expected_outcome(run_dir, expectation, project_root=tmp_path)

    assert result["passed"] is True


def test_expected_outcome_rejects_semantic_value_regression(tmp_path: Path):
    run_dir = _write_run(tmp_path)
    expectation = _write_expectation(tmp_path)
    deliverables = run_dir / "deliverables"
    matrix = json.loads((deliverables / "isa_values.json").read_text())
    matrix["sample"]["rows"][0]["sample description"] = "wrong material"
    (deliverables / "isa_values.json").write_text(json.dumps(matrix))
    metadata = json.loads((deliverables / "metadata.json").read_text())
    metadata["isa_values"] = matrix
    metadata["isa_structure"]["sample"]["rows"] = matrix["sample"]["rows"]
    (deliverables / "metadata.json").write_text(json.dumps(metadata))
    workbook = load_workbook_for_test(deliverables / "metadata_fairds.xlsx")
    workbook["sample"].cell(row=2, column=2, value="wrong material")
    workbook.save(deliverables / "metadata_fairds.xlsx")

    result = evaluate_expected_outcome(
        run_dir, expectation, project_root=tmp_path
    )

    assert result["passed"] is False
    assert "value_distribution_mismatch" in {
        error["code"] for error in result["expectation_errors"]
    }
