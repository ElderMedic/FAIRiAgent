import asyncio
import json
import logging
from pathlib import Path
from types import SimpleNamespace

from fastapi import FastAPI
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from fairifier.apps.api.routers.v1 import (
    _build_demo_document_response,
    _build_word_entries,
    _resolve_default_demo_document_key,
    get_artifact,
    list_artifacts,
    memory_cloud,
    router,
)
from fairifier.apps.api.services.runner import (
    _serialisable_artifacts,
    _start_full_output_capture,
    _stop_full_output_capture,
    _persist_run_outputs,
    _serialisable_artifacts,
)
from fairifier.apps.api.storage.sqlite_store import (
    SQLiteProjectStore,
)
from fairifier.apps.api.models import DemoDocumentResponse
from fairifier.utils.json_logger import JSONLogger

SESSION_HEADERS = {
    "X-FAIRifier-Session-Id": "11111111-1111-4111-8111-111111111111",
    "X-FAIRifier-Session-Started-At": "2026-04-01T10:00:00+00:00",
}


def test_serialisable_artifacts_falls_back_to_output_filenames():
    assert _serialisable_artifacts(
        {
            "metadata_json": "{}",
            "runtime_config": "{}",
            "validation_report": "ok",
        }
    ) == [
        "metadata.json",
        "runtime_config.json",
        "validation_report.txt",
    ]


def test_serialisable_artifacts_prefers_actual_output_files(tmp_path):
    (tmp_path / "metadata.json").write_text("{}", encoding="utf-8")
    (tmp_path / "metadata_fairds.xlsx").write_bytes(b"xlsx")
    (tmp_path / "runtime_config.json").write_text("{}", encoding="utf-8")
    (tmp_path / ".hidden").write_text("skip", encoding="utf-8")

    assert _serialisable_artifacts(
        {"metadata_json": "{}"},
        output_dir=str(tmp_path),
    ) == [
        "metadata.json",
        "metadata_fairds.xlsx",
        "runtime_config.json",
    ]


def test_list_artifacts_includes_nested_files_and_downloads(
    tmp_path,
):
    output_dir = tmp_path / "output"
    output_dir.mkdir()
    (output_dir / "workflow_report.json").write_text(
        "{}", encoding="utf-8"
    )
    nested_dir = output_dir / "mineru_doc"
    nested_dir.mkdir()
    nested_file = nested_dir / "report.md"
    nested_file.write_text(
        "# converted", encoding="utf-8"
    )
    (output_dir / ".hidden").write_text(
        "skip", encoding="utf-8"
    )
    hidden_nested = nested_dir / ".skip.txt"
    hidden_nested.write_text(
        "skip", encoding="utf-8"
    )

    store = SQLiteProjectStore(
        str(tmp_path / "projects.db")
    )
    app = FastAPI()
    app.state.store = store
    app.include_router(router)

    store.create_project(
        "proj-1",
        {
            "project_id": "proj-1",
            "project_name": "Test Project",
            "session_id": SESSION_HEADERS[
                "X-FAIRifier-Session-Id"
            ],
            "session_started_at": SESSION_HEADERS[
                "X-FAIRifier-Session-Started-At"
            ],
            "status": "completed",
            "output_dir": str(output_dir),
        },
    )

    try:
        request = SimpleNamespace(
            headers=SESSION_HEADERS,
            query_params={},
            app=SimpleNamespace(
                state=SimpleNamespace(store=store)
            ),
        )
        payload = asyncio.run(
            list_artifacts("proj-1", request)
        )
        assert payload["project_id"] == "proj-1"
        assert payload["artifacts"] == [
            {
                "name": "mineru_doc/report.md",
                "size": len("# converted"),
                "available": True,
            },
            {
                "name": "workflow_report.json",
                "size": len("{}"),
                "available": True,
            },
        ]

        download = asyncio.run(
            get_artifact(
                "proj-1",
                "mineru_doc/report.md",
                request,
            )
        )
        assert download.path == str(nested_file)
        assert download.filename == "mineru_doc/report.md"
        assert (
            Path(download.path).read_text(encoding="utf-8")
            == "# converted"
        )
    finally:
        store.close()


def test_persist_run_outputs_writes_core_downloadable_files(
    tmp_path,
):
    json_logger = JSONLogger(
        component="test", enable_stdout=False
    )
    json_logger.log_processing_start(
        "document.pdf", "proj-2"
    )
    json_logger.log_processing_end(
        "proj-2", "completed", 1.25
    )

    errors = _persist_run_outputs(
        project_id="proj-2",
        result={
            "artifacts": {
                "metadata_json": '{"ok": true}',
                "validation_report": "looks good",
            }
        },
        output_dir=str(tmp_path),
        json_logger=json_logger,
    )

    assert errors == []
    assert (
        tmp_path / "deliverables" / "metadata.json"
    ).read_text(encoding="utf-8") == '{"ok": true}'
    assert (
        tmp_path / "reports" / "validation_report.txt"
    ).read_text(encoding="utf-8") == "looks good"
    processing_log = (
        tmp_path / "logs" / "processing_log.jsonl"
    ).read_text(encoding="utf-8")
    assert "processing_started" in processing_log
    assert "processing_completed" in processing_log
    assert "artifact_saved" in processing_log
    assert "workflow_result_summary" in processing_log


def test_persist_run_outputs_writes_react_scratchpad_and_full_errors(
    tmp_path,
):
    json_logger = JSONLogger(
        component="test", enable_stdout=False
    )
    result_errors = [f"error-{i}" for i in range(15)]
    errors = _persist_run_outputs(
        project_id="proj-scratch",
        result={
            "status": "completed",
            "errors": result_errors,
            "artifacts": {
                "react_scratchpad": {
                    "ISAValueMapper": {
                        "iterations": 2,
                        "tools_called": ["search_workspace"],
                    }
                }
            },
        },
        output_dir=str(tmp_path),
        json_logger=json_logger,
    )

    assert errors == []
    scratch = json.loads(
        (tmp_path / "logs" / "react_scratchpad.json").read_text(
            encoding="utf-8"
        )
    )
    assert scratch["ISAValueMapper"]["iterations"] == 2
    processing_log = (
        tmp_path / "logs" / "processing_log.jsonl"
    ).read_text(encoding="utf-8")
    for item in result_errors:
        assert item in processing_log


def test_persist_run_outputs_writes_auto_repair_trace_for_gate(
    tmp_path,
):
    json_logger = JSONLogger(
        component="test", enable_stdout=False
    )

    errors = _persist_run_outputs(
        project_id="proj-auto",
        result={
            "artifacts": {
                "metadata_json": '{"isa_structure": {"study": {"fields": []}}}',
                "auto_repair_trace": {
                    "mode": "deterministic_exact_patch",
                    "summary": {
                        "accepted_patch_count": 0,
                        "metadata_mutated": False,
                    },
                },
            }
        },
        output_dir=str(tmp_path),
        json_logger=json_logger,
    )

    assert errors == []
    trace = json.loads(
        (tmp_path / "reports" / "auto_repair_trace.json").read_text(encoding="utf-8")
    )
    assert trace["mode"] == "deterministic_exact_patch"
    assert trace["summary"]["accepted_patch_count"] == 0


def test_persist_run_outputs_exports_fairds_workbook_from_isa_values(
    tmp_path,
):
    json_logger = JSONLogger(
        component="test", enable_stdout=False
    )
    metadata = {
        "isa_structure": {
            "study": {
                "columns": ["study title"],
                "rows": [{"study title": "fallback"}],
            }
        }
    }
    isa_values = {
        "study": {
            "columns": ["study title"],
            "rows": [{"study title": "Pea cold stress response"}],
        }
    }

    errors = _persist_run_outputs(
        project_id="proj-fairds",
        result={
            "artifacts": {
                "metadata_json": json.dumps(metadata),
                "isa_values_json": json.dumps(isa_values),
            }
        },
        output_dir=str(tmp_path),
        json_logger=json_logger,
        fair_ds_api_url="",
    )

    assert errors == []
    workbook_path = tmp_path / "deliverables" / "metadata_fairds.xlsx"
    assert workbook_path.exists()
    workbook = load_workbook(workbook_path, data_only=True)
    try:
        worksheet = next(
            sheet
            for sheet in workbook.worksheets
            if sheet.title.lower() == "study"
            or sheet.title.lower().startswith("study - ")
        )
        headers = [
            worksheet.cell(1, col).value
            for col in range(1, worksheet.max_column + 1)
        ]
        row = {
            headers[col - 1]: worksheet.cell(2, col).value
            for col in range(1, worksheet.max_column + 1)
        }
        assert row["study title"] == "Pea cold stress response"
    finally:
        workbook.close()
    log_entries = [
        json.loads(line)
        for line in (tmp_path / "logs" / "processing_log.jsonl")
        .read_text(encoding="utf-8")
        .splitlines()
        if line.strip()
    ]
    summary = next(
        entry for entry in log_entries
        if entry.get("event") == "workflow_result_summary"
    )
    assert "deliverables/metadata.json" in summary["artifact_names"]
    assert "deliverables/isa_values.json" in summary["artifact_names"]
    assert "deliverables/metadata_fairds.xlsx" in summary["artifact_names"]
    assert summary["artifact_keys"] == ["isa_values_json", "metadata_json"]



def test_serialisable_artifacts_falls_back_to_output_filenames():
    assert _serialisable_artifacts(
        {
            "metadata_json": "{}",
            "auto_repair_trace": "{}",
            "validation_report": "ok",
        }
    ) == [
        "metadata.json",
        "auto_repair_trace.json",
        "validation_report.txt",
    ]


def test_serialisable_artifacts_prefers_actual_output_files(tmp_path):
    (tmp_path / "metadata.json").write_text("{}", encoding="utf-8")
    (tmp_path / "metadata_fairds.xlsx").write_bytes(b"xlsx")
    (tmp_path / "runtime_config.json").write_text("{}", encoding="utf-8")
    (tmp_path / ".hidden").write_text("skip", encoding="utf-8")

    assert _serialisable_artifacts(
        {"metadata_json": "{}"},
        output_dir=str(tmp_path),
    ) == [
        "metadata.json",
        "metadata_fairds.xlsx",
        "runtime_config.json",
    ]


def test_full_output_capture_writes_root_logger_messages(tmp_path):
    log_path, handle, handler = _start_full_output_capture(
        str(tmp_path),
        project_id="proj-3",
        file_path="document.pdf",
    )
    try:
        logging.getLogger("fairifier.agents.test").warning("agent log line")
    finally:
        _stop_full_output_capture(handler, handle)

    assert log_path is not None
    text = log_path.read_text(encoding="utf-8")
    assert "Project ID: proj-3" in text
    assert "Input: document.pdf" in text
    assert "agent log line" in text
    assert (tmp_path / "logs").is_dir()
    assert not (tmp_path / "deliverables").exists()
    assert not (tmp_path / "reports").exists()
    assert not (tmp_path / "workspace").exists()


def test_default_demo_document_key_falls_back_to_available_sample():
    documents = [
      DemoDocumentResponse(
          key="earthworm_paper",
          label="Earthworm BioRxiv Paper",
          filename="earthworm_4n_paper_bioRxiv.pdf",
          description="Representative PDF example.",
          size_bytes=123,
      )
    ]

    assert (
        _resolve_default_demo_document_key(documents)
        == "earthworm_paper"
    )


def test_demo_document_response_lists_multi_source_bundle(tmp_path):
    paper = tmp_path / "paper.md"
    workbook = tmp_path / "supplement.xlsx"
    paper.write_text("# Demo", encoding="utf-8")
    workbook.write_bytes(b"xlsx")

    response = _build_demo_document_response(
        "multi_source",
        {
            "path": paper,
            "paths": [paper, workbook],
            "label": "Multi-source demo",
            "description": "Paper plus workbook",
        },
    )

    assert response.filename == "paper.md"
    assert response.files == ["paper.md", "supplement.xlsx"]


def test_memory_cloud_separates_run_and_user_memory(
    monkeypatch,
    tmp_path,
):
    store = SQLiteProjectStore(
        str(tmp_path / "projects.db")
    )
    app = FastAPI()
    app.state.store = store
    app.include_router(router)

    store.create_project(
        "proj-memory",
        {
            "project_id": "proj-memory",
            "project_name": "Memory Project",
            "session_id": SESSION_HEADERS[
                "X-FAIRifier-Session-Id"
            ],
            "session_started_at": SESSION_HEADERS[
                "X-FAIRifier-Session-Started-At"
            ],
            "status": "completed",
        },
    )

    class FakeMem0:
        def is_available(self):
            return True

        def list_memories(self, session_id, agent_id=None):
            if session_id == "proj-memory":
                return [
                    {
                        "memory": "nanotoxicology uses soil package",
                        "metadata": {"agent_id": "KnowledgeRetriever"},
                    },
                ]
            if session_id == SESSION_HEADERS[
                "X-FAIRifier-Session-Id"
            ]:
                return [
                    {
                        "memory": "earthworm studies use ENVO ontology",
                        "metadata": {"agent_id": "DocumentParser"},
                    },
                    {
                        "memory": "soil studies prefer MIxS fields",
                        "metadata": {"agent_id": "JSONGenerator"},
                    },
                ]
            return []

    monkeypatch.setattr(
        "fairifier.services.mem0_service.get_mem0_service",
        lambda: FakeMem0(),
    )
    monkeypatch.setattr(
        "fairifier.config.config.memory_scope_id",
        None,
    )

    try:
        request = SimpleNamespace(
            headers=SESSION_HEADERS,
            query_params={},
            app=SimpleNamespace(
                state=SimpleNamespace(store=store)
            ),
        )
        response = asyncio.run(
            memory_cloud("proj-memory", request)
        )
        payload = response.model_dump()
        assert payload["memory_enabled"] is True
        assert payload["session_total"] == 1
        assert payload["scope_total"] == 3
        assert {
            entry["text"] for entry in payload["session_words"]
        } >= {"nanotoxicology", "soil", "package"}
        assert {
            entry["text"] for entry in payload["scope_words"]
        } >= {
            "nanotoxicology",
            "earthworm",
            "ontology",
            "prefer",
        }
    finally:
        store.close()


def test_memory_cloud_requires_matching_session(
    monkeypatch,
    tmp_path,
):
    store = SQLiteProjectStore(str(tmp_path / "projects.db"))
    app = FastAPI()
    app.state.store = store
    app.include_router(router)

    store.create_project(
        "proj-a",
        {
            "project_id": "proj-a",
            "project_name": "Protected Project",
            "session_id": SESSION_HEADERS[
                "X-FAIRifier-Session-Id"
            ],
            "session_started_at": SESSION_HEADERS[
                "X-FAIRifier-Session-Started-At"
            ],
            "status": "completed",
        },
    )

    class FakeMem0:
        def is_available(self):
            return True

        def list_memories(self, session_id, agent_id=None):
            return []

    monkeypatch.setattr(
        "fairifier.services.mem0_service.get_mem0_service",
        lambda: FakeMem0(),
    )

    foreign_headers = {
        "X-FAIRifier-Session-Id": "22222222-2222-4222-8222-222222222222",
        "X-FAIRifier-Session-Started-At": "2026-04-01T11:00:00+00:00",
    }

    try:
        foreign_request = SimpleNamespace(
            headers=foreign_headers,
            query_params={},
            app=SimpleNamespace(
                state=SimpleNamespace(store=store)
            ),
        )
        missing_request = SimpleNamespace(
            headers=SESSION_HEADERS,
            query_params={},
            app=SimpleNamespace(
                state=SimpleNamespace(store=store)
            ),
        )

        try:
            asyncio.run(
                memory_cloud("proj-a", foreign_request)
            )
            raise AssertionError("Expected HTTPException for foreign session")
        except HTTPException as exc:
            assert exc.status_code == 404

        try:
            asyncio.run(
                memory_cloud("missing-project", missing_request)
            )
            raise AssertionError("Expected HTTPException for missing project")
        except HTTPException as exc:
            assert exc.status_code == 404
    finally:
        store.close()


def test_memory_word_entries_include_singletons():
    entries = _build_word_entries(
        [
            ("nanotoxicology", "KnowledgeRetriever"),
            ("package", "KnowledgeRetriever"),
        ]
    )

    assert [entry.model_dump() for entry in entries] == [
        {
            "text": "nanotoxicology",
            "value": 1,
            "category": "KnowledgeRetriever",
        },
        {
            "text": "package",
            "value": 1,
            "category": "KnowledgeRetriever",
        },
    ]
