"""Standalone FAIR-DS metadata Excel generator with multi-row support.

Generates ``metadata_fairds.xlsx`` from the ``isa_structure`` in
``metadata.json``.

Features
--------
* **Local generation** — no API dependency; can produce a fully formatted
  workbook from the ISA structure alone.
* **Multi-row support** — each ISA sheet gets one Excel row per entity
  (the ``columns`` + ``rows`` format).
* **Entity splitting** — semicolon-separated values, "Experiment N" patterns,
  and other heuristics expand single merged rows into per-entity rows.
* **Backward compatibility** — the flat ``fields`` list is still accepted.
* **FAIR-DS API integration** — when an API URL is provided and reachable,
  ``POST /api/isa`` may supply a partial workbook (often only investigation /
  study). Local post-processing creates any missing ISA sheets and replaces
  placeholder rows from the compiled ``columns`` × ``rows`` matrix.
* **Help sheet** — lists used metadata types in FAIR-DS layout (term,
  definition, requirement, package, example), built from ``isa_structure``
  fields even when data rows come from a compiled ``isa_values`` sidecar.
* **Graceful degradation** — if ``openpyxl`` is not installed, the function
  falls back to the API path; if neither works it returns ``None``.
"""

from __future__ import annotations

import io
import json
import logging
import re
from copy import copy
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from fairifier.utils.entity_splitter import split_entities_in_isa_structure
from fairifier.utils.isa_order import ISA_LEVEL_ORDER

ISA_WORKBOOK_LEVEL_ORDER = ISA_LEVEL_ORDER

logger = logging.getLogger(__name__)

_XML_ILLEGAL_RE = re.compile(r"[\x00-\x08\x0B-\x0C\x0E-\x1F\uFFFE\uFFFF]")
_HEADER_BADGE_RE = re.compile(r"\s*\(([MRO])\)\s*$")
_API_FIELD_KEYS = (
    "field_name",
    "value",
    "package_source",
    "requirement",
    "isa_sheet",
    "isa_level",
    "origin",
    "status",
    "confidence",
)

_LOCAL_TERM_CATALOG: Optional[Dict[str, Dict[str, Dict[str, str]]]] = None


def _excel_safe_value(value: Any) -> str:
    return _XML_ILLEGAL_RE.sub("", str(value))

# ── Row resolution ──────────────────────────────────────────────


def _resolve_rows(level_data: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Resolve data rows from an ISA sheet's *level_data*.

    Prefers ``rows`` (``columns`` + ``rows`` multi-row format); falls back
    to the flat ``fields`` list (backward-compatible single-row).

    Returns
    -------
    List[Dict[str, Any]]
        Each dict maps a **lowercase field name** to its string value.
        Empty list if no data is present.
    """
    rows: List[Dict[str, Any]] = level_data.get("rows", [])
    if rows:
        return rows

    fields: List[Dict[str, Any]] = level_data.get("fields", [])
    if fields:
        row: Dict[str, Any] = {}
        for f in fields:
            name = (f.get("field_name") or "").strip().lower()
            val = f.get("value")
            if name and val is not None:
                row[name] = str(val)
        if row:
            return [row]

    return []


# ── Header-column map builder ────────────────────────────────────


def _build_header_column_map(ws: "openpyxl.worksheet.worksheet.Worksheet") -> Dict[str, int]:
    """Build ``{lowercase_header: column_1_indexed}`` from Excel row 1."""
    header_col: Dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        h = ws.cell(row=1, column=col_idx).value
        if h:
            # Read legacy FAIRiAgent workbooks defensively, but never make the
            # non-standard requirement badge part of a FAIR-DS term name.
            normalized = _HEADER_BADGE_RE.sub("", str(h)).strip().lower()
            header_col[normalized] = col_idx
    return header_col


# ── Local Excel generation (no API needed) ───────────────────────


# ── requirement-label styling ───────────────────────────────────────
def _first_text(*values: Any) -> str:
    for value in values:
        if value is None or isinstance(value, (dict, list)):
            continue
        text = str(value).strip()
        if text:
            return text
    return ""


def _normalize_term_name(name: Any) -> str:
    text = _first_text(name)
    return _HEADER_BADGE_RE.sub("", text).strip()


def _empty_term_catalog() -> Dict[str, Dict[str, Dict[str, str]]]:
    return {"by_label": {}, "by_sheet_label": {}, "by_package_label": {}}


def _store_term_record(
    catalog: Dict[str, Dict[str, Dict[str, str]]],
    record: Dict[str, str],
    *,
    label: str,
    sheet: str = "",
    package: str = "",
) -> None:
    label_key = label.strip().lower()
    if not label_key:
        return
    catalog["by_label"].setdefault(label_key, record)
    if sheet:
        catalog["by_sheet_label"][f"{sheet.strip().lower()}::{label_key}"] = record
    if package:
        catalog["by_package_label"][f"{package.strip().lower()}::{label_key}"] = record


def _lookup_term_record(
    catalog: Optional[Dict[str, Dict[str, Dict[str, str]]]],
    *,
    name: str,
    sheet: str = "",
    package: str = "",
) -> Dict[str, str]:
    if not catalog or not name:
        return {}
    key = name.strip().lower()
    if package:
        hit = catalog.get("by_package_label", {}).get(f"{package.strip().lower()}::{key}")
        if hit:
            return hit
    if sheet:
        hit = catalog.get("by_sheet_label", {}).get(f"{sheet.strip().lower()}::{key}")
        if hit:
            return hit
    return catalog.get("by_label", {}).get(key) or {}


def _index_package_metadata(
    package: Dict[str, Any],
    catalog: Dict[str, Dict[str, Dict[str, str]]],
) -> None:
    package_name = _first_text(package.get("packageName"), package.get("name"))
    for item in package.get("metadata") or []:
        if not isinstance(item, dict):
            continue
        term = item.get("term") if isinstance(item.get("term"), dict) else {}
        label = _normalize_term_name(
            item.get("label") or term.get("label") or item.get("field_name")
        )
        if not label:
            continue
        sheet = _first_text(
            item.get("sheetName"),
            item.get("level"),
            item.get("isa_sheet"),
            term.get("sheetName"),
        ).lower()
        record = {
            "label": label,
            "definition": _first_text(item.get("definition"), term.get("definition")),
            "example": _first_text(item.get("example"), term.get("example")),
            "requirement": _first_text(item.get("requirement"), term.get("requirement")).upper(),
            "package": _first_text(item.get("packageName"), package_name),
        }
        _store_term_record(
            catalog,
            record,
            label=label,
            sheet=sheet,
            package=record["package"],
        )


def _index_fairds_terms(
    terms: Optional[Dict[str, Any]],
) -> Dict[str, Dict[str, Dict[str, str]]]:
    catalog = _empty_term_catalog()
    if not isinstance(terms, dict):
        return catalog
    for name, info in terms.items():
        if not isinstance(info, dict):
            continue
        nested = info.get("term") if isinstance(info.get("term"), dict) else {}
        label = _normalize_term_name(info.get("label") or nested.get("label") or name)
        if not label:
            continue
        record = {
            "label": label,
            "definition": _first_text(info.get("definition"), nested.get("definition")),
            "example": _first_text(info.get("example"), nested.get("example")),
            "requirement": _first_text(info.get("requirement"), nested.get("requirement")).upper(),
            "package": _first_text(
                info.get("packageName"),
                info.get("package"),
                nested.get("packageName"),
            ),
        }
        sheet = _first_text(
            info.get("sheetName"), info.get("level"), info.get("isa_sheet")
        ).lower()
        _store_term_record(
            catalog,
            record,
            label=label,
            sheet=sheet,
            package=record["package"],
        )
    return catalog


def _load_local_term_catalog(
    force_refresh: bool = False,
) -> Dict[str, Dict[str, Dict[str, str]]]:
    """Index explicitly configured local extension packages for Help lookup."""
    global _LOCAL_TERM_CATALOG
    if _LOCAL_TERM_CATALOG is not None and not force_refresh:
        return _LOCAL_TERM_CATALOG

    catalog = _empty_term_catalog()
    try:
        from ..config import config

        sources = list(tuple(getattr(config, "local_package_paths", ()) or ()))
        files: List[Path] = []
        for source in sources:
            path = Path(source).expanduser()
            if not path.is_absolute():
                path = Path(config.project_root) / path
            candidates = sorted(path.glob("*_package.json")) if path.is_dir() else [path]
            for candidate in candidates:
                if candidate.is_file() and candidate not in files:
                    files.append(candidate)
        for package_file in files:
            try:
                payload = json.loads(package_file.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError) as exc:
                logger.debug("Skipping local package %s: %s", package_file, exc)
                continue
            if isinstance(payload, dict):
                _index_package_metadata(payload, catalog)
    except Exception as exc:
        logger.debug("Local Help term catalog unavailable: %s", exc)

    _LOCAL_TERM_CATALOG = catalog
    return catalog


def _merge_term_catalogs(
    base: Optional[Dict[str, Dict[str, Dict[str, str]]]],
    extra: Optional[Dict[str, Dict[str, Dict[str, str]]]],
) -> Dict[str, Dict[str, Dict[str, str]]]:
    merged = _empty_term_catalog()
    for bucket in merged:
        merged[bucket].update((base or {}).get(bucket, {}))
        # The live FAIR-DS catalog is authoritative. Local extension records
        # may fill omissions but must never shadow server package/term metadata.
        for key, record in ((extra or {}).get(bucket, {}) or {}).items():
            current = merged[bucket].get(key) or {}
            merged[bucket][key] = {
                field: value
                for field in set(current) | set(record)
                if (value := record.get(field) or current.get(field)) is not None
            }
    return merged


def _merge_fields_into_fill(
    fill_structure: Dict[str, Any],
    field_source: Dict[str, Any],
) -> Dict[str, Any]:
    """Keep compiled columns×rows while restoring ``fields`` for Help / API."""
    merged: Dict[str, Any] = {}
    for key, value in (fill_structure or {}).items():
        merged[key] = dict(value) if isinstance(value, dict) else value

    for level, block in (field_source or {}).items():
        if level == "_field_definitions" and isinstance(block, list):
            merged["_field_definitions"] = block
            continue
        if not isinstance(block, dict):
            continue
        dest = merged.get(level)
        if not isinstance(dest, dict):
            dest = {}
            merged[level] = dest
        else:
            dest = dict(dest)
            merged[level] = dest
        if block.get("fields"):
            dest["fields"] = block["fields"]
        if block.get("description") and not dest.get("description"):
            dest["description"] = block["description"]
        if block.get("column_metadata") and not dest.get("column_metadata"):
            dest["column_metadata"] = block["column_metadata"]
    return merged


def _slim_isa_for_api(isa_structure: Dict[str, Any]) -> Dict[str, Any]:
    """POST ``fields`` (not the columns×rows matrix) so FAIR-DS can fill Help."""
    slim: Dict[str, Any] = {}
    for level, block in (isa_structure or {}).items():
        if not isinstance(block, dict) or str(level).startswith("_"):
            continue
        if level not in ISA_LEVEL_ORDER:
            # Auxiliary sheets are added locally after the FAIR-DS API returns
            # its five-level workbook; older servers may reject unknown levels.
            continue
        fields = [
            {key: field[key] for key in _API_FIELD_KEYS if key in field and field[key] is not None}
            for field in (block.get("fields") or [])
            if isinstance(field, dict) and _normalize_term_name(field.get("field_name"))
        ]
        entry: Dict[str, Any] = {}
        if block.get("description"):
            entry["description"] = block["description"]
        if fields:
            entry["fields"] = fields
            slim[level] = entry
        elif block.get("columns") or block.get("rows"):
            if block.get("columns"):
                entry["columns"] = block["columns"]
            if block.get("rows"):
                entry["rows"] = block["rows"]
            slim[level] = entry
    return slim


def _help_section_title(level_name: str) -> str:
    return f"Below are the metadata terms that are used in the {level_name} sheet"


def _help_record_from_field(
    field: Dict[str, Any],
    *,
    sheet: str,
    catalog: Dict[str, Dict[str, Dict[str, str]]],
) -> Optional[Dict[str, str]]:
    term = field.get("term") if isinstance(field.get("term"), dict) else {}
    name = _normalize_term_name(
        field.get("field_name") or field.get("label") or term.get("label")
    )
    if not name:
        return None
    package = _first_text(field.get("package_source"), field.get("package"), field.get("packageName"))
    looked = _lookup_term_record(catalog, name=name, sheet=sheet, package=package)
    requirement = _first_text(field.get("requirement"), looked.get("requirement")).upper()
    return {
        "name": name,
        "definition": _first_text(
            field.get("definition"),
            term.get("definition"),
            looked.get("definition"),
        ),
        "requirement": requirement,
        "package": _first_text(package, looked.get("package")),
        "example": _first_text(
            field.get("example"),
            term.get("example"),
            looked.get("example"),
        ),
    }


def _collect_help_terms_for_level(
    level_name: str,
    field_source: Dict[str, Any],
    fill_structure: Optional[Dict[str, Any]],
    catalog: Dict[str, Dict[str, Dict[str, str]]],
) -> List[Dict[str, str]]:
    source_block = field_source.get(level_name)
    fill_block = (fill_structure or {}).get(level_name)
    if not isinstance(source_block, dict):
        source_block = {}
    if not isinstance(fill_block, dict):
        fill_block = {}

    terms: List[Dict[str, str]] = []
    seen: set[str] = set()

    def _append(record: Optional[Dict[str, str]]) -> None:
        if not record:
            return
        key = record["name"].strip().lower()
        if not key or key in seen:
            return
        seen.add(key)
        terms.append(record)

    for field in source_block.get("fields") or []:
        if isinstance(field, dict):
            _append(_help_record_from_field(field, sheet=level_name, catalog=catalog))

    for field in field_source.get("_field_definitions") or []:
        if not isinstance(field, dict):
            continue
        sheet = _first_text(field.get("isa_sheet"), field.get("isa_level")).lower()
        if sheet and sheet != level_name:
            continue
        _append(_help_record_from_field(field, sheet=level_name, catalog=catalog))

    columns: List[Any] = []
    columns.extend(source_block.get("columns") or [])
    columns.extend(fill_block.get("columns") or [])
    for column in columns:
        name = _normalize_term_name(column)
        if not name:
            continue
        looked = _lookup_term_record(catalog, name=name, sheet=level_name)
        _append(
            {
                "name": name,
                "definition": _first_text(looked.get("definition")),
                "requirement": _first_text(looked.get("requirement")).upper(),
                "package": _first_text(looked.get("package")),
                "example": _first_text(looked.get("example")),
            }
        )
    return terms


def _write_help_sheet(
    wb: Any,
    field_source: Dict[str, Any],
    fill_structure: Optional[Dict[str, Any]] = None,
    term_catalog: Optional[Dict[str, Dict[str, Dict[str, str]]]] = None,
) -> None:
    """Replace the Help sheet with a FAIR-DS-style used-term catalog."""
    try:
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        return

    catalog = _merge_term_catalogs(_load_local_term_catalog(), term_catalog)
    sections: List[Tuple[str, List[Dict[str, str]]]] = []
    for level_name in ISA_WORKBOOK_LEVEL_ORDER:
        terms = _collect_help_terms_for_level(
            level_name, field_source, fill_structure, catalog
        )
        if terms:
            sections.append((level_name, terms))
    if not sections:
        return

    if "Help" in wb.sheetnames:
        wb.remove(wb["Help"])
    ws = wb.create_sheet(title="Help")
    wrap = Alignment(vertical="top", wrap_text=True)
    title_font = Font(size=11)

    row_idx = 1
    for section_i, (level_name, terms) in enumerate(sections):
        if section_i:
            row_idx += 1
        cell = ws.cell(row=row_idx, column=1, value=_help_section_title(level_name))
        cell.font = title_font
        row_idx += 2
        for term in terms:
            ws.cell(row=row_idx, column=1, value=_excel_safe_value(term["name"]))
            ws.cell(row=row_idx, column=2, value=_excel_safe_value(term["definition"]))
            ws.cell(row=row_idx, column=3, value=_excel_safe_value(term["requirement"]))
            ws.cell(row=row_idx, column=4, value=_excel_safe_value(term["package"]))
            ws.cell(row=row_idx, column=5, value=_excel_safe_value(term["example"]))
            for col in range(1, 6):
                ws.cell(row=row_idx, column=col).alignment = wrap
            row_idx += 1

    widths = (42, 70, 16, 28, 48)
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _ensure_help_sheet(
    xlsx_bytes: bytes,
    field_source: Dict[str, Any],
    fill_structure: Optional[Dict[str, Any]] = None,
    term_catalog: Optional[Dict[str, Dict[str, Dict[str, str]]]] = None,
) -> bytes:
    """Rewrite Help so API workbooks are not left with header-only stubs."""
    try:
        import openpyxl
    except ImportError:
        return xlsx_bytes
    try:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    except Exception:
        return xlsx_bytes
    _write_help_sheet(wb, field_source, fill_structure, term_catalog)
    try:
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    except Exception:
        return xlsx_bytes


def _generate_xlsx_local(
    isa_structure: Dict[str, Any],
) -> Optional[bytes]:
    """Generate a complete ``metadata_fairds.xlsx`` from *isa_structure*.

    Uses **openpyxl** only.  Returns raw ``.xlsx`` bytes or ``None`` if
    the library is unavailable or an error occurs.
    """
    try:
        import openpyxl
        from openpyxl.styles import (
            Alignment,
            Border,
            Font,
            PatternFill,
            Side,
        )
        from openpyxl.utils import get_column_letter
    except ImportError:
        logger.warning("openpyxl not installed; cannot generate Excel locally")
        return None

    # ── Style definitions ───────────────────────────────────────
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_align = Alignment(vertical="top", wrap_text=False)
    thin_side = Side(style="thin")
    thin_border = Border(
        left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
    )
    # ── end style definitions ───────────────────────────────────

    # ── Build requirement lookup ─────────────────────────────────
    wb = openpyxl.Workbook()
    # Remove the auto-created empty sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    isa_level_order = list(ISA_WORKBOOK_LEVEL_ORDER)
    # ── ISA data sheets ─────────────────────────────────────────
    for level_name in isa_level_order:
        level_data = isa_structure.get(level_name, {})
        if not isinstance(level_data, dict):
            continue

        columns: List[str] = level_data.get("columns", [])
        if not columns:
            logger.debug("Skipping '%s' sheet: no columns defined", level_name)
            continue

        rows = _resolve_rows(level_data)
        ws = wb.create_sheet(title=_local_sheet_title(level_name, level_data))

        # Row 1: raw FAIR-DS term labels. Requirement is schema metadata and
        # belongs in Help, not in the field name.
        for col_idx, col_name in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        # Row 2+: data rows
        for row_idx, row_data in enumerate(rows, start=2):
            for col_idx, col_name in enumerate(columns, start=1):
                key = col_name.strip().lower()
                value = row_data.get(key)
                if value is not None:
                    cell = ws.cell(
                        row=row_idx, column=col_idx,
                        value=_excel_safe_value(value),
                    )
                    cell.alignment = cell_align
                    cell.border = thin_border

        # Column widths (auto-fit, clamped 10–60)
        for col_idx, col_name in enumerate(columns, start=1):
            col_letter = get_column_letter(col_idx)
            lengths = [len(str(col_name))]
            for row_data in rows:
                val_str = str(row_data.get(col_name.strip().lower(), ""))
                lengths.append(len(val_str))
            width = min(max(max(lengths) + 2, 10), 60)
            ws.column_dimensions[col_letter].width = width

        # Freeze header row
        ws.freeze_panes = "A2"

        logger.debug(
            "Generated '%s' sheet: %d columns x %d rows",
            level_name,
            len(columns),
            len(rows),
        )

    # ── Help sheet (FAIR-DS used-term catalog) ─────────────────
    _write_help_sheet(wb, isa_structure, isa_structure)

    try:
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    except Exception as exc:
        logger.warning("Failed to save Excel workbook: %s", exc)
        return None


# ── Post-process: fill rows into API-generated workbook ─────────


def _level_from_sheet_title(title: str) -> Optional[str]:
    """Resolve bare and FAIR-DS ``level - package`` worksheet titles."""
    normalized = str(title).strip().lower()
    compact = normalized.replace(" ", "")
    for level_name in ISA_LEVEL_ORDER:
        aliases = {level_name, level_name.replace("unit", " unit")}
        for alias in aliases:
            if normalized == alias or normalized.startswith(f"{alias} - "):
                return level_name
        if compact == level_name or compact.startswith(f"{level_name}-"):
            return level_name
    return None


def _local_sheet_title(level_name: str, level_data: Dict[str, Any]) -> str:
    """Match FAIR-DS gold workbooks: ``ISA level - selected package``."""
    packages: List[str] = []
    normalized_packages: set[str] = set()
    for field in level_data.get("fields") or []:
        if not isinstance(field, dict):
            continue
        package = _first_text(
            field.get("package_source"),
            field.get("package"),
            field.get("packageName"),
        )
        normalized = package.lower()
        if package and normalized not in normalized_packages:
            packages.append(package)
            normalized_packages.add(normalized)
    package = next((item for item in packages if item.lower() != "default"), None)
    package = package or (packages[0] if packages else "default")
    return f"{level_name} - {package}"[:31]


def _clear_data_rows(ws: "openpyxl.worksheet.worksheet.Worksheet") -> None:
    """Remove all data rows while keeping the header row intact."""
    if ws.max_row <= 1:
        return
    ws.delete_rows(2, ws.max_row - 1)


def _populate_sheet_from_level(
    ws: "openpyxl.worksheet.worksheet.Worksheet",
    level_data: Dict[str, Any],
    *,
    replace_existing_rows: bool,
) -> int:
    """Write ``columns``/``rows`` from *level_data* into *ws*.

    Returns the number of cell values written.
    """
    columns = [
        str(col).strip()
        for col in (level_data.get("columns") or [])
        if str(col).strip()
    ]
    rows = _resolve_rows(level_data)
    if not columns and not rows:
        return 0

    # The compiled matrix is authoritative. FAIR-DS API workbook templates can
    # contain legacy or server-version-specific columns (for example a contact
    # role column) that are not part of the selected matrix. Rebuild the header
    # to the exact matrix columns while retaining the API's header style.
    if replace_existing_rows and columns:
        header_style = copy(ws.cell(row=1, column=1)._style)
        header_alignment = copy(ws.cell(row=1, column=1).alignment)
        if ws.max_row > 1:
            _clear_data_rows(ws)
        if ws.max_column:
            ws.delete_cols(1, ws.max_column)
        for col_idx, col_name in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell._style = copy(header_style)
            cell.alignment = copy(header_alignment)

    header_col = _build_header_column_map(ws)
    if not header_col:
        # Brand-new sheet (or headerless): seed headers from columns, then rows.
        for col_idx, col_name in enumerate(columns, start=1):
            ws.cell(row=1, column=col_idx, value=col_name)
            header_col[col_name.strip().lower()] = col_idx
    else:
        # Extend headers with any columns present in the authoritative matrix.
        for col_name in columns:
            normalized = col_name.strip().lower()
            if normalized in header_col:
                continue
            col_idx = ws.max_column + 1
            ws.cell(row=1, column=col_idx, value=col_name)
            header_col[normalized] = col_idx

    if not columns:
        for row_data in rows:
            for key in row_data.keys():
                normalized = str(key).strip().lower()
                if not normalized or normalized in header_col:
                    continue
                col_idx = ws.max_column + 1
                ws.cell(row=1, column=col_idx, value=normalized)
                header_col[normalized] = col_idx

    if not rows:
        return 0

    if replace_existing_rows and ws.max_row > 1:
        _clear_data_rows(ws)

    # Only skip when the sheet already has data and the caller asked not to
    # replace (legacy empty-sheet fill).
    if (not replace_existing_rows) and ws.max_row > 1:
        return 0

    total_filled = 0
    for row_idx, row_data in enumerate(rows, start=2):
        if not isinstance(row_data, dict):
            continue
        for key, value in row_data.items():
            col = header_col.get(str(key).strip().lower())
            if col is not None and value is not None:
                val_str = _excel_safe_value(value)
                if not val_str.strip():
                    # Do not overwrite a previously written non-empty cell value with an empty string
                    existing = ws.cell(row=row_idx, column=col).value
                    if existing and str(existing).strip():
                        continue
                ws.cell(row=row_idx, column=col, value=val_str)
                total_filled += 1
    return total_filled


def _fill_missing_data_rows(
    xlsx_bytes: bytes,
    isa_structure: Dict[str, Any],
) -> bytes:
    """Post-process Excel: complete ISA sheets from the authoritative matrix.

    The FAIR-DS ``POST /api/isa`` endpoint commonly returns only
    ``investigation`` / ``study`` sheets (sometimes with placeholder study
    rows).  This function:

    1. Creates any missing ISA sheets that have columns/rows in
       *isa_structure* (observationunit / sample / assay).
    2. Fills empty sheets and **replaces** API placeholder data rows with
       the compiled ``columns`` × ``rows`` matrix — one Excel row per entity.

    Returns
    -------
    bytes
        The (possibly modified) ``.xlsx`` bytes.
    """
    try:
        import openpyxl
    except ImportError:
        return xlsx_bytes

    try:
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    except Exception:
        return xlsx_bytes

    # FAIR-DS ISOSA output has five data levels. Some older API versions add
    # a Person worksheet; contributor fields belong as repeated Investigation
    # rows in the curated workbook convention used by this project.
    for name in list(wb.sheetnames):
        normalized = str(name).strip().lower()
        if normalized == "person" or normalized.startswith("person - "):
            wb.remove(wb[name])

    existing_by_level: Dict[str, str] = {}
    for name in wb.sheetnames:
        level_name = _level_from_sheet_title(name)
        if level_name and level_name not in existing_by_level:
            existing_by_level[level_name] = name

    for level_name in ISA_WORKBOOK_LEVEL_ORDER:
        level_data = isa_structure.get(level_name)
        if not isinstance(level_data, dict):
            continue
        columns = level_data.get("columns") or []
        rows = _resolve_rows(level_data)
        fields = level_data.get("fields") or []
        if not (columns or rows or fields):
            continue

        sheet_title = existing_by_level.get(level_name)
        created = False
        if sheet_title is None:
            sheet_title = _local_sheet_title(level_name, level_data)
            ws = wb.create_sheet(title=sheet_title)
            existing_by_level[level_name] = sheet_title
            created = True
        else:
            ws = wb[sheet_title]

        # Always prefer the compiled matrix over API placeholder rows.
        # Newly created sheets start empty; existing API sheets may already
        # contain incomplete investigation/study stubs that must be replaced.
        total_filled = _populate_sheet_from_level(
            ws,
            level_data,
            replace_existing_rows=True,
        )
        if total_filled or created:
            logger.debug(
                "%s ISA sheet '%s' from matrix (%d values)",
                "Created" if created else "Filled",
                sheet_title,
                total_filled,
            )

    # Canonical order: investigation → study → OU → sample → assay → Help.
    ordered_titles: List[str] = []
    for level_name in ISA_WORKBOOK_LEVEL_ORDER:
        title = existing_by_level.get(level_name)
        if title and title in wb.sheetnames:
            ordered_titles.append(title)
    for name in wb.sheetnames:
        if name not in ordered_titles:
            ordered_titles.append(name)
    for target_idx, title in enumerate(ordered_titles):
        current_idx = wb.sheetnames.index(title)
        if current_idx != target_idx:
            wb.move_sheet(title, offset=target_idx - current_idx)

    try:
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    except Exception:
        return xlsx_bytes


# ── Public API ──────────────────────────────────────────────────


def try_export_fairds_metadata_excel(
    output_dir: Path,
    *,
    fair_ds_api_url: Optional[str] = None,
) -> Optional[Path]:
    """Write ``metadata_fairds.xlsx`` when ``metadata.json`` has fillable ISA data.

    Workflow
    --------
    1. Read ``metadata.json``, extract ``isa_structure``.
    2. Apply entity-splitting heuristics to expand single merged rows.
    3. If a FAIR-DS API URL is available **and** the server is reachable,
       call ``POST /api/isa`` with ``fields`` (so Help can list metadata
       types) to obtain a workbook with column headers. Then fill data
       rows from the compiled ``columns`` × ``rows`` matrix and rewrite
       Help from field + package catalogs.
    4. If the API is unavailable or not configured, generate the workbook
       entirely locally with ``openpyxl``.
    5. Return the path to the written ``.xlsx`` file, or ``None`` on
       failure / missing data.

    Parameters
    ----------
    output_dir : Path
        Directory containing ``metadata.json`` (and where the ``.xlsx``
        will be written).
    fair_ds_api_url : str, optional
        Override the configured FAIR-DS API URL.  ``None`` (default) uses
        ``config.fair_ds_api_url``.

    Returns
    -------
    Path or None
        Absolute path to the written ``metadata_fairds.xlsx`` file.
    """
    from ..config import config
    from ..output_paths import (
        FAIRDS_METADATA_EXCEL_FILENAME,
        deliverables_dir,
        resolve_isa_values_read_path,
        resolve_metadata_output_read_path,
    )
    from ..services.fair_data_station import FAIRDataStationClient

    url = config.fair_ds_api_url if fair_ds_api_url is None else fair_ds_api_url

    meta_path = resolve_metadata_output_read_path(Path(output_dir))
    if not meta_path:
        logger.debug(
            "FAIR-DS Excel export skipped: no metadata.json in %s",
            output_dir,
        )
        return None

    try:
        with open(meta_path, encoding="utf-8") as fh:
            data: Dict[str, Any] = json.load(fh)
    except (OSError, json.JSONDecodeError) as exc:
        logger.warning(
            "FAIR-DS Excel export skipped: cannot read %s: %s",
            meta_path,
            exc,
        )
        return None

    isa_structure = data.get("isa_structure")
    if not isinstance(isa_structure, dict) or not isa_structure:
        logger.debug(
            "FAIR-DS Excel export skipped: missing or empty isa_structure"
        )
        return None

    # Prefer dedicated compiled matrix (§12.1). Do not re-split when the
    # sidecar already carries columns×rows — that matrix is the sole projection.
    isa_values_path = resolve_isa_values_read_path(Path(output_dir))
    fill_structure: Dict[str, Any]
    used_compiled_sidecar = False
    if isa_values_path and isa_values_path.exists():
        try:
            with open(isa_values_path, encoding="utf-8") as fh:
                isa_values = json.load(fh)
            if isinstance(isa_values, dict) and any(
                isinstance(block, dict) and (block.get("columns") or block.get("rows"))
                for block in isa_values.values()
            ):
                fill_structure = isa_values
                used_compiled_sidecar = True
            else:
                fill_structure = split_entities_in_isa_structure(isa_structure)
        except (OSError, json.JSONDecodeError):
            fill_structure = split_entities_in_isa_structure(isa_structure)
    else:
        fill_structure = split_entities_in_isa_structure(isa_structure)

    # Check for fillable content.
    has_content = False
    for _level, block in fill_structure.items():
        if isinstance(block, dict) and (
            block.get("fields") or block.get("rows")
        ):
            has_content = True
            break
    if not has_content:
        logger.debug(
            "FAIR-DS Excel export skipped: no fields/rows under isa_structure"
        )
        return None

    # Keep field metadata from metadata.json even when data rows come from
    # the compiled isa_values sidecar. FAIR-DS Help is populated from fields.
    if not used_compiled_sidecar:
        isa_structure = split_entities_in_isa_structure(isa_structure)
        fill_structure = isa_structure
    export_structure = _merge_fields_into_fill(fill_structure, isa_structure)

    # ── Step 2: Generate Excel ─────────────────────────────────
    xlsx_bytes: Optional[bytes] = None
    api_used = False
    api_term_catalog: Optional[Dict[str, Dict[str, Dict[str, str]]]] = None

    if url and str(url).strip():
        base = str(url).strip().rstrip("/")
        client = FAIRDataStationClient(base)
        if client.is_available():
            try:
                xlsx_bytes = client.generate_excel_from_isa_structure(
                    _slim_isa_for_api(export_structure)
                )
                api_used = True
                try:
                    api_term_catalog = _index_fairds_terms(client.get_terms())
                except Exception as exc:
                    logger.debug("FAIR-DS term catalog unavailable for Help: %s", exc)
            except Exception as exc:
                logger.warning(
                    "FAIR-DS API Excel generation failed: %s", exc
                )
        else:
            logger.debug(
                "FAIR-DS API not reachable at %s; falling back to local generation",
                base,
            )
    else:
        logger.debug(
            "FAIR-DS API URL not set; using local Excel generation"
        )

    if xlsx_bytes is not None and api_used:
        # The API returned headers (and perhaps a few data rows);
        # fill the remaining data locally, then restore a complete Help catalog.
        xlsx_bytes = _fill_missing_data_rows(xlsx_bytes, fill_structure)
        xlsx_bytes = _ensure_help_sheet(
            xlsx_bytes,
            export_structure,
            fill_structure,
            api_term_catalog,
        )
    elif xlsx_bytes is None:
        # API was not used or failed — generate entirely locally.
        xlsx_bytes = _generate_xlsx_local(export_structure)
        if xlsx_bytes is not None:
            xlsx_bytes = _ensure_help_sheet(
                xlsx_bytes,
                export_structure,
                fill_structure,
            )

    if xlsx_bytes is None:
        logger.warning(
            "FAIR-DS Excel export failed: no xlsx bytes generated"
        )
        return None

    out = deliverables_dir(Path(output_dir)) / FAIRDS_METADATA_EXCEL_FILENAME
    out.parent.mkdir(parents=True, exist_ok=True)

    try:
        out.write_bytes(xlsx_bytes)
    except OSError as exc:
        logger.warning(
            "FAIR-DS Excel export could not write %s: %s",
            out,
            exc,
        )
        return None

    logger.info(
        "FAIR-DS metadata Excel written: %s (%s bytes)",
        out.name,
        len(xlsx_bytes),
    )
    return out


__all__ = [
    "try_export_fairds_metadata_excel",
    "split_entities_in_isa_structure",
    "_generate_xlsx_local",
    "_fill_missing_data_rows",
    "_resolve_rows",
]
